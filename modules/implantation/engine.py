"""Moteur du module Suivi implantation : lecture T1 et stocks ERP, statuts, taux, cessions et export (logique reprise de la v7.1)."""
import io
import re
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from datetime import date

# ─────────────────────────────────────────────────────────────────────────────
# CONSTANTES
# ─────────────────────────────────────────────────────────────────────────────
TODAY      = pd.Timestamp(date.today())
TODAY_STR  = date.today().strftime("%d %b %Y")
TODAY_FILE = date.today().strftime("%Y%m%d")

STOCK_COLS_REQUIRED = [
    "Site", "Libellé site", "Code article", "Libellé article",
    "Code etat", "Nouveau stock", "Ral", "Pcb",
    "Code marketing", "Nom fourn.", "Libellé rayon", "Libellé famille",
    "Date dernière entrée", "Type saisonnalité",
]

ETAT_ACTIF    = "2"
ETAT_PURGE    = "P"
ETAT_ANOMALIE = {"B", "S", "F", "6", "5", "1"}
ETAT_LABEL    = {
    "B": "Rayon générique", "S": "Suspendu",
    "F": "Fin de vie",      "6": "Déréférencé", "5": "Autre",
}

ALERTES = {
    "✅ Implanté":        "#34C759",
    "🔵 Appro en cours":  "#007AFF",
    "🛒 Passer commande": "#FF9500",
    "🚩 Anomalie référ.": "#FF9500",
}
ACTION_LABEL = {
    "✅ Implanté":        "—",
    "🔵 Appro en cours":  "Accélérer livraison",
    "🛒 Passer commande": "Passer commande fournisseur",
    "🚩 Anomalie référ.": "Vérifier référencement magasin",
}


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────
def fmt_n(n) -> str:
    try:
        return f"{int(n):,}".replace(",", "\u202f")
    except Exception:
        return str(n)

def color_taux(t: float) -> str:
    if t >= 80: return "#34C759"
    if t >= 50: return "#FF9500"
    return "#FF3B30"

def scorecard_cls(t: float) -> str:
    if t >= 80: return "ok"
    if t >= 50: return "warn"
    return "ko"

def extract_site_from_filename(filename: str) -> str | None:
    m = re.search(r'_(\d{5})_', filename)
    return m.group(1) if m else None

def extract_date_from_filename(filename: str) -> pd.Timestamp:
    m8 = re.search(r'_(\d{8})_', filename)
    if m8:
        for fmt in ('%Y%m%d', '%d%m%Y'):
            try:
                return pd.Timestamp(pd.to_datetime(m8.group(1), format=fmt))
            except Exception:
                continue
    return TODAY

def _sem_sort(s) -> int:
    cleaned = re.sub(r"[Ss]", "", str(s).strip())
    return int(cleaned) if cleaned.isdigit() else 99

def get_alerte(row) -> str:
    etat  = str(row.get("Code etat", "")).strip()
    stock = row.get("Nouveau stock")
    ral   = float(row.get("Ral", 0) or 0)
    if not etat or etat in ("nan", ETAT_PURGE) or etat in ETAT_ANOMALIE:
        return "🚩 Anomalie référ."
    if pd.isna(stock):
        return "🚩 Anomalie référ."
    if float(stock) != 0:
        return "✅ Implanté"
    return "🔵 Appro en cours" if ral > 0 else "🛒 Passer commande"


# ─────────────────────────────────────────────────────────────────────────────
# PARSERS
# ─────────────────────────────────────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def parse_t1(file_bytes: bytes, filename: str):
    buf = io.BytesIO(file_bytes)
    try:
        df = pd.read_excel(buf, header=None, dtype=str) \
             if filename.lower().endswith((".xlsx", ".xls")) \
             else pd.read_csv(buf, header=None, sep=None, engine="python",
                              encoding="latin1", dtype=str, on_bad_lines="skip")
    except Exception as e:
        return None, f"Lecture T1 : {e}"

    first = str(df.iloc[0, 0]).strip().replace(".0", "")
    has_header = not first.isdigit()
    if has_header:
        df.columns = df.iloc[0].astype(str).str.strip().str.upper()
        df = df.iloc[1:].reset_index(drop=True)
    else:
        df.columns = ["ARTICLE"] + [f"_COL{i}" for i in range(1, len(df.columns))]

    df.columns = (df.columns.astype(str).str.strip().str.upper()
                  .str.replace("\ufeff", "", regex=False)
                  .str.replace("\xa0", " ", regex=False))

    if "ARTICLE" not in df.columns:
        return None, "Colonne ARTICLE introuvable dans le fichier T1"

    df["SKU"] = (df["ARTICLE"].astype(str).str.strip()
                 .str.replace(r"\.0$", "", regex=True).str.zfill(8).str[:8])
    df = df[df["SKU"].str.match(r"^\d{8}$", na=False)].drop_duplicates("SKU").copy()

    for col, val in {"LIBELLÉ ARTICLE": "", "LIBELLÉ FOURNISSEUR ORIGINE": "",
                     "MODE APPRO": "", "SEMAINE RECEPTION": "", "DATE LIV.": ""}.items():
        if col not in df.columns:
            df[col] = val

    df["SEMAINE RECEPTION"] = df["SEMAINE RECEPTION"].astype(str).str.strip().replace("nan", "")
    def _sem_to_num(s):
        if s is None or (isinstance(s, float) and np.isnan(s)):
            return 99
        cleaned = re.sub(r"[Ss]", "", str(s).strip())
        return int(cleaned) if cleaned.isdigit() else 99

    df["SEM_NUM"] = df["SEMAINE RECEPTION"].apply(_sem_to_num)
    df["ORIGINE"] = df["MODE APPRO"].apply(
        lambda m: "IM" if "IMPORT" in str(m).upper() else "LO"
    )
    return df, None


@st.cache_data(show_spinner=False)
def parse_stock_file(file_bytes: bytes, filename: str, sku_scope: tuple):
    try:
        df = pd.read_csv(io.BytesIO(file_bytes), sep=";", encoding="latin1",
                         low_memory=False, on_bad_lines="skip", dtype=str)
    except Exception as e:
        return None, f"Lecture {filename} : {e}"

    missing = [c for c in STOCK_COLS_REQUIRED if c not in df.columns]
    if missing:
        return None, f"{filename} — colonnes manquantes : {', '.join(missing)}"

    for col in ["Libellé site", "Libellé article", "Nom fourn.", "Libellé rayon",
                "Code etat", "Code marketing", "Type saisonnalité"]:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip()

    df["SKU"]       = (df["Code article"].astype(str).str.strip()
                       .str.replace(r"\.0$", "", regex=True).str.zfill(8).str[:8])
    df["Code site"] = df["Site"].astype(str).str.strip()

    site_from_name = extract_site_from_filename(filename)
    if site_from_name and df["Code site"].isin(["nan", ""]).all():
        df["Code site"] = site_from_name

    df = df[df["Code etat"] != ETAT_PURGE].copy()
    if sku_scope:
        df = df[df["SKU"].isin(sku_scope)].copy()
    df = df[df["SKU"].str.match(r"^\d{8}$", na=False)].copy()

    df["Nouveau stock"]  = pd.to_numeric(df["Nouveau stock"], errors="coerce")
    df["Ral"]            = pd.to_numeric(df["Ral"], errors="coerce").fillna(0).astype(int)
    df["Pcb"]            = pd.to_numeric(df["Pcb"], errors="coerce").fillna(1)
    df["Origine"]        = df["Code marketing"].apply(lambda m: "IM" if str(m).strip().upper() == "IM" else "LO")
    df["_date_fichier"]  = extract_date_from_filename(filename)

    return df, None


def consolidate_stock(uploaded_files, sku_scope: tuple):
    frames, errors = [], []
    for f in uploaded_files:
        df_site, err = parse_stock_file(f.read(), f.name, sku_scope)
        if err:
            errors.append(err)
        elif df_site is not None and not df_site.empty:
            frames.append(df_site)
    return (pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()), errors



# ─────────────────────────────────────────────────────────────────────────────
# DATASET ET MÉTRIQUES
# ─────────────────────────────────────────────────────────────────────────────
def build(t1_df, df_stock, orig_sel, sem_sel):
    """Grille magasin × SKU du périmètre T1 filtré, avec statut, action et métriques réseau."""
    site_ref = (df_stock[["Code site", "Libellé site"]].drop_duplicates("Code site")
                .set_index("Code site")["Libellé site"].to_dict())
    all_codes = sorted(site_ref.keys())
    mask_t1 = t1_df["ORIGINE"].isin(orig_sel)
    if sem_sel:
        mask_t1 = mask_t1 & t1_df["SEMAINE RECEPTION"].isin(sem_sel)
    t1_scope = t1_df[mask_t1].copy()
    sku_scope = t1_scope["SKU"].unique()
    if len(sku_scope) == 0:
        return None

    sku_dans_stock = set(df_stock[df_stock["Code site"].isin(all_codes)]["SKU"].unique())
    n_sku = len(sku_scope)
    n_sku_trouves = len([s for s in sku_scope if s in sku_dans_stock])
    n_mag = len(all_codes)

    stk_filt = df_stock[df_stock["Code site"].isin(all_codes) & df_stock["SKU"].isin(sku_scope)].copy()
    grid = pd.DataFrame(pd.MultiIndex.from_product([all_codes, sku_scope], names=["Code site", "SKU"]).tolist(),
                        columns=["Code site", "SKU"])
    KEEP = ["Code site", "SKU", "Nouveau stock", "Ral", "Code etat",
            "Origine", "Libellé article", "Nom fourn.", "Libellé rayon", "Libellé famille", "Pcb"]
    merged = grid.merge(stk_filt[[c for c in KEEP if c in stk_filt.columns]], on=["Code site", "SKU"], how="left")
    t1_ref = t1_scope.set_index("SKU")[[
        "LIBELLÉ ARTICLE", "LIBELLÉ FOURNISSEUR ORIGINE", "MODE APPRO", "SEMAINE RECEPTION", "DATE LIV.", "ORIGINE", "SEM_NUM"
    ]].rename(columns={
        "LIBELLÉ ARTICLE": "T1_lib", "LIBELLÉ FOURNISSEUR ORIGINE": "Fournisseur T1",
        "MODE APPRO": "Mode Appro", "SEMAINE RECEPTION": "Sem. Réception",
        "DATE LIV.": "Date Livraison", "ORIGINE": "Origine_T1", "SEM_NUM": "SEM_NUM",
    })
    merged = merged.merge(t1_ref.reset_index(), on="SKU", how="left")
    merged["Libellé article"] = merged["Libellé article"].fillna("").astype(str)
    merged["Libellé article"] = merged.apply(
        lambda r: r["Libellé article"] if r["Libellé article"] else r.get("T1_lib", ""), axis=1)
    merged["Origine"] = merged.apply(
        lambda r: r.get("Origine") if pd.notna(r.get("Origine")) and str(r.get("Origine")) not in ("nan", "")
        else r.get("Origine_T1", "LO"), axis=1)
    merged.drop(columns=["T1_lib", "Origine_T1"], errors="ignore", inplace=True)
    merged["Magasin"] = merged["Code site"].map(site_ref).fillna(merged["Code site"])
    merged["Code etat"] = merged["Code etat"].fillna("").astype(str)
    merged["Ral"] = pd.to_numeric(merged["Ral"], errors="coerce").fillna(0)
    merged["Pcb"] = pd.to_numeric(merged["Pcb"], errors="coerce").fillna(1)
    merged["Alerte"] = merged.apply(get_alerte, axis=1)
    merged["Action"] = merged["Alerte"].map(ACTION_LABEL)

    merged_actif = merged[merged["Code etat"].str.strip() == ETAT_ACTIF]
    n_base_taux = len(merged_actif)
    n_impl = int((merged["Alerte"] == "✅ Implanté").sum())

    def taux_mag(mag):
        dm = merged_actif[merged_actif["Magasin"] == mag]
        return int((dm["Alerte"] == "✅ Implanté").sum() / len(dm) * 100) if len(dm) else 0

    pivot_mag = (merged.groupby(["Magasin", "Alerte"]).size().unstack(fill_value=0)
                 .reindex(columns=list(ALERTES.keys()), fill_value=0).reset_index())
    pivot_mag.columns.name = None
    pivot_mag["Taux (%)"] = pivot_mag["Magasin"].apply(taux_mag)

    rayon_pivot = pd.DataFrame()
    if "Libellé rayon" in merged.columns:
        try:
            rayon_pivot = (merged_actif.groupby(["Libellé rayon", "Magasin"])
                           .apply(lambda x: int((x["Alerte"] == "✅ Implanté").sum() / len(x) * 100) if len(x) else 0)
                           .reset_index(name="Taux (%)")
                           .pivot(index="Libellé rayon", columns="Magasin", values="Taux (%)")
                           .fillna(0).astype(int))
        except Exception:
            rayon_pivot = pd.DataFrame()

    non_impl = merged[merged["Alerte"].isin(["🛒 Passer commande", "🔵 Appro en cours"])]
    cnt = non_impl.groupby("SKU")["Magasin"].count()
    return dict(
        merged=merged, site_ref=site_ref, all_codes=all_codes, t1_scope=t1_scope, sku_scope=sku_scope,
        n_sku=n_sku, n_sku_trouves=n_sku_trouves, n_sku_absents=n_sku - n_sku_trouves, n_mag=n_mag,
        n_impl=n_impl, n_appro=int((merged["Alerte"] == "🔵 Appro en cours").sum()),
        n_cmd=int((merged["Alerte"] == "🛒 Passer commande").sum()),
        n_anomalie=int((merged["Alerte"] == "🚩 Anomalie référ.").sum()),
        taux_reseau=int(n_impl / n_base_taux * 100) if n_base_taux else 0,
        n_sku_im=int((t1_scope["ORIGINE"] == "IM").sum()), n_sku_lo=int((t1_scope["ORIGINE"] == "LO").sum()),
        pivot_mag=pivot_mag, rayon_pivot=rayon_pivot, sku_rupt=cnt[cnt == n_mag].index.tolist(),
    )


# ─────────────────────────────────────────────────────────────────────────────
# CESSIONS INTER-MAGASINS (réserve de 2 PCB chez le cédant)
# ─────────────────────────────────────────────────────────────────────────────
def _sku_stock(df_stock, sku):
    d = df_stock[df_stock["SKU"] == sku].copy()
    d["Nouveau stock"] = pd.to_numeric(d["Nouveau stock"], errors="coerce").fillna(0)
    d["Pcb"] = pd.to_numeric(d.get("Pcb", pd.Series(1, index=d.index)), errors="coerce").fillna(1).clip(lower=1)
    d["Reserve_2pcb"] = (d["Pcb"] * 2).astype(int)
    return d


def auto_cessions(df_stock, site_ref, sku_rupt):
    """Pour les ruptures communes réseau : meilleur cédant hors réseau filtré."""
    out = []
    for sku in sku_rupt:
        d = _sku_stock(df_stock, sku)
        best = d[d["Nouveau stock"] > d["Reserve_2pcb"]].sort_values("Nouveau stock", ascending=False)
        lib = d["Libellé article"].iloc[0] if len(d) else sku
        if not best.empty:
            b = best.iloc[0]
            out.append({"SKU": sku, "Libellé": lib, "Cédant": site_ref.get(b["Code site"], b["Code site"]),
                        "Stock cédant": int(b["Nouveau stock"]), "Qté cessible": int(b["Nouveau stock"]) - int(b["Reserve_2pcb"])})
    return pd.DataFrame(out)


def manual_cessions(df_stock, site_ref, all_codes, sku_scope, mag_detresse, seuil_det, min_1pcb):
    """Plan de cessions vers les magasins en détresse (seulement les cessions possibles)."""
    det_codes = [c for c in all_codes if site_ref.get(c, c) in mag_detresse]
    ced_codes = [c for c in all_codes if c not in det_codes]
    sugg = []
    for sku in sku_scope:
        d = _sku_stock(df_stock, sku)
        if d.empty:
            continue
        lib = d["Libellé article"].iloc[0] if "Libellé article" in d.columns else sku
        det_rows = d[d["Code site"].isin(det_codes) & (d["Nouveau stock"] <= seuil_det)]
        if det_rows.empty:
            continue
        ced = d[d["Code site"].isin(ced_codes) & (d["Nouveau stock"] > d["Reserve_2pcb"])].sort_values("Nouveau stock", ascending=False)
        for _, dr in det_rows.iterrows():
            if ced.empty:
                continue
            b = ced.iloc[0]
            qty = int(b["Nouveau stock"]) - int(b["Reserve_2pcb"])
            if qty < 1:
                continue
            sugg.append({"SKU": sku, "Libellé": lib, "Magasin détresse": site_ref.get(dr["Code site"], dr["Code site"]),
                         "Stock détresse": int(dr["Nouveau stock"]), "Cédant": site_ref.get(b["Code site"], b["Code site"]),
                         "Stock cédant": int(b["Nouveau stock"]), "Réserve (2 PCB)": int(b["Reserve_2pcb"]), "Qté cessible": qty})
    df = pd.DataFrame(sugg)
    if df.empty:
        return df
    if min_1pcb:
        df = df[df["Qté cessible"] >= (df["Réserve (2 PCB)"] / 2).clip(lower=1)]
    return df.sort_values("Qté cessible", ascending=False).reset_index(drop=True)


# ─────────────────────────────────────────────────────────────────────────────
# EXPORTS EXCEL (valeurs uniquement)
# ─────────────────────────────────────────────────────────────────────────────
COLS_DET = ["Magasin", "SKU", "Libellé article", "Origine", "Code etat", "Nouveau stock", "Ral",
            "Mode Appro", "Sem. Réception", "Fournisseur T1", "Alerte", "Action"]


def build_export(merged, pivot_mag) -> bytes:
    from openpyxl.styles import Alignment as XAlignment, Font as XFont, PatternFill as XFill
    from openpyxl.utils import get_column_letter as gcl
    fills = {"✅ Implanté": ("D1FAE5", "065F46"), "🔵 Appro en cours": ("DBEAFE", "1D4ED8"),
             "🛒 Passer commande": ("FEF3C7", "92400E"), "🚩 Anomalie référ.": ("FFFDE7", "795548")}
    buf = io.BytesIO()
    cols = [c for c in COLS_DET if c in merged.columns]
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        cols_s = ["Magasin"] + [c for c in ALERTES if c in pivot_mag.columns] + ["Taux (%)"]
        pivot_mag[cols_s].sort_values("Taux (%)", ascending=False).to_excel(writer, sheet_name="Synthèse Réseau", index=False)
        merged[cols].to_excel(writer, sheet_name="Détail Complet", index=False)
        merged[merged["Alerte"] != "✅ Implanté"][cols].sort_values(["Alerte", "Magasin"]).to_excel(
            writer, sheet_name="Alertes & Actions", index=False)
        FH, FT = XFill("solid", fgColor="24235C"), XFont(bold=True, color="FFFFFF", name="Arial", size=11)
        for ws in writer.book.worksheets:
            for cell in ws[1]:
                cell.fill, cell.font, cell.alignment = FH, FT, XAlignment(horizontal="center")
            for col in ws.columns:
                ws.column_dimensions[gcl(col[0].column)].width = min(max((len(str(c.value)) for c in col if c.value), default=10) + 4, 50)
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if str(cell.value) in fills:
                        bg, fg = fills[str(cell.value)]
                        cell.fill, cell.font = XFill("solid", fgColor=bg), XFont(color=fg, name="Arial", size=10)
    return buf.getvalue()


def build_cessions_export(df_cess) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df_cess.to_excel(w, sheet_name="Plan Cessions", index=False)
    return buf.getvalue()
