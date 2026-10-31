"""Export Excel du flash (valeurs uniquement, sans formules), charte CFAO Consumer.

Onglets : Flash (réseau) · PGC · Non-alimentaire · Frais · Retours attendus.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config as C
from .engine import short_rayon, short_site

NAVY, ORANGE, GREY, RED, TEAL, INK = "24235C", "F09F39", "F3F4F8", "E62244", "08868D", "363C5A"
F = "Arial"

PCT = '+0.0%;[Red]-0.0%;0.0%'
PTS = '+0.0" pt";[Red]-0.0" pt";0.0" pt"'
MIO = '#,##0.0,," M";[Red]-#,##0.0,," M"'
DMIO = '+#,##0.0,," M";[Red]-#,##0.0,," M";0.0" M"'
TAUX = '0.0%'

# colonnes standard d'un bloc KPI : (en-tête, clé, format)
KPI_COLS = [
    ("CA", "ca", MIO), ("CA vs N-1", "vs_n1", PCT),
    ("Marge", "marge", MIO), ("Marge vs N-1", "marge_vs_n1", PCT), ("Écart marge", "dmarge", DMIO),
    ("Taux de marge", "taux", TAUX), ("Écart taux", "dpts", PTS),
    ("Effet volume", "effet_volume", DMIO), ("Effet taux", "effet_taux", DMIO),
]
TRAFIC_COLS = [("Débit vs N-1", "debit_vs_n1", PCT), ("Panier vs N-1", "panier_vs_n1", PCT)]

_thin = Side(style="thin", color="DADDE6")


class Sheet:
    def __init__(self, ws, title, subtitle):
        self.ws = ws
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 2
        ws["B1"] = title
        ws["B1"].font = Font(name=F, size=16, bold=True, color="FFFFFF")
        ws["B2"] = subtitle
        ws["B2"].font = Font(name=F, size=10, color="FFFFFF")
        for r in (1, 2):
            for c in range(1, 16):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=NAVY)
        ws.row_dimensions[1].height = 26
        self.r = 4

    def section(self, text):
        c = self.ws.cell(self.r, 2, text)
        c.font = Font(name=F, size=12, bold=True, color=NAVY)
        c.border = Border(bottom=Side(style="medium", color=ORANGE))
        self.r += 1

    def note(self, text):
        c = self.ws.cell(self.r, 2, text)
        c.font = Font(name=F, size=9, italic=True, color="6B7090")
        self.r += 1

    def table(self, head, rows, fmts=None, red_flags=None):
        """rows : listes de valeurs. fmts : formats par colonne (None = texte)."""
        ws = self.ws
        for j, h in enumerate(head):
            c = ws.cell(self.r, 2 + j, h)
            c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=ORANGE)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[self.r].height = 28
        self.r += 1
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                c = ws.cell(self.r, 2 + j, None if v is None else (float(v) if hasattr(v, "dtype") else v))
                c.font = Font(name=F, size=10, color=INK, bold=(j == 0))
                c.border = Border(bottom=_thin)
                if i % 2:
                    c.fill = PatternFill("solid", fgColor=GREY)
                if fmts and j < len(fmts) and fmts[j]:
                    c.number_format = fmts[j]
                    c.alignment = Alignment(horizontal="right")
                else:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            if red_flags and red_flags[i]:
                ws.cell(self.r, 2).font = Font(name=F, size=10, bold=True, color=RED)
            self.r += 1
        self.r += 1

    def widths(self, w):
        for j, x in enumerate(w):
            self.ws.column_dimensions[get_column_letter(2 + j)].width = x


def _kpi_row(label, k, trafic=False):
    cols = KPI_COLS + (TRAFIC_COLS if trafic else [])
    return [label] + [k.get(key) for _, key, _ in cols]


def _kpi_head(first, trafic=False):
    cols = KPI_COLS + (TRAFIC_COLS if trafic else [])
    return [first] + [h for h, _, _ in cols], [None] + [f for _, _, f in cols]


def _sem_row(label, k):
    s = k.get("sem") or {}
    return [label, s.get("ca"), s.get("vs_n1"), s.get("marge"), s.get("marge_vs_n1"), s.get("taux"), s.get("dpts"), s.get("vs_budget")]


SEM_HEAD = ["", "CA semaine", "vs N-1", "Marge semaine", "Marge vs N-1", "Taux", "Écart taux", "vs budget"]
SEM_FMT = [None, MIO, PCT, MIO, PCT, TAUX, PTS, PCT]


def build_excel(R: dict) -> bytes:
    d = R["date"].strftime("%d/%m/%Y")
    wb = Workbook()

    # ------------------------------------------------------------ Flash réseau
    sh = Sheet(wb.active, "Flash commercial", f"Ventes du {d} · HT · comparaison au même jour de l'an dernier")
    wb.active.title = "Flash"
    sh.section("L'essentiel")
    for m in R["messages"]:
        sh.note("• " + m)
    sh.r += 1
    head, fmt = _kpi_head("Périmètre", trafic=True)
    rows = [_kpi_row("Réseau", R["reseau"], True), _kpi_row("Hors magasins convertis", R["hors_convertis"], False)]
    sh.section("Réseau – hier"); sh.table(head, rows, fmt)

    head, fmt = _kpi_head("Secteur")
    sh.section("Secteurs – hier")
    sh.table(head, [_kpi_row(n, k) for n, k in R["secteurs"].items()], fmt)
    sh.section("Secteurs – semaine en cours")
    sh.table(SEM_HEAD, [_sem_row("Réseau", {"sem": R["reseau_sem"]})] + [_sem_row(n, k) for n, k in R["secteurs"].items()], SEM_FMT)

    head, fmt = _kpi_head("Format", trafic=True)
    sh.section("Formats")
    sh.table(head, [_kpi_row(n, k, True) for n, k in R["formats"].items()], fmt)
    head, fmt = _kpi_head("Zone de pricing", trafic=True)
    sh.section("Zones de pricing")
    sh.table(head, [_kpi_row(n, k, True) for n, k in R["zones"].items()], fmt)

    head, fmt = _kpi_head("Magasin", trafic=True)
    sh.section("Magasins (en rouge : en alerte)")
    sh.table(["Magasin", "Format", "Zone"] + head[1:],
             [[x["nom"], x["format"], x["zone"]] + _kpi_row("", x, True)[1:] for x in R["sites"]],
             [None, None, None] + fmt[1:], red_flags=[x["alerte"] for x in R["sites"]])
    sh.section("Magasins convertis")
    sh.table(["Magasin", "Converti le"] + head[1:],
             [[x["nom"], x["converti_le"].strftime("%d/%m/%Y") if x.get("converti_le") else ""] + _kpi_row("", x, True)[1:] for x in R["convertis"]],
             [None, None] + fmt[1:])

    sh.section("Alertes rayon × magasin (triées par perte de marge)")
    sh.note(f"Seuils : CA vs N-1 < {C.SEUILS['ca_vs_n1']:.0%}, taux < {C.SEUILS['taux_pts']:.0f} pts, marge vs N-1 < {C.SEUILS['marge_vs_n1']:.0%} ; CA N-1 ≥ {C.SEUILS['ca_n1_min']:,} FCFA".replace(",", " "))
    sh.table(["Rayon", "Magasin", "Secteur", "CA", "CA vs N-1", "Marge", "Écart marge", "Taux", "Écart taux", "Diagnostic", "Qui", "Depuis"],
             [[short_rayon(a["rayon"]), short_site(a["site"]), a["secteur"], a["ca"], a["vs_n1"], a["marge"], a["dmarge"],
               a["taux"], a["dpts"], a["diagnostic"], a["qui"], "Aussi sur la semaine" if a.get("persistant") else "Hier"]
              for a in R["alertes"]],
             [None, None, None, MIO, PCT, MIO, DMIO, TAUX, PTS, None, None, None])
    sh.widths([26, 16, 14, 11, 11, 11, 11, 11, 11, 11, 30, 18, 18, 11, 11])

    # ------------------------------------------------------------ secteurs
    for sec in ("PGC", "Non-alimentaire", "Frais"):
        if sec not in R["secteurs"]:
            continue
        k = R["secteurs"][sec]
        sh = Sheet(wb.create_sheet(sec), f"Secteur {sec}", f"Ventes du {d} · HT · vs même jour N-1")
        head, fmt = _kpi_head("")
        sh.section("Hier et tendance de la semaine")
        sh.table(head, [_kpi_row("Hier", k)], fmt)
        sh.table(SEM_HEAD, [_sem_row("Semaine", k)], SEM_FMT)
        sh.section("Par format")
        head, fmt = _kpi_head("Format")
        sh.table(head + ["Hors seuil"],
                 [_kpi_row(f, v["kpi"]) + [v["texte"]] for f, v in R["croise"].get(sec, {}).items()], fmt + [None])
        sh.section("Rayons à traiter")
        head, fmt = _kpi_head("Rayon")
        sh.table(head + ["Piste"], [_kpi_row(r["nom"], r) + [r.get("action", "")] for r in R["rayons"].get(sec, [])], fmt + [None])
        al = [a for a in R["alertes"] if a["secteur"] == sec]
        sh.section("Magasins à appeler")
        sh.table(["Rayon", "Magasin", "CA vs N-1", "Écart marge", "Écart taux", "Diagnostic", "Qui"],
                 [[short_rayon(a["rayon"]), short_site(a["site"]), a["vs_n1"], a["dmarge"], a["dpts"], a["diagnostic"], a["qui"]] for a in al]
                 or [["Aucune alerte", "", None, None, None, "", ""]],
                 [None, None, PCT, DMIO, PTS, None, None])
        sh.widths([24, 16, 11, 11, 11, 11, 11, 11, 11, 11, 30, 18])

    # ------------------------------------------------------------ retours
    sh = Sheet(wb.create_sheet("Retours attendus"), "Retours attendus",
               f"Réponse le {_jour(R['date_reponse'])} {R['date_reponse'].strftime('%d/%m')} avant {C.HEURE_REPONSE}")
    sh.table(["Sujet", "Question", "Qui répond", "Type", "Réponse"],
             [[x["sujet"], x["question"], x["qui"], x["type"], ""] for x in R["retours"]])
    sh.widths([34, 60, 30, 18, 40])

    for ws in wb.worksheets:
        ws.freeze_panes = "B4"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _jour(d) -> str:
    return ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"][d.weekday()]
