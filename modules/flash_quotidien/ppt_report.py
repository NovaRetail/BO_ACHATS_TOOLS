"""PowerPoint du flash, construit sur le masque CFAO fourni (.pptx).

Le masque n'est pas codé en dur : on lit ses mises en page par nom (avec repli),
on garde sa couverture et on recopie son pied de page et son numéro de slide.
"""
from __future__ import annotations

import copy
import datetime as dt
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import config as C
from .engine import kfcfa, mfcfa, pct, pts, short_rayon, short_site

# ------------------------------------------------------------------ charte CFAO Consumer 2025
NAVY, ORANGE, ORANGE2 = "24235C", "F09F39", "F5AD3E"
GREY, LINE, INK, MUTE = "F3F4F8", "DADDE6", "363C5A", "6B7090"
RED, TEAL, WHITE = "E62244", "08868D", "FFFFFF"
JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]

X0, W = 0.61, 12.1   # marge gauche et largeur utile du masque
Y0 = 1.45            # sous le bandeau de titre


def rgb(h):
    return RGBColor.from_string(h)


def tone(v, bad_if_neg=True):
    if v is None:
        return MUTE
    return RED if (v < 0) == bad_if_neg else TEAL


# ------------------------------------------------------------------ masque
def _layout(prs, *names):
    for n in names:
        for lay in prs.slide_layouts:
            if lay.name.lower() == n.lower():
                return lay
    for n in names:
        for lay in prs.slide_layouts:
            if n.lower() in lay.name.lower():
                return lay
    return prs.slide_layouts[0]


def _keep_only_cover(prs):
    """Garde la 1re slide du masque si c'est une couverture, supprime les autres."""
    sld = prs.slides._sldIdLst
    ids = list(sld)
    keep = None
    if ids and "couverture" in prs.slides[0].slide_layout.name.lower():
        keep = ids[0]
    for s in ids:
        if s is keep:
            continue
        prs.part.drop_rel(s.get(qn("r:id")))
        sld.remove(s)
    return keep is not None


def _clone_footer(slide, layout, text):
    for ph in layout.placeholders:
        t = str(ph.placeholder_format.type)
        if "FOOTER" in t or "SLIDE_NUMBER" in t:
            el = copy.deepcopy(ph._element)
            slide.shapes._spTree.append(el)
            if "FOOTER" in t:
                tf = slide.shapes[-1].text_frame
                p = tf.paragraphs[0]
                for r in list(p.runs)[1:]:
                    r._r.getparent().remove(r._r)
                if p.runs:
                    p.runs[0].text = text
                else:
                    p.add_run().text = text


def _title(slide, light, bold):
    ph = slide.shapes.title
    tf = ph.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = light; r1.font.bold = False
    r2 = p.add_run(); r2.text = bold


# ------------------------------------------------------------------ primitives
def tb(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, italic=False):
    """Zone de texte. runs : texte, ou liste de paragraphes = listes de (texte, options)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.04))
    paras = [[(runs, {})]] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, o in para:
            r = p.add_run(); r.text = text
            f = r.font
            f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
        if i and "space" in (para[0][1] if para else {}):
            p.space_before = Pt(para[0][1]["space"])
    return box


def rect(slide, x, y, w, h, fill, rounded=True, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.08
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def circle(slide, x, y, d, fill, text, size=14, color=WHITE):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = rgb(fill); c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = rgb(color)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def table(slide, x, y, w, col_w, rows, header, size=13, row_h=0.42, styles=None):
    """Tableau au style du masque : en-tête orange, lignes alternées gris clair."""
    n, m = len(rows) + 1, len(header)
    gt = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(row_h * n))
    t = gt.table
    tblPr = t._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # « aucun style, grille du tableau »
    tot = sum(col_w)
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(w * cw / tot)
    for i in range(n):
        t.rows[i].height = Inches(row_h)
        for j in range(m):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = rgb(ORANGE)
                val, o = header[j], {"color": WHITE, "bold": True, "italic": True}
            else:
                cell.fill.fore_color.rgb = rgb(GREY if i % 2 == 0 else WHITE)
                val = rows[i - 1][j]
                o = (styles[i - 1][j] if styles and styles[i - 1][j] else {})
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 or o.get("left") else PP_ALIGN.RIGHT if o.get("num") else PP_ALIGN.LEFT
            r = p.add_run(); r.text = str(val)
            r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", False); r.font.italic = o.get("italic", False)
            r.font.color.rgb = rgb(o.get("color", INK))
    _no_borders(t)
    return gt


def _no_borders(t):
    for row in t.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tcPr.find(qn(tag))
                if ln is None:
                    ln = tcPr.makeelement(qn(tag), {})
                    tcPr.insert(0, ln)
                ln.set("w", "0")
                for ch in list(ln):
                    ln.remove(ch)
                ln.append(ln.makeelement(qn("a:noFill"), {}))


def bar_chart(slide, x, y, w, h, cats, vals, title, fmt='+0.0" M";-0.0" M"', labels=None):
    cd = CategoryChartData(); cd.categories = cats; cd.add_series(title, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size = Pt(13); tp.runs[0].font.bold = True; tp.runs[0].font.color.rgb = rgb(NAVY)
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(12); dl.font.bold = True; dl.font.color.rgb = rgb(INK)
    ser = plot.series[0]
    for i, v in enumerate(vals):
        pt_ = ser.points[i]
        pt_.format.fill.solid(); pt_.format.fill.fore_color.rgb = rgb(TEAL if v >= 0 else RED)
        if labels:  # libellé texte au format français (virgule), quel que soit le poste
            tf = pt_.data_label.text_frame
            tf.text = labels[i]
            r = tf.paragraphs[0].runs[0]
            r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = rgb(INK)
            pt_.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(12); ca.tick_labels.font.color.rgb = rgb(INK)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.color.rgb = rgb(LINE)
    return gf


# ------------------------------------------------------------------ slides
def _content(prs, title_light, title_bold, footer):
    lay = _layout(prs, "Titre seul", "Title Only", "Contenu")
    s = prs.slides.add_slide(lay)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    _title(s, title_light, title_bold)
    _clone_footer(s, lay, footer)
    return s


def _cover(prs, kept, R):
    d = R["date"]
    sub = f"Ventes du {JOURS[d.weekday()]} {d:%d/%m/%Y} · Category management, pricing et données"
    if kept:
        s = prs.slides[0]
    else:
        s = prs.slides.add_slide(_layout(prs, "14_Couverture", "Couverture"))
    if s.shapes.title is not None:
        _title(s, "Flash ", "commercial")
    sub_box = None
    for sh in s.shapes:
        if sh.has_text_frame and not sh.is_placeholder and "sous-titre" in sh.text_frame.text.lower():
            sub_box = sh
    if sub_box is not None:
        p = sub_box.text_frame.paragraphs[0]
        for r in list(p.runs)[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = sub
        sub_box.width = Inches(9.0)
        sub_box.text_frame.word_wrap = False
        p.alignment = PP_ALIGN.LEFT
    else:
        tb(s, 0.94, 4.86, 9.0, 0.4, sub, size=14, color=WHITE)


def _essentiel(prs, R, footer):
    r, w = R["reseau"], R["reseau_sem"]
    s = _content(prs, "L'essentiel : ", f"ventes {pct(r['vs_n1'], 0)}, marge {pct(r['marge_vs_n1'], 0)}", footer)
    cards = [
        ("Chiffre d'affaires", mfcfa(r["ca"]), pct(r["vs_n1"]), r["vs_n1"]),
        ("Marge", mfcfa(r["marge"]), pct(r["marge_vs_n1"]), r["marge_vs_n1"]),
        ("Taux de marge", pct(r["taux"], sign=False), pts(r["dpts"]), r["dpts"]),
        ("Fréquentation", pct(r.get("debit_vs_n1")), f"panier {pct(r.get('panier_vs_n1'))}", r.get("panier_vs_n1")),
    ]
    cw, gap = 2.86, 0.22
    for i, (lab, val, dlt, v) in enumerate(cards):
        x = X0 + i * (cw + gap)
        rect(s, x, Y0 + 0.1, cw, 1.6, GREY)
        tb(s, x + 0.2, Y0 + 0.2, cw - 0.4, 0.35, lab, size=13, color=MUTE)
        tb(s, x + 0.2, Y0 + 0.55, cw - 0.4, 0.6, val, size=30, bold=True, color=NAVY)
        tb(s, x + 0.2, Y0 + 1.15, cw - 0.4, 0.4, dlt, size=14, bold=True, color=tone(v))
    tb(s, X0, 3.45, 6.4, 0.4, "À retenir", size=18, bold=True, color=NAVY)
    for i, msg in enumerate(R["messages"]):
        y = 4.0 + i * 0.95
        circle(s, X0, y, 0.42, ORANGE, str(i + 1), size=14)
        tb(s, X0 + 0.6, y - 0.04, 5.9, 0.85, msg, size=15, color=INK)
    cats = list(R["secteurs"].keys())
    vals = [round(R["secteurs"][c]["dmarge"] / 1e6, 2) for c in cats]
    labels = [mfcfa(R["secteurs"][c]["dmarge"], sign=True) for c in cats]
    bar_chart(s, 7.35, 3.4, 5.35, 3.6, cats, vals, "Marge vs N-1 par secteur (M FCFA)", labels=labels)
    s.notes_slide.notes_text_frame.text = R["lecture"]


def _secteurs(prs, R, footer):
    sec = R["secteurs"]
    perdants = [n for n, k in sec.items() if k["dmarge"] < 0]
    if perdants:
        t = f"le {perdants[0].lower()} perd de la marge" if len(perdants) == 1 else f"{len(perdants)} secteurs perdent de la marge"
    else:
        best = max(sec, key=lambda n: sec[n]["dmarge"])
        t = f"{best} tire la marge"
    s = _content(prs, "Les secteurs : ", t, footer)
    cw, gap = 3.9, 0.2
    for i, n in enumerate(C.SECTEURS):
        k = sec[n]
        x = X0 + i * (cw + gap)
        rect(s, x, Y0 + 0.05, cw, 5.35, GREY)
        hdr = rect(s, x, Y0 + 0.05, cw, 0.6, NAVY)
        tb(s, x + 0.2, Y0 + 0.1, cw - 0.4, 0.5, n.upper(), size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + 0.25, Y0 + 0.85, cw - 0.5, 0.35, "Marge vs N-1", size=13, color=MUTE)
        tb(s, x + 0.25, Y0 + 1.15, cw - 0.5, 0.65, pct(k["marge_vs_n1"]), size=32, bold=True, color=tone(k["marge_vs_n1"]))
        tb(s, x + 0.25, Y0 + 1.8, cw - 0.5, 0.35, f"soit {mfcfa(k['dmarge'], sign=True)} FCFA", size=13, color=MUTE)
        lines = [
            ("Chiffre d'affaires", pct(k["vs_n1"]), tone(k["vs_n1"])),
            ("Taux de marge", f"{pct(k['taux'], sign=False)}  ({pts(k['dpts'])})", tone(k["dpts"])),
            ("Semaine vs budget", pct(k["sem"].get("vs_budget")), tone(k["sem"].get("vs_budget"))),
        ]
        for j, (lab, val, col) in enumerate(lines):
            y = Y0 + 2.35 + j * 0.5
            tb(s, x + 0.25, y, 1.9, 0.4, lab, size=13, color=INK)
            tb(s, x + 2.0, y, cw - 2.25, 0.4, val, size=13, bold=True, color=col, align=PP_ALIGN.RIGHT)
        ray = R["rayons"][n][:2]
        tb(s, x + 0.25, Y0 + 3.95, cw - 0.5, 0.35, "À regarder", size=13, bold=True, color=NAVY)
        if ray:
            txt = [[(f"{r['nom']} : marge {kfcfa(r['dmarge'])}, taux {pts(r['dpts'])}", {})] for r in ray]
        else:
            txt = [[("Aucun rayon hors seuil", {"color": MUTE})]]
        tb(s, x + 0.25, Y0 + 4.3, cw - 0.5, 1.0, txt, size=13, color=INK)


def _ou_agir(prs, R, footer):
    al = R["sites_alerte"]
    s = _content(prs, "Où agir : ", f"{len(al)} magasins en alerte" if al else "aucun magasin en alerte", footer)
    tb(s, X0, Y0, 7.4, 0.4, "Magasins en alerte", size=16, bold=True, color=NAVY)
    rows, st = [], []
    for r in al:
        rows.append([r["nom"], r["format"], pct(r["vs_n1"]), pct(r["marge_vs_n1"]), pct(r["taux"], sign=False), pts(r["dpts"])])
        st.append([None, {"color": MUTE}, {"num": True, "color": tone(r["vs_n1"])}, {"num": True, "color": tone(r["marge_vs_n1"]), "bold": True},
                   {"num": True}, {"num": True, "color": tone(r["dpts"])}])
    if not rows:
        rows, st = [["Aucun magasin hors seuil", "", "", "", "", ""]], None
    table(s, X0, Y0 + 0.45, 7.4, [2.2, 1.8, 1.1, 1.1, 1.1, 1.1], rows,
          ["Magasin", "Format", "CA", "Marge", "Taux", "Écart taux"], styles=st)
    # écarts secteur x format
    y = Y0 + 0.55 + 0.42 * (len(rows) + 1) + 0.35
    tb(s, X0, y, 7.4, 0.4, "Écarts hors seuil par secteur et format", size=16, bold=True, color=NAVY)
    para = []
    for sec, d in R["croise"].items():
        items = [f"{f} {v['texte']}" for f, v in d.items() if v["texte"]]
        if items:
            para.append([(f"{sec} : ", {"bold": True, "color": NAVY}), (" ; ".join(items), {})])
    tb(s, X0, y + 0.45, 7.4, 1.6, para or "Aucun écart hors seuil", size=13, color=INK)
    # magasins convertis
    xr = 8.35
    rect(s, xr, Y0, 4.36, 5.4, GREY)
    tb(s, xr + 0.25, Y0 + 0.15, 3.9, 0.4, "Magasins convertis", size=16, bold=True, color=NAVY)
    tb(s, xr + 0.25, Y0 + 0.55, 3.9, 0.55, "Comparés à leur ancien format Supeco : on lit surtout la tendance.", size=12, color=MUTE)
    for i, r in enumerate(R["convertis"]):
        y = Y0 + 1.25 + i * 1.3
        tb(s, xr + 0.25, y, 3.9, 0.35, f"{r['nom']}  ·  converti le {r['converti_le']:%d/%m}", size=14, bold=True, color=INK)
        lines = [[("Marge ", {"color": MUTE}), (pct(r["marge_vs_n1"]), {"bold": True, "color": tone(r["marge_vs_n1"])}),
                  ("   Débit ", {"color": MUTE}), (pct(r.get("debit_vs_n1")), {"bold": True, "color": tone(r.get("debit_vs_n1"))})],
                 [("Panier ", {"color": MUTE}), (pct(r.get("panier_vs_n1")), {"bold": True, "color": tone(r.get("panier_vs_n1"))}),
                  ("   Budget sem. ", {"color": MUTE}),
                  (pct(r["sem"].get("vs_budget")) if r["budgete"] else "n.d.", {"bold": True, "color": tone(r["sem"].get("vs_budget")) if r["budgete"] else MUTE})]]
        tb(s, xr + 0.25, y + 0.38, 3.9, 0.75, lines, size=13, color=INK)


def _retours(prs, R, footer):
    rep = R["date_reponse"]
    att = [x for x in R["retours"] if x["type"] == "Réponse attendue"]
    s = _content(prs, "Retours attendus : ",
                 f"{len(att)} questions, réponse le {JOURS[rep.weekday()]} {rep:%d/%m} avant {C.HEURE_REPONSE}", footer)
    rows, st = [], []
    for x in R["retours"]:
        info = x["type"] != "Réponse attendue"
        rows.append([x["sujet"], x["question"], x["qui"], "Pour information" if info else "Réponse attendue"])
        o = {"italic": True, "color": MUTE} if info else {}
        st.append([dict(o, bold=not info), o or None, dict(o, bold=not info, color=MUTE if info else NAVY), dict(o, size=12)])
    table(s, X0, Y0 + 0.1, W, [2.9, 5.2, 2.6, 1.6], rows, ["Sujet", "Question", "Qui répond", "Type"],
          size=13, row_h=0.62, styles=st)
    y = Y0 + 0.2 + 0.62 * (len(rows) + 1) + 0.2
    tb(s, X0, min(y, 6.45), W, 0.45,
       "Les réponses ferment le sujet ou le reportent à l'édition suivante. Les sujets « pour information » "
       "relèvent d'autres directions.", size=12, italic=True, color=MUTE)


def _lecture(prs, R, footer):
    s = _content(prs, "Comment lire ", "ce flash", footer)
    se = C.SEUILS
    blocs = [
        ("Marge et taux de marge", "La marge est en FCFA, le taux en % du CA. Le taux peut baisser pendant que la marge "
                                   "progresse si les ventes augmentent."),
        ("Effet volume, effet taux", "L'évolution de la marge se découpe en deux : ce que rapportent les ventes en plus "
                                     "(volume) et ce que coûte la baisse du taux (prix, promo, mix)."),
        ("Débit et panier", "Le débit compte les tickets, c'est la fréquentation. Le panier est le CA par ticket. "
                            "Débit en baisse : trafic. Panier en baisse : offre ou ruptures."),
        ("Seuils d'alerte", f"CA sous {pct(se['ca_vs_n1'], 0)}, marge sous {pct(se['marge_vs_n1'], 0)} ou taux en "
                            f"baisse de plus de {abs(se['taux_pts']):.0f} pts vs N-1. Seuls ces écarts sont affichés."),
    ]
    bw, bh = 5.95, 1.55
    for i, (t, d) in enumerate(blocs):
        x = X0 + (i % 2) * (bw + 0.2)
        y = Y0 + 0.05 + (i // 2) * (bh + 0.2)
        rect(s, x, y, bw, bh, GREY)
        circle(s, x + 0.2, y + 0.22, 0.42, NAVY, str(i + 1), size=14)
        tb(s, x + 0.8, y + 0.18, bw - 1.0, 0.4, t, size=15, bold=True, color=NAVY)
        tb(s, x + 0.8, y + 0.58, bw - 1.0, 0.95, d, size=13, color=INK)
    y = Y0 + 3.75
    tb(s, X0, y, W, 0.4, "Le rituel", size=16, bold=True, color=NAVY)
    etapes = [("Chaque matin", "le flash de la veille, comparé au même jour N-1"),
              (f"Le lendemain avant {C.HEURE_REPONSE}", "les réponses aux retours attendus"),
              ("Chaque lundi", "la revue de la semaine et les causes des écarts qui durent")]
    ew = 3.85
    for i, (t, d) in enumerate(etapes):
        x = X0 + i * (ew + 0.275)
        rect(s, x, y + 0.5, ew, 1.15, WHITE, line=ORANGE)
        tb(s, x + 0.2, y + 0.6, ew - 0.4, 0.4, t, size=14, bold=True, color=ORANGE)
        tb(s, x + 0.2, y + 0.98, ew - 0.4, 0.6, d, size=13, color=INK)
        if i < 2:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + ew + 0.04), Inches(y + 0.95), Inches(0.2), Inches(0.25))
            a.fill.solid(); a.fill.fore_color.rgb = rgb(ORANGE); a.line.fill.background()


def build_ppt(R: dict, masque, with_lecture: bool = True) -> bytes:
    """R : résultat de engine.compute. masque : chemin ou flux du .pptx CFAO."""
    if isinstance(masque, (bytes, bytearray)):
        masque = io.BytesIO(masque)
    prs = Presentation(masque)
    kept = _keep_only_cover(prs)
    d = R["date"]
    footer = f"Flash commercial · ventes du {d:%d/%m/%Y}"
    _cover(prs, kept, R)
    _essentiel(prs, R, footer)
    _secteurs(prs, R, footer)
    _ou_agir(prs, R, footer)
    _retours(prs, R, footer)
    if with_lecture:
        _lecture(prs, R, footer)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
