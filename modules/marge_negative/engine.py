"""Moteur du module Marge négative (ex « Diagnostic Rentabilité Réseau ») : lecture de l'export PBI,
calculs et export Excel, repris tels quels de la page v2.2."""
import streamlit as st
import pandas as pd
import numpy as np
import re
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── HELPERS ──────────────────────────────────────────────────────────────────
def fmt(n):
    if pd.isna(n) or n is None: return "—"
    a = abs(n)
    if a >= 1_000_000: return f"{n/1_000_000:.1f} M"
    if a >= 1_000:     return f"{int(n/1_000)} K"
    return f"{int(n):,}"

def fmt_pct(v, dec=1):
    if pd.isna(v) or v is None: return "—"
    return f"{v:.{dec}f}%"

def fmt_delta(v):
    if pd.isna(v) or v is None: return "—"
    return f"{v:+.1f} pts"

def get_format(site_name):
    s = str(site_name)
    if "Supeco" in s: return "Supeco"
    if "Hyper"  in s: return "Hyper"
    return "Market"

def short_name(s):
    s = str(s)
    return s.split(" - ", 1)[-1].strip() if " - " in s else s

def extract_code(s):
    if pd.isna(s): return None
    return str(s).split(" - ", 1)[0].strip() if " - " in str(s) else None

def extract_periode(df_raw):
    try:
        for val in df_raw.iloc[:, 0].astype(str):
            m = re.search(r"après le (\d{2}/\d{2}/\d{4}) et est avant le (\d{2}/\d{2}/\d{4})", val)
            if m:
                from datetime import datetime
                d1 = datetime.strptime(m.group(1), "%d/%m/%Y")
                d2 = datetime.strptime(m.group(2), "%d/%m/%Y")
                nb = (d2 - d1).days
                return f"{m.group(1)} → {m.group(2)}", nb
    except: pass
    return "Période inconnue", 1

# ─── CHARGEMENT ───────────────────────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def load_data(byt, fname):
    ext = fname.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xls"):
        df = pd.read_excel(BytesIO(byt), dtype=str)
    else:
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                df = pd.read_csv(BytesIO(byt), sep=";", encoding=enc, dtype=str)
                break
            except: continue

    periode, nb_jours = extract_periode(df)

    if "CA Promo" in df.columns and "CA HT Promo" not in df.columns:
        df = df.rename(columns={"CA Promo": "CA HT Promo"})

    num_cols = ["CA", "Marge", "CA Hors Promo", "Marge Hors Promo",
                "CA HT Promo", "Marge Promo", "Qté Vente",
                "Casse (Valeur)", "Casse (Qté)", "%Marge", "%CA Poids Promo"]
    for col in num_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df["lib_art"]   = df["Article"].apply(lambda s: short_name(s) if pd.notna(s) else None)
    df["code_art"]  = df["Article"].apply(extract_code)
    df["lib_site"]  = df["Site nom long"].apply(lambda s: short_name(s) if pd.notna(s) else None)
    df["lib_rayon"] = df["Rayon"].apply(lambda s: short_name(s) if pd.notna(s) else None)
    df["lib_fam"]   = df["Famille"].apply(lambda s: short_name(s) if pd.notna(s) else None)
    df["format"]    = df["Site nom long"].apply(lambda s: get_format(s) if pd.notna(s) else None)

    df_clean = df[
        df["lib_art"].notna()   & (df["lib_art"]   != "Total") &
        df["lib_site"].notna()  & (df["lib_site"]  != "Total") &
        df["lib_rayon"].notna() &
        ~df["Rayon"].astype(str).str.startswith("Filtres") &
        (df["lib_rayon"] != "Total") &
        df["lib_fam"].notna()   & (df["lib_fam"]   != "Total")
    ].copy()

    return df_clean, periode, nb_jours


def compute(df_f):
    """Toutes les agrégations du module, sur le périmètre filtré."""
    nb_sites_actifs = df_f[df_f["CA"] > 0]["lib_site"].nunique()

    def agg_base(df_in, by):
        g = df_in.groupby(by).agg(
            CA=("CA","sum"), Marge=("Marge","sum"),
            CA_Promo=("CA HT Promo","sum"), CA_HP=("CA Hors Promo","sum"),
            Marge_HP=("Marge Hors Promo","sum"), Marge_Promo=("Marge Promo","sum"),
            Casse=("Casse (Valeur)","sum")
        ).reset_index()
        g["TxMarge"]       = (g["Marge"] / g["CA"] * 100).where(g["CA"] > 0)
        g["PdsPromo"]      = (g["CA_Promo"] / g["CA"] * 100).where(g["CA"] > 0)
        g["TxMarge_HP"]    = (g["Marge_HP"] / g["CA_HP"] * 100).where(g["CA_HP"] > 0)
        g["TxMarge_Promo"] = (g["Marge_Promo"] / g["CA_Promo"] * 100).where(g["CA_Promo"] > 0)
        g["TxCasse"]       = (g["Casse"].abs() / g["CA"] * 100).where(g["CA"] > 0)
        return g

    agg_fmt  = agg_base(df_f, "format").sort_values("TxMarge", ascending=False)
    agg_rax  = agg_base(df_f, "lib_rayon")
    agg_rax["PoidsCA"] = (agg_rax["CA"] / agg_rax["CA"].sum() * 100)
    agg_rax  = agg_rax.sort_values("TxMarge")

    agg_site = df_f.groupby(["lib_site","format"]).agg(
        CA=("CA","sum"), Marge=("Marge","sum"),
        CA_Promo=("CA HT Promo","sum"), Casse=("Casse (Valeur)","sum")
    ).reset_index()
    agg_site["TxMarge"]  = (agg_site["Marge"] / agg_site["CA"] * 100).where(agg_site["CA"] > 0)
    agg_site["PdsPromo"] = (agg_site["CA_Promo"] / agg_site["CA"] * 100).where(agg_site["CA"] > 0)
    agg_site["TxCasse"]  = (agg_site["Casse"].abs() / agg_site["CA"] * 100).where(agg_site["CA"] > 0)
    agg_site = agg_site[agg_site["CA"] > 0].sort_values("TxMarge", ascending=False).reset_index(drop=True)
    moy_marge_site = agg_site["TxMarge"].mean()

    mat = df_f.groupby(["lib_site","lib_rayon"]).agg(CA=("CA","sum"), Marge=("Marge","sum")).reset_index()
    mat["TxMarge"] = (mat["Marge"] / mat["CA"] * 100).where(mat["CA"] > 0)
    mat_pivot = mat.pivot_table(index="lib_rayon", columns="lib_site", values="TxMarge").round(1)

    art_site = df_f[df_f["CA"] > 0].groupby(["Article","lib_site","format"]).agg(
        CA=("CA","sum"), Marge=("Marge","sum"), Qte=("Qté Vente","sum")
    ).reset_index()
    art_site["TxMarge_site"] = (art_site["Marge"] / art_site["CA"] * 100).where(art_site["CA"] > 0)
    art_site["lib_court"]    = art_site["lib_site"]

    # ── agg_art avec code_art ─────────────────────────────────────────────────────
    agg_art = df_f.groupby(["Article", "lib_art", "code_art", "lib_rayon", "lib_fam"]).agg(
        CA=("CA","sum"), Marge=("Marge","sum"),
        CA_Promo=("CA HT Promo","sum"), Qte=("Qté Vente","sum")
    ).reset_index()
    agg_art["TxMarge"]  = (agg_art["Marge"] / agg_art["CA"] * 100).where(agg_art["CA"] > 0)
    agg_art["PdsPromo"] = (agg_art["CA_Promo"] / agg_art["CA"] * 100).where(agg_art["CA"] > 0)

    flop100 = agg_art[agg_art["CA"] > 5000].nsmallest(100, "TxMarge").copy().reset_index(drop=True)
    flop100["Rang"] = range(1, len(flop100) + 1)

    def build_bloc(article_full, fmt_name):
        rows = art_site[(art_site["Article"] == article_full) & (art_site["format"] == fmt_name)].sort_values("TxMarge_site")
        if rows.empty: return "—"
        parts = []
        for _, r in rows.iterrows():
            tm = r["TxMarge_site"]; qty = r["Qte"]
            if pd.notna(tm):
                parts.append(f"{r['lib_court']}: {tm:.1f}% | Qty: {int(qty):,}" if pd.notna(qty) else f"{r['lib_court']}: {tm:.1f}%")
        return "  |  ".join(parts) if parts else "—"

    flop100["Bloc_Hyper"]  = flop100["Article"].apply(lambda a: build_bloc(a, "Hyper"))
    flop100["Bloc_Market"] = flop100["Article"].apply(lambda a: build_bloc(a, "Market"))
    flop100["Bloc_Supeco"] = flop100["Article"].apply(lambda a: build_bloc(a, "Supeco"))

    # KPIs
    ca_total    = df_f["CA"].sum()
    marge_total = df_f["Marge"].sum()
    ca_promo    = df_f["CA HT Promo"].sum()
    ca_hp       = df_f["CA Hors Promo"].sum()
    m_promo     = df_f["Marge Promo"].sum()
    m_hp        = df_f["Marge Hors Promo"].sum()
    casse_total = df_f["Casse (Valeur)"].sum()
    tx_marge    = marge_total / ca_total * 100  if ca_total > 0 else 0
    tx_m_promo  = m_promo / ca_promo * 100      if ca_promo > 0 else 0
    tx_m_hp     = m_hp / ca_hp * 100            if ca_hp > 0    else 0
    poids_promo = ca_promo / ca_total * 100     if ca_total > 0 else 0
    delta_hp_p  = tx_m_hp - tx_m_promo
    tx_casse    = abs(casse_total) / ca_total * 100 if ca_total > 0 else 0
    nb_art_neg  = int((agg_art["TxMarge"] < 0).sum())
    nb_flop_neg = int((flop100["TxMarge"] < 0).sum())

    # ── Données marges négatives par site (utilisé dans Tab Streamlit + export) ───
    art_neg_site = art_site[art_site["TxMarge_site"] < 0].copy()
    art_neg_site = art_neg_site.merge(
        df_f[["Article","lib_art","code_art","lib_rayon","lib_fam"]].drop_duplicates("Article"),
        on="Article", how="left"
    )
    # Résoudre les colonnes dupliquées issues du merge
    for col in ["lib_art","code_art","lib_rayon","lib_fam"]:
        col_x = col + "_x"
        col_y = col + "_y"
        if col_x in art_neg_site.columns:
            art_neg_site[col] = art_neg_site[col_x].combine_first(art_neg_site[col_y])
            art_neg_site.drop(columns=[col_x, col_y], inplace=True)
    art_neg_site = art_neg_site.sort_values(["lib_rayon","lib_art","TxMarge_site"]).reset_index(drop=True)
    agg_fam = df_f.groupby("lib_fam").agg(CA=("CA","sum"), Marge=("Marge","sum")).reset_index()
    agg_fam["TxMarge"] = (agg_fam["Marge"] / agg_fam["CA"] * 100).where(agg_fam["CA"] > 0)
    fam_sous_seuil = agg_fam[(agg_fam["TxMarge"] < 8) & (agg_fam["CA"] > 500_000)]
    casse_col_v = "Casse (Valeur)"
    casse_col_q = "Casse (Qté)"
    has_qty = casse_col_q in df_f.columns

    agg_casse_site = df_f.groupby(["lib_site","format"]).agg(
        CA=("CA","sum"), Casse_V=(casse_col_v,"sum"),
        **({} if not has_qty else {"Casse_Q": (casse_col_q,"sum")})
    ).reset_index()
    agg_casse_site["TxCasse"] = (agg_casse_site["Casse_V"].abs() / agg_casse_site["CA"] * 100).where(agg_casse_site["CA"] > 0)
    agg_casse_site = agg_casse_site[agg_casse_site["CA"] > 0].sort_values("TxCasse", ascending=False).reset_index(drop=True)

    agg_casse_rax = df_f.groupby("lib_rayon").agg(CA=("CA","sum"), Casse_V=(casse_col_v,"sum")).reset_index()
    agg_casse_rax["TxCasse"] = (agg_casse_rax["Casse_V"].abs() / agg_casse_rax["CA"] * 100).where(agg_casse_rax["CA"] > 0)
    agg_casse_rax = agg_casse_rax.sort_values("TxCasse", ascending=False)

    agg_casse_art = df_f.groupby(["Article","lib_art","code_art","lib_rayon","lib_fam"]).agg(
        CA=("CA","sum"), Casse_V=(casse_col_v,"sum"),
        **({} if not has_qty else {"Casse_Q": (casse_col_q,"sum")})
    ).reset_index()
    agg_casse_art["TxCasse"] = (agg_casse_art["Casse_V"].abs() / agg_casse_art["CA"] * 100).where(agg_casse_art["CA"] > 0)
    top30_casse = agg_casse_art[agg_casse_art["Casse_V"].abs() > 0].nlargest(30, "Casse_V").reset_index(drop=True)
    top30_casse["Rang"] = range(1, len(top30_casse)+1)

    casse_reseau      = df_f[casse_col_v].sum()
    tx_casse_reseau   = abs(casse_reseau) / ca_total * 100 if ca_total > 0 else 0
    nb_sites_alerte   = int((agg_casse_site["TxCasse"] > 1).sum())
    moy_tx_casse_site = agg_casse_site["TxCasse"].mean()
    return {k: v for k, v in locals().items() if k not in ('df_in', 'by')}


def build_excel(R, periode, nb_jours):
    """Classeur 6 onglets (valeurs uniquement)."""
    agg_fmt = R['agg_fmt']
    agg_rax = R['agg_rax']
    agg_site = R['agg_site']
    mat_pivot = R['mat_pivot']
    flop100 = R['flop100']
    art_neg_site = R['art_neg_site']
    ca_total = R['ca_total']
    marge_total = R['marge_total']
    tx_marge = R['tx_marge']
    tx_m_promo = R['tx_m_promo']
    tx_m_hp = R['tx_m_hp']
    poids_promo = R['poids_promo']
    delta_hp_p = R['delta_hp_p']
    tx_casse = R['tx_casse']
    top30_casse = R['top30_casse']
    agg_casse_site = R['agg_casse_site']
    agg_casse_rax = R['agg_casse_rax']
    has_qty = R['has_qty']
    casse_reseau = R['casse_reseau']
    tx_casse_reseau = R['tx_casse_reseau']
    moy_tx_casse_site = R['moy_tx_casse_site']
    nb_sites_alerte = R['nb_sites_alerte']
    wb_exp = Workbook()
    C_HDR="1B2A4A"; C_SUB="2E4B7A"; C_WH="FFFFFF"; C_DK="1A1A2E"
    C_HYP="154360"; C_MKT="145A32"; C_SUP="6E2F8A"

    def xfill(h): return PatternFill("solid", fgColor=h)
    def xbdr():
        s = Side(style="thin", color="CCCCCC")
        return Border(left=s, right=s, top=s, bottom=s)
    def xctr(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
    def xrgt(): return Alignment(horizontal="right",  vertical="center")
    def xlft(w=False): return Alignment(horizontal="left", vertical="center", wrap_text=w)

    def write_header_row(ws, row_num, headers, widths, bg=C_SUB):
        for i,(h,w) in enumerate(zip(headers,widths)):
            c = ws.cell(row=row_num, column=i+1, value=h)
            c.font=Font("Calibri",size=10,bold=True,color=C_WH)
            c.fill=xfill(bg); c.alignment=xctr(); c.border=xbdr()
            ws.column_dimensions[get_column_letter(i+1)].width = w
        ws.row_dimensions[row_num].height = 24

    def title_block(ws, txt, span=10):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
        c = ws.cell(row=1,column=1,value=txt)
        c.font=Font("Calibri",size=13,bold=True,color=C_WH); c.fill=xfill(C_HDR); c.alignment=xctr()
        ws.row_dimensions[1].height=30
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=span)
        c2=ws.cell(row=2,column=1,value=f"  Période : {periode}  ·  {nb_jours} jour(s)")
        c2.font=Font("Calibri",size=9,italic=True,color="AABBCC"); c2.fill=xfill(C_HDR); c2.alignment=xlft()
        ws.row_dimensions[2].height=16; ws.row_dimensions[3].height=6

    # ── Onglet 1 : Synthèse
    ws1 = wb_exp.active; ws1.title = "Synthèse Réseau"
    title_block(ws1, "DIAGNOSTIC RENTABILITÉ RÉSEAU — SYNTHÈSE", span=8)
    r=4
    kpi_data = [
        ("CA Réseau (FCFA)",f"{ca_total:,.0f}"),("Marge Brute (FCFA)",f"{marge_total:,.0f}"),
        ("Taux de Marge",f"{tx_marge:.1f}%"),("Taux Marge HP",f"{tx_m_hp:.1f}%"),
        ("Taux Marge Promo",f"{tx_m_promo:.1f}%"),("Effet Promo (pts)",f"−{delta_hp_p:.1f}"),
        ("Poids Promo",f"{poids_promo:.1f}%"),("Taux Casse",f"{tx_casse:.2f}%"),
    ]
    for ci,(lbl,val) in enumerate(kpi_data):
        c=ws1.cell(row=r,column=ci+1,value=lbl); c.font=Font("Calibri",size=9,bold=True,color=C_WH)
        c.fill=xfill(C_SUB); c.alignment=xctr(); c.border=xbdr()
        c2=ws1.cell(row=r+1,column=ci+1,value=val); c2.font=Font("Calibri",size=12,bold=True,color=C_DK)
        c2.fill=xfill("FFFFFF"); c2.alignment=xctr(); c2.border=xbdr()
        ws1.column_dimensions[get_column_letter(ci+1)].width=18
    ws1.row_dimensions[r].height=20; ws1.row_dimensions[r+1].height=28; r+=3

    ws1.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
    c=ws1.cell(row=r,column=1,value="  PERFORMANCE PAR FORMAT")
    c.font=Font("Calibri",size=10,bold=True,color=C_WH); c.fill=xfill(C_SUB); c.alignment=xlft(); ws1.row_dimensions[r].height=22; r+=1
    write_header_row(ws1,r,["Format","CA (FCFA)","Marge (FCFA)","Tx Marge","Tx Marge HP","Tx Marge Promo","Pds Promo","Tx Casse"],[14,18,18,12,14,16,12,12]); r+=1
    for _,(fd) in agg_fmt.iterrows():
        bg_f={"Hyper":"D6EAF8","Market":"D5F5E3","Supeco":"E8DAEF"}.get(fd["format"],"FFFFFF")
        for ci,(v,fmt3) in enumerate([(fd["format"],None),(fd["CA"],"#,##0"),(fd["Marge"],"#,##0"),
            (fd["TxMarge"]/100 if pd.notna(fd["TxMarge"]) else None,"0.0%"),
            (fd.get("TxMarge_HP",None)/100 if pd.notna(fd.get("TxMarge_HP")) else None,"0.0%"),
            (fd.get("TxMarge_Promo",None)/100 if pd.notna(fd.get("TxMarge_Promo")) else None,"0.0%"),
            (fd["PdsPromo"]/100 if pd.notna(fd["PdsPromo"]) else None,"0.0%"),
            (fd["TxCasse"]/100 if pd.notna(fd["TxCasse"]) else None,"0.0%")]):
            c=ws1.cell(row=r,column=ci+1,value=v); c.font=Font("Calibri",size=10,color=C_DK)
            c.fill=xfill(bg_f); c.border=xbdr()
            if fmt3: c.number_format=fmt3
            c.alignment=xrgt() if ci in [1,2] else xctr()
        ws1.row_dimensions[r].height=20; r+=1

    r+=1
    ws1.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
    c=ws1.cell(row=r,column=1,value="  PALMARÈS MAGASINS")
    c.font=Font("Calibri",size=10,bold=True,color=C_WH); c.fill=xfill(C_SUB); c.alignment=xlft(); ws1.row_dimensions[r].height=22; r+=1
    write_header_row(ws1,r,["Rang","Magasin","Format","CA (FCFA)","Marge (FCFA)","Tx Marge","Pds Promo","Tx Casse"],[5,24,10,18,18,12,12,12]); r+=1
    for ri3,(_,sd) in enumerate(agg_site.iterrows()):
        bg_s="F7F7F7" if ri3%2==0 else "FFFFFF"
        for ci3,(v,f3) in enumerate([(ri3+1,None),(sd["lib_site"],None),(sd["format"],None),
            (sd["CA"],"#,##0"),(sd["Marge"],"#,##0"),
            (sd["TxMarge"]/100 if pd.notna(sd["TxMarge"]) else None,"0.0%"),
            (sd["PdsPromo"]/100 if pd.notna(sd["PdsPromo"]) else None,"0.0%"),
            (sd["TxCasse"]/100 if pd.notna(sd["TxCasse"]) else None,"0.0%")]):
            c=ws1.cell(row=r,column=ci3+1,value=v); c.font=Font("Calibri",size=10,color=C_DK)
            c.fill=xfill(bg_s); c.border=xbdr()
            if f3: c.number_format=f3
            c.alignment=xctr() if ci3==0 else xrgt() if ci3 in [3,4] else xctr()
        ws1.row_dimensions[r].height=20; r+=1
    ws1.freeze_panes="A4"

    # ── Onglet 2 : Récap rayon
    ws2=wb_exp.create_sheet("Récap Rayon")
    title_block(ws2,"RÉCAPITULATIF PAR RAYON",span=10)
    write_header_row(ws2,4,["Rayon","CA (FCFA)","Poids CA","Marge (FCFA)","Tx Marge","Tx Marge HP","Tx Marge Promo","Écart HP−Promo","Pds Promo","Tx Casse"],[22,16,10,16,12,14,16,16,12,12])
    for ri4,(_,rd4) in enumerate(agg_rax.iterrows()):
        r4=ri4+5; bg4="F7F7F7" if ri4%2==0 else "FFFFFF"
        ecart4=rd4["TxMarge_HP"]-rd4["TxMarge_Promo"] if pd.notna(rd4.get("TxMarge_Promo")) else None
        for ci4,(v,f4) in enumerate([(rd4["lib_rayon"],None),(rd4["CA"],"#,##0"),
            (rd4["PoidsCA"]/100,"0.0%"),(rd4["Marge"],"#,##0"),
            (rd4["TxMarge"]/100 if pd.notna(rd4["TxMarge"]) else None,"0.0%"),
            (rd4["TxMarge_HP"]/100 if pd.notna(rd4.get("TxMarge_HP")) else None,"0.0%"),
            (rd4["TxMarge_Promo"]/100 if pd.notna(rd4.get("TxMarge_Promo")) else None,"0.0%"),
            (ecart4/100 if ecart4 is not None and pd.notna(ecart4) else None,"0.0%"),
            (rd4["PdsPromo"]/100 if pd.notna(rd4["PdsPromo"]) else None,"0.0%"),
            (rd4["TxCasse"]/100 if pd.notna(rd4["TxCasse"]) else None,"0.0%")]):
            c=ws2.cell(row=r4,column=ci4+1,value=v); c.font=Font("Calibri",size=10,color=C_DK)
            c.fill=xfill(bg4); c.border=xbdr()
            if f4: c.number_format=f4
            c.alignment=xrgt() if ci4 in [1,3] else xctr()
        ws2.row_dimensions[r4].height=20
    ws2.freeze_panes="A5"

    # ── Onglet 3 : Matrice
    ws3=wb_exp.create_sheet("Matrice Marge")
    title_block(ws3,"MATRICE TAUX DE MARGE — RAYON × MAGASIN",span=len(mat_pivot.columns)+2)
    c=ws3.cell(row=4,column=1,value="Rayon"); c.font=Font("Calibri",size=10,bold=True,color=C_WH)
    c.fill=xfill(C_SUB); c.alignment=xctr(); c.border=xbdr(); ws3.column_dimensions["A"].width=22
    for ci5,col_name in enumerate(mat_pivot.columns):
        c=ws3.cell(row=4,column=ci5+2,value=short_name(col_name))
        c.font=Font("Calibri",size=9,bold=True,color=C_WH); c.fill=xfill(C_SUB); c.alignment=xctr(); c.border=xbdr()
        ws3.column_dimensions[get_column_letter(ci5+2)].width=16
    ws3.row_dimensions[4].height=36
    for ri5,rayon_idx in enumerate(mat_pivot.index):
        r5=ri5+5
        c0=ws3.cell(row=r5,column=1,value=rayon_idx); c0.font=Font("Calibri",size=10,bold=True,color=C_DK)
        c0.fill=xfill("F7F7F7"); c0.alignment=xlft(); c0.border=xbdr()
        for ci5,col_name in enumerate(mat_pivot.columns):
            v5=mat_pivot.loc[rayon_idx,col_name]
            c=ws3.cell(row=r5,column=ci5+2)
            if pd.notna(v5): c.value=v5/100; c.number_format="0.0%"
            c.font=Font("Calibri",size=11,bold=True,color=C_DK); c.alignment=xctr(); c.border=xbdr(); c.fill=xfill("FFFFFF")
        ws3.row_dimensions[r5].height=28
    ws3.freeze_panes="B5"

    # ── Onglet 4 : Flop 100
    ws4=wb_exp.create_sheet("Flop 100")
    title_block(ws4,f"FLOP {len(flop100)} — DESTRUCTEURS DE MARGE",span=13)
    hdrs4=["#","Code Art.","Article","Rayon","Famille","CA (FCFA)","Marge (FCFA)","Tx Marge","Pds Promo","Qté",
           "🔵 HYPER","🟢 MARKET","🟣 SUPECO"]
    wdths4=[5,12,38,16,24,13,13,10,10,8,46,50,54]
    bloc_bg={10:C_HYP,11:C_MKT,12:C_SUP}
    for ci6,(h,w) in enumerate(zip(hdrs4,wdths4)):
        bg6=bloc_bg.get(ci6,C_SUB)
        c=ws4.cell(row=4,column=ci6+1,value=h); c.font=Font("Calibri",size=9,bold=True,color=C_WH)
        c.fill=xfill(bg6); c.alignment=xctr(); c.border=xbdr()
        ws4.column_dimensions[get_column_letter(ci6+1)].width=w
    ws4.row_dimensions[4].height=28
    bloc_fill={10:"D6EAF8",11:"D5F5E3",12:"E8DAEF"}
    for ri6,(_,rd6) in enumerate(flop100.iterrows()):
        r6=ri6+5; bg6="F7F7F7" if ri6%2==0 else "FFFFFF"
        tm6=rd6["TxMarge"]; pp6=rd6["PdsPromo"]
        vals6=[rd6["Rang"],rd6["code_art"],rd6["lib_art"],rd6["lib_rayon"],rd6["lib_fam"],
               rd6["CA"],rd6["Marge"],
               tm6/100 if pd.notna(tm6) else None,
               pp6/100 if pd.notna(pp6) else None,
               int(rd6["Qte"]) if pd.notna(rd6["Qte"]) else None,
               rd6["Bloc_Hyper"],rd6["Bloc_Market"],rd6["Bloc_Supeco"]]
        fmts6=[None,None,None,None,None,"#,##0","#,##0","0.0%","0.0%","#,##0",None,None,None]
        for ci6,(v,f6) in enumerate(zip(vals6,fmts6)):
            c=ws4.cell(row=r6,column=ci6+1,value=v)
            cell_bg=bloc_fill.get(ci6,bg6) if (ci6>=10 and v and v!="—") else bg6
            c.font=Font("Calibri",size=10 if ci6<10 else 9,color=C_DK)
            c.fill=xfill(cell_bg); c.border=xbdr()
            if f6: c.number_format=f6
            if ci6==0: c.font=Font("Calibri",size=10,bold=True,color=C_DK); c.alignment=xctr()
            elif ci6 in [5,6]: c.alignment=xrgt()
            elif ci6 in [7,8,9]: c.alignment=xctr()
            elif ci6>=10: c.alignment=xlft(w=True)
            else: c.alignment=xlft(w=(ci6 in [2,4]))
        ws4.row_dimensions[r6].height=30
    ws4.freeze_panes="A5"

    # ── Onglet 5 : Analyse Casse
    ws5=wb_exp.create_sheet("Analyse Casse")
    title_block(ws5,"ANALYSE CASSE RÉSEAU",span=9)
    kpi_casse=[("Casse Réseau (FCFA)",f"{abs(casse_reseau):,.0f}"),
               ("Tx Casse Réseau",f"{tx_casse_reseau:.2f}%"),
               ("Sites > 1% casse",str(nb_sites_alerte)),
               ("Tx Casse Moyen",f"{moy_tx_casse_site:.2f}%")]
    for ci_k,(lbl,val) in enumerate(kpi_casse):
        ck=ws5.cell(row=4,column=ci_k+1,value=lbl); ck.font=Font("Calibri",size=9,bold=True,color=C_WH)
        ck.fill=xfill(C_SUB); ck.alignment=xctr(); ck.border=xbdr()
        ck2=ws5.cell(row=5,column=ci_k+1,value=val); ck2.font=Font("Calibri",size=12,bold=True,color=C_DK)
        ck2.fill=xfill("FFFFFF"); ck2.alignment=xctr(); ck2.border=xbdr()
        ws5.column_dimensions[get_column_letter(ci_k+1)].width=22
    ws5.row_dimensions[4].height=20; ws5.row_dimensions[5].height=28

    r5s=7
    ws5.merge_cells(start_row=r5s,start_column=1,end_row=r5s,end_column=7)
    c=ws5.cell(row=r5s,column=1,value="  CLASSEMENT MAGASINS")
    c.font=Font("Calibri",size=10,bold=True,color=C_WH); c.fill=xfill(C_SUB); c.alignment=xlft(); ws5.row_dimensions[r5s].height=22; r5s+=1
    hdrs_cs=["Rang","Magasin","Format","CA (FCFA)","Casse (FCFA)","Casse (Qté)","Tx Casse"]
    write_header_row(ws5,r5s,hdrs_cs,[5,28,10,16,16,14,12]); r5s+=1
    for ri_s,(_,sd) in enumerate(agg_casse_site.iterrows()):
        bg_s="FFF8F0" if sd["TxCasse"]>1 else ("F7F7F7" if ri_s%2==0 else "FFFFFF")
        casse_q_val=int(sd["Casse_Q"]) if has_qty and pd.notna(sd.get("Casse_Q")) else None
        for ci_s,(v,f_s) in enumerate([(ri_s+1,None),(sd["lib_site"],None),(sd["format"],None),
            (sd["CA"],"#,##0"),(abs(sd["Casse_V"]),"#,##0"),(casse_q_val,"#,##0"),
            (sd["TxCasse"]/100 if pd.notna(sd["TxCasse"]) else None,"0.00%")]):
            c=ws5.cell(row=r5s,column=ci_s+1,value=v); c.font=Font("Calibri",size=10,color=C_DK)
            c.fill=xfill(bg_s); c.border=xbdr()
            if f_s: c.number_format=f_s
            c.alignment=xctr() if ci_s==0 else xrgt() if ci_s in [3,4,5] else xctr()
        ws5.row_dimensions[r5s].height=20; r5s+=1

    r5s+=1
    ws5.merge_cells(start_row=r5s,start_column=1,end_row=r5s,end_column=5)
    c=ws5.cell(row=r5s,column=1,value="  CASSE PAR RAYON")
    c.font=Font("Calibri",size=10,bold=True,color=C_WH); c.fill=xfill(C_SUB); c.alignment=xlft(); ws5.row_dimensions[r5s].height=22; r5s+=1
    write_header_row(ws5,r5s,["Rayon","CA (FCFA)","Casse (FCFA)","Tx Casse","% Casse Réseau"],[22,16,16,12,16]); r5s+=1
    for ri_r,(_,rd) in enumerate(agg_casse_rax.iterrows()):
        bg_r="F7F7F7" if ri_r%2==0 else "FFFFFF"
        pct_res=abs(rd["Casse_V"])/abs(casse_reseau) if casse_reseau!=0 else 0
        for ci_r,(v,f_r) in enumerate([(rd["lib_rayon"],None),(rd["CA"],"#,##0"),
            (abs(rd["Casse_V"]),"#,##0"),
            (rd["TxCasse"]/100 if pd.notna(rd["TxCasse"]) else None,"0.00%"),(pct_res,"0.0%")]):
            c=ws5.cell(row=r5s,column=ci_r+1,value=v); c.font=Font("Calibri",size=10,color=C_DK)
            c.fill=xfill(bg_r); c.border=xbdr()
            if f_r: c.number_format=f_r
            c.alignment=xlft() if ci_r==0 else xrgt() if ci_r in [1,2] else xctr()
        ws5.row_dimensions[r5s].height=20; r5s+=1

    r5s+=1
    ws5.merge_cells(start_row=r5s,start_column=1,end_row=r5s,end_column=9)
    c=ws5.cell(row=r5s,column=1,value="  TOP 30 ARTICLES — valeur de casse la plus élevée")
    c.font=Font("Calibri",size=10,bold=True,color=C_WH); c.fill=xfill(C_SUB); c.alignment=xlft(); ws5.row_dimensions[r5s].height=22; r5s+=1
    hdrs_a=["#","Code Art.","Article","Rayon","Famille","CA (FCFA)","Casse (FCFA)","Casse (Qté)","Tx Casse"]
    wdths_a=[5,12,38,18,24,16,16,14,12]
    write_header_row(ws5,r5s,hdrs_a,wdths_a); r5s+=1
    for ri_a,(_,ra) in enumerate(top30_casse.iterrows()):
        bg_a="F7F7F7" if ri_a%2==0 else "FFFFFF"
        casse_q_a=int(ra["Casse_Q"]) if has_qty and pd.notna(ra.get("Casse_Q")) else None
        for ci_a,(v,f_a) in enumerate([(ra["Rang"],None),(ra["code_art"],None),(ra["lib_art"],None),
            (ra["lib_rayon"],None),(ra["lib_fam"],None),
            (ra["CA"],"#,##0"),(abs(ra["Casse_V"]),"#,##0"),(casse_q_a,"#,##0"),
            (ra["TxCasse"]/100 if pd.notna(ra["TxCasse"]) else None,"0.00%")]):
            c=ws5.cell(row=r5s,column=ci_a+1,value=v); c.font=Font("Calibri",size=10,color=C_DK)
            c.fill=xfill(bg_a); c.border=xbdr()
            if f_a: c.number_format=f_a
            if ci_a==0: c.alignment=xctr()
            elif ci_a in [5,6,7]: c.alignment=xrgt()
            elif ci_a==8: c.alignment=xctr()
            else: c.alignment=xlft(w=(ci_a in [2,4]))
        ws5.row_dimensions[r5s].height=20; r5s+=1
    ws5.freeze_panes="A4"

    # ── Onglet 6 : Marges négatives par site ──────────────────────────
    ws6 = wb_exp.create_sheet("Marges Négatives par Site")
    title_block(ws6, "ARTICLES À MARGE NÉGATIVE — DÉTAIL PAR MAGASIN", span=11)

    # KPIs bandeau
    nb_lignes_neg    = len(art_neg_site)
    nb_art_uniq_neg  = art_neg_site["Article"].nunique()
    nb_sites_neg     = art_neg_site["lib_site"].nunique()
    perte_neg_total  = art_neg_site["Marge"].sum()
    ca_neg_total     = art_neg_site["CA"].sum()

    kpi_neg = [
        ("Lignes article × site",  str(nb_lignes_neg)),
        ("Articles distincts",      str(nb_art_uniq_neg)),
        ("Magasins touchés",        str(nb_sites_neg)),
        ("Pertes nettes (FCFA)",   f"{abs(perte_neg_total):,.0f}"),
        ("CA exposé (FCFA)",       f"{ca_neg_total:,.0f}"),
    ]
    for ci_n, (lbl, val) in enumerate(kpi_neg):
        ck = ws6.cell(row=4, column=ci_n + 1, value=lbl)
        ck.font = Font("Calibri", size=9, bold=True, color=C_WH)
        ck.fill = xfill(C_SUB); ck.alignment = xctr(); ck.border = xbdr()
        ck2 = ws6.cell(row=5, column=ci_n + 1, value=val)
        ck2.font = Font("Calibri", size=12, bold=True, color=C_DK)
        ck2.fill = xfill("FFFFFF"); ck2.alignment = xctr(); ck2.border = xbdr()
        ws6.column_dimensions[get_column_letter(ci_n + 1)].width = 22
    ws6.row_dimensions[4].height = 20
    ws6.row_dimensions[5].height = 28
    ws6.row_dimensions[6].height = 8  # espace vide

    # En-têtes tableau
    r6s = 7
    hdrs6  = ["#", "Code Art.", "Article", "Rayon", "Famille",
              "Magasin", "Format", "CA (FCFA)", "Marge (FCFA)", "Tx Marge", "Qté"]
    wdths6 = [5,   12,          38,        18,      24,
              28,       10,       14,          14,           10,        8]
    write_header_row(ws6, r6s, hdrs6, wdths6)
    r6s += 1

    # Palette de fond par rayon (alternance de teintes douces)
    rayon_palette = [
        "EBF5FB","E8F8F5","FEF9E7","F5EEF8","FDEDEC",
        "EAF4FB","E9F7EF","FDFEFE","F4ECF7","FEF5E7",
    ]
    rayon_color_map = {}
    for rayon_name in art_neg_site["lib_rayon"].dropna().unique():
        idx = len(rayon_color_map) % len(rayon_palette)
        rayon_color_map[rayon_name] = rayon_palette[idx]

    for ri6, row6 in art_neg_site.reset_index(drop=True).iterrows():
        rayon_v   = row6.get("lib_rayon")
        fam_v     = row6.get("lib_fam")
        art_lib_v = row6.get("lib_art")
        code_v    = row6.get("code_art")
        tm_v      = row6["TxMarge_site"]
        qte_v     = int(row6["Qte"]) if pd.notna(row6.get("Qte")) else None

        # Fond par sévérité de marge
        if pd.notna(tm_v) and tm_v < -5:
            bg6 = "FDECEA"   # rouge pâle — marge < -5%
        elif pd.notna(tm_v) and tm_v < 0:
            bg6 = "FFF3E0"   # orange pâle — marge entre -5% et 0%
        else:
            bg6 = rayon_color_map.get(rayon_v, "FFFFFF")

        vals6 = [
            ri6 + 1,
            code_v,
            art_lib_v,
            rayon_v,
            fam_v,
            row6["lib_site"],
            row6["format"],
            row6["CA"],
            row6["Marge"],
            (tm_v / 100) if pd.notna(tm_v) else None,
            qte_v,
        ]
        fmts6 = [
            None, None, None, None, None, None, None,
            "#,##0", "#,##0", "0.0%", "#,##0"
        ]

        for ci6_v, (v, f6) in enumerate(zip(vals6, fmts6)):
            c = ws6.cell(row=r6s, column=ci6_v + 1, value=v)
            c.fill   = xfill(bg6)
            c.border = xbdr()
            if f6:
                c.number_format = f6

            # Police : taux de marge en rouge gras si négatif
            if ci6_v == 9 and pd.notna(tm_v) and tm_v < 0:
                c.font = Font("Calibri", size=10, bold=True, color="C0392B")
            else:
                c.font = Font("Calibri", size=10, color=C_DK)

            # Alignement
            if ci6_v == 0:
                c.alignment = xctr()
            elif ci6_v in [7, 8, 10]:
                c.alignment = xrgt()
            elif ci6_v == 9:
                c.alignment = xctr()
            else:
                c.alignment = xlft(w=(ci6_v in [2, 4]))

        ws6.row_dimensions[r6s].height = 20
        r6s += 1

    ws6.freeze_panes = f"A8"
    ws6.auto_filter.ref = (
        f"A7:{get_column_letter(len(hdrs6))}{r6s - 1}"
    )

    # ── Sauvegarde ────────────────────────────────────────────────────
    buf = BytesIO()
    wb_exp.save(buf)
    buf.seek(0)
    return buf
