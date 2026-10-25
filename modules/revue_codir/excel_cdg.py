"""Classeur des chefs de groupe : chiffres en valeurs (aucune formule) et cellules jaunes « Causes et actions ».

write_workbook(R, path_or_buffer)  -> classeur à déposer sur Teams à 08:15
read_comments(files)               -> {clé famille: {...}, '__general__': {secteur: texte}} relu au gel de 10:10
"""
from __future__ import annotations

import io
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from . import engine as E

NAVY, ORANGE, GREY, YEL = '24235C', 'F19F39', 'ECEDF2', 'FFF2CC'
F = 'Arial'
CAUSES = ['Prix / promo', 'Rupture / supply', 'Dégagement DLC', 'Effet de base N-1', 'Mix / gros volumes',
          'Fréquentation', 'Concurrence', 'Opération commerciale', 'Donnée à vérifier', 'Autre']
FIELDS = ['Cause', 'Commentaire', 'Action', 'Responsable', 'Échéance', 'Décision DG']
HEAD = Font(name=F, size=10, bold=True, italic=True, color='FFFFFF')
NORM = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
RED, GREEN = '9A1C14', '1B6B34'
thin = Side(style='thin', color='C9CCDB')
ybd = Side(style='thin', color='BF8F00')


def _title(ws, r, txt, sub=None):
    ws.cell(r, 1, txt).font = Font(name=F, size=15, bold=True, color=NAVY)
    if sub:
        ws.cell(r + 1, 1, sub).font = Font(name=F, size=9, italic=True, color='4E5370')


def _band(ws, r, txt, ncol=9):
    for c in range(1, ncol + 1):
        ws.cell(r, c).fill = PatternFill('solid', fgColor=NAVY)
    ws.cell(r, 1, txt.upper()).font = Font(name=F, size=11, bold=True, color='FFFFFF')
    ws.row_dimensions[r].height = 20


def _header(ws, r, labels, c0=1):
    for j, t in enumerate(labels):
        x = ws.cell(r, c0 + j, t)
        x.font = HEAD
        x.fill = PatternFill('solid', fgColor=ORANGE)
        x.alignment = Alignment(horizontal='left' if j == 0 else 'center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 28


def _val(ws, r, c, v, fmt=None, bold=False, sign=False, fill=None):
    x = ws.cell(r, c, v)
    x.font = Font(name=F, size=10, bold=bold, color=(RED if sign and isinstance(v, (int, float)) and v < 0 else GREEN if sign and isinstance(v, (int, float)) and v > 0 else '1E2140'))
    if fmt:
        x.number_format = fmt
    x.alignment = Alignment(horizontal='right' if isinstance(v, (int, float)) else 'left', vertical='center')
    x.border = Border(bottom=thin)
    if fill:
        x.fill = PatternFill('solid', fgColor=fill)
    return x


M1 = '#,##0.0;-#,##0.0;0.0'
MS = '+#,##0.0;-#,##0.0;0.0'
PC = '+0.0%;-0.0%;0.0%'
PT = '+0.0" pt";-0.0" pt";0.0" pt"'


def _yellow(ws, r, c1, c2=None):
    c2 = c2 or c1
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    for c in range(c1, c2 + 1):
        x = ws.cell(r, c)
        x.fill = PatternFill('solid', fgColor=YEL)
        x.protection = Protection(locked=False)
        x.border = Border(left=ybd, right=ybd, top=ybd, bottom=ybd)
        x.alignment = Alignment(wrap_text=True, vertical='top')
        x.font = NORM


def write_workbook(R: dict, out) -> None:
    wb = openpyxl.Workbook()
    meta = R['meta']
    sem = f"S{meta['week']}"

    # ---------------- mode d'emploi
    ws = wb.active
    ws.title = "Mode d'emploi"
    ws.sheet_view.showGridLines = False
    _title(ws, 1, f'Revue de performance · {sem}', f"{meta['periode']} · Vs N-1 · M FCFA HT · données au {meta['data_date']:%d/%m/%Y} 08:00")
    lines = [
        ('Ce que vous devez faire', True),
        ('1. Ouvrez l’onglet de votre secteur (PGC, Frais, Non-alimentaire).', False),
        ('2. Lisez le résultat et les familles qui reculent le plus.', False),
        ('3. Remplissez les cellules jaunes « Causes et actions » des 3 familles prioritaires (les 2 suivantes sont facultatives).', False),
        ('4. Ajoutez un commentaire général d’une ou deux phrases (300 caractères maximum).', False),
        ('', False),
        ('Heure limite : 10:10. Le rapport part au directeur à 10:30 quoi qu’il arrive.', True),
        ('Sans commentaire, le rapport indique « En attente : présenté en séance par [chef de groupe] ».', False),
        ('', False),
        ('Repères', True),
        ('Écart de marge = marge N − marge N-1 (M FCFA). Effet volume = (CA N − CA N-1) × taux N-1. Effet taux = CA N × (taux N − taux N-1).', False),
        ('Familles retenues : CA d’au moins 1 M FCFA en N ou en N-1, classées par écart de marge en valeur.', False),
        ('Ne modifiez pas les autres cellules : le fichier est relu automatiquement.', False),
    ]
    for i, (t, b) in enumerate(lines):
        ws.cell(4 + i, 1, t).font = Font(name=F, size=11, bold=b, color=NAVY if b else '1E2140')
    ws.column_dimensions['A'].width = 120

    # ---------------- listes
    ls = wb.create_sheet('Listes')
    for i, c in enumerate(CAUSES):
        ls.cell(1 + i, 1, c)
    ls.cell(1, 2, 'Oui')
    ls.cell(2, 2, 'Non')
    ls.sheet_state = 'hidden'

    # ---------------- onglets secteur
    for s in E.SECTEUR_ORDER:
        ws = wb.create_sheet(s, index=len(wb.sheetnames) - 1)
        ws.sheet_view.showGridLines = False
        o = R['sect_detail'][s]
        _title(ws, 1, f'{s} · {sem}', f"{E.CHEFS[s]} · {meta['periode']} · Vs N-1 · M FCFA HT")
        r = 4
        _band(ws, r, "Chiffre d'affaires et marge (Vs N-1)")
        r += 1
        _header(ws, r, ['Indicateur', 'Réel', 'N-1', 'Δ', 'Δ %'])
        r += 1
        rows = [("Chiffre d'affaires", o['CA'] / 1e6, o['CA1'] / 1e6, o['dCA'] / 1e6, o['ca'], M1),
                ('Marge', o['M'] / 1e6, o['M1'] / 1e6, o['dM'] / 1e6, o['M'] / o['M1'] - 1 if o['M1'] else None, M1),
                ('Taux de marge', o['t'], o['t1'], o['dt'], None, '0.0%')]
        for lab, a, b, d, p, f in rows:
            _val(ws, r, 1, lab, bold=True)
            _val(ws, r, 2, a, f, bold=True)
            _val(ws, r, 3, b, f)
            _val(ws, r, 4, d, PT if lab == 'Taux de marge' else MS, bold=True, sign=True)
            if p is not None:
                _val(ws, r, 5, p, PC, sign=True)
            r += 1
        r += 1
        ws.cell(r, 1, 'Lecture : ' + R['texts']['lecture'][s]).font = Font(name=F, size=10, italic=True, color='33385A')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 30
        r += 2

        _band(ws, r, 'Écart de marge par format (Vs N-1, M FCFA)')
        r += 1
        _header(ws, r, ['Format'] + E.FORMAT_ORDER)
        r += 1
        _val(ws, r, 1, 'Écart de marge', bold=True)
        for j, f in enumerate(E.FORMAT_ORDER):
            _val(ws, r, 2 + j, R['familles'][s]['fmt'][f] / 1e6, MS, bold=True, sign=True)
        r += 2

        _band(ws, r, 'Familles : écart de marge (dont volume / dont taux)')
        r += 1
        _header(ws, r, ['Famille', 'Rayon', 'CA (M)', 'CA Vs N-1', 'Écart marge (M)', 'dont volume', 'dont taux', 'Évol. taux'])
        r += 1
        for part, lab in [('top', 'Top contributeurs'), ('flop', 'Principaux détracteurs')]:
            ws.cell(r, 1, lab).font = Font(name=F, size=10, bold=True, color=GREEN if part == 'top' else RED)
            r += 1
            for f_ in R['familles'][s][part]:
                _val(ws, r, 1, f_['label'], bold=True)
                _val(ws, r, 2, f_['rlabel'])
                _val(ws, r, 3, f_['CA'] / 1e6, M1)
                _val(ws, r, 4, f_['ca'] if f_['CA1'] else None, PC, sign=True)
                _val(ws, r, 5, f_['dM'] / 1e6, MS, bold=True, sign=True)
                _val(ws, r, 6, f_['effV'] / 1e6, MS, sign=True)
                _val(ws, r, 7, f_['effT'] / 1e6, MS, sign=True)
                _val(ws, r, 8, f_['dt'], PT, sign=True)
                r += 1
        r += 1

        _band(ws, r, 'Causes et actions (à compléter avant 10:10)')
        r += 1
        ws.cell(r, 1, 'Commentaire général du secteur (300 caractères maximum)').font = BOLD
        r += 1
        _yellow(ws, r, 1, 9)
        ws.row_dimensions[r].height = 44
        ws.cell(r, 11, '__general__').font = Font(name=F, size=8, color='FFFFFF')
        dvg = DataValidation(type='textLength', operator='lessThanOrEqual', formula1='300', showErrorMessage=True, error='300 caractères maximum.')
        ws.add_data_validation(dvg)
        dvg.add(ws.cell(r, 1))
        r += 2
        _header(ws, r, ['Famille', 'Écart (M)', 'Cause', 'Commentaire (150 car.)', '', 'Action (150 car.)', '', 'Responsable', 'Échéance', 'Décision DG'])
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        r += 1
        dvc = DataValidation(type='list', formula1='=Listes!$A$1:$A$10', allow_blank=True, showErrorMessage=True, error='Choisir dans la liste.')
        dvo = DataValidation(type='list', formula1='=Listes!$B$1:$B$2', allow_blank=True, showErrorMessage=True)
        dvt = DataValidation(type='textLength', operator='lessThanOrEqual', formula1='150', showErrorMessage=True, error='150 caractères maximum.')
        dvd = DataValidation(type='date', operator='greaterThan', formula1='45000', allow_blank=True, showErrorMessage=True, error='Saisir une date (jj/mm/aaaa).')
        for dv in (dvc, dvo, dvt, dvd):
            ws.add_data_validation(dv)
        for i, f_ in enumerate(R['familles'][s]['flop']):
            _val(ws, r, 1, f_['label'] + ('' if i < 3 else ' (facultatif)'), bold=i < 3)
            _val(ws, r, 2, f_['dM'] / 1e6, MS, sign=True)
            _yellow(ws, r, 3)
            _yellow(ws, r, 4, 5)
            _yellow(ws, r, 6, 7)
            _yellow(ws, r, 8)
            _yellow(ws, r, 9)
            _yellow(ws, r, 10)
            ws.cell(r, 9).number_format = 'dd/mm/yyyy'
            dvc.add(ws.cell(r, 3))
            dvt.add(ws.cell(r, 4))
            dvt.add(ws.cell(r, 6))
            dvd.add(ws.cell(r, 9))
            dvo.add(ws.cell(r, 10))
            ws.cell(r, 11, f_['key']).font = Font(name=F, size=8, color='FFFFFF')
            ws.row_dimensions[r].height = 36
            r += 1

        for c, w in zip('ABCDEFGHIJK', [34, 22, 18, 22, 16, 22, 16, 18, 13, 12, 4]):
            ws.column_dimensions[c].width = w
        ws.column_dimensions['K'].hidden = True
        ws.protection.sheet = True
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        ws.freeze_panes = 'A4'

    wb.active = 0
    wb.save(out)


def _txt(v):
    if v is None:
        return ''
    if isinstance(v, (datetime, date)):
        return v.strftime('%d/%m/%Y')
    return str(v).strip()


def read_comments(files) -> dict:
    """Relit un ou plusieurs classeurs complétés. La dernière saisie non vide l'emporte."""
    out = {'__general__': {}}
    for f in files:
        src = io.BytesIO(f.read()) if hasattr(f, 'read') else f
        wb = openpyxl.load_workbook(src, data_only=True)
        for s in E.SECTEUR_ORDER:
            if s not in wb.sheetnames:
                continue
            ws = wb[s]
            for r in range(1, ws.max_row + 1):
                key = ws.cell(r, 11).value
                if not key:
                    continue
                if key == '__general__':
                    t = _txt(ws.cell(r, 1).value)
                    if t:
                        out['__general__'][s] = t
                    continue
                vals = {'Cause': _txt(ws.cell(r, 3).value), 'Commentaire': _txt(ws.cell(r, 4).value), 'Action': _txt(ws.cell(r, 6).value),
                        'Responsable': _txt(ws.cell(r, 8).value), 'Échéance': _txt(ws.cell(r, 9).value), 'Décision DG': _txt(ws.cell(r, 10).value)}
                if any(vals.values()):
                    out[key] = vals
    return out


def completeness(R: dict) -> dict:
    """Nombre de familles prioritaires (3 par secteur) avec cause, commentaire et action renseignés."""
    res = {}
    for s in E.SECTEUR_ORDER:
        pr = R['familles'][s]['flop'][:3]
        ok = sum(1 for f in pr if all(f.get('cdg', {}).get(k) for k in ('Cause', 'Commentaire', 'Action')))
        res[s] = (ok, len(pr))
    return res
