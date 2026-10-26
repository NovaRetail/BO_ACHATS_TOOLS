"""Rapport PowerPoint à la Direction Générale, généré sur le masque de l'entreprise.

build_report(R, masque, out) : 4 pages principales + annexe (A1 à A4).
Mise en page « exécutive » : un message par page (titre), gros chiffres, graphiques natifs modifiables,
un encadré « À retenir », beaucoup d'air. Tous les chiffres viennent de `R` (engine.compute).
"""
from __future__ import annotations

import copy

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import engine as E

# ------------------------------------------------------------------ charte
NAVY, ORANGE = '24235C', 'F19F39'
INK, MUTE, LINE, SOFT, PAPER = '1E2140', '6B7090', 'DADDE6', 'F4F5F8', 'FFFFFF'
NEG, POS, SMALL = 'C0392B', '1E8449', 'B5B9CC'
NEG_T, POS_T = 'B3261E', '1B7A45'
NEG_BG, POS_BG = 'FBEAE8', 'E7F4EC'
N1_BAR = '9FB0D6'
X0, W = 0.61, 12.1           # marge gauche et largeur utile
Y0 = 1.5                     # haut de la zone de contenu (sous le bandeau du masque)


def rgb(h):
    return RGBColor.from_string(h)


def sc(v, small=False):
    return MUTE if small else (NEG_T if v < 0 else POS_T)


# ------------------------------------------------------------------ outils
def _drop_all_slides(prs):
    sld = prs.slides._sldIdLst
    for s in list(sld):
        prs.part.drop_rel(s.get(qn('r:id')))
        sld.remove(s)


def _layout(prs, name):
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    raise KeyError(name)


def _clone_footer(slide, layout, footer_text):
    """Recopie le pied de page et le numéro de page du masque (python-pptx ne les crée pas)."""
    for ph in layout.placeholders:
        name = str(ph.placeholder_format.type)
        if 'FOOTER' in name or 'SLIDE_NUMBER' in name:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))
            if 'FOOTER' in name:
                slide.shapes[-1].text_frame.text = footer_text


def tb(slide, x, y, w, h, runs, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       italic=False, spacing=None, line=None):
    """runs : texte simple, ou liste de paragraphes = listes de (texte, {options})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        if line:
            p.line_spacing = line
        for text, o in para:
            r = p.add_run()
            r.text = text
            f = r.font
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.italic = o.get('italic', italic)
            f.color.rgb = rgb(o.get('color', color))
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def hline(slide, x, y, w, color=LINE, weight=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(weight)
    return ln


def eyebrow(slide, x, y, w, text, color=MUTE):
    """Petit sur-titre en capitales, avec un trait orange."""
    rect(slide, x, y + 0.07, 0.22, 0.04, ORANGE)
    tb(slide, x + 0.3, y, w - 0.3, 0.22, [[(text.upper(), {'bold': True, 'size': 9.5, 'color': color})]])


def pill(slide, x, y, text, v, small=False, w=None, size=11):
    w = w or (0.3 + 0.1 * len(text))
    fill = SOFT if small else (NEG_BG if v < 0 else POS_BG)
    s = rect(slide, x, y, w, 0.3, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = rgb(sc(v, small))
    return s


def takeaway(slide, x, y, w, h, text, label='À retenir'):
    """Encadré de conclusion : fond bleu nuit, libellé orange."""
    rect(slide, x, y, w, h, NAVY)
    tb(slide, x + 0.25, y, w - 0.5, h, [[(label.upper() + '   ', {'bold': True, 'size': 10, 'color': ORANGE}),
                                         (text, {'size': 12.5, 'color': 'FFFFFF', 'bold': True})]], anchor=MSO_ANCHOR.MIDDLE)


def table(slide, x, y, w, rows, col_w, header, size=11, row_h=0.42, aligns=None, colors=None, bolds=None, fills=None):
    """Tableau épuré : en-tête bleu nuit, lignes blanches séparées d'un filet."""
    nr, nc = len(rows) + 1, len(header)
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for attr in ('bandRow', 'firstRow'):
        tblPr.set(attr, '0')
    st = tblPr.find(qn('a:tableStyleId'))
    if st is not None:
        st.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    aligns = aligns or (['l'] + ['r'] * (nc - 1))

    def cell(i, j, text, fill, color, bold, sz):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(fill)
        c.margin_left = c.margin_right = Inches(0.08)
        c.margin_top = c.margin_bottom = Inches(0.03)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tcPr = c._tc.get_or_add_tcPr()
        for edge in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
            ln = tcPr.find(qn(edge))
            if ln is not None:
                tcPr.remove(ln)
        for edge, show in (('a:lnL', False), ('a:lnR', False), ('a:lnT', False), ('a:lnB', i > 0)):
            ln = tcPr.makeelement(qn(edge), {'w': '9525' if show else '0'})
            if show:
                sf = ln.makeelement(qn('a:solidFill'), {})
                clr = sf.makeelement(qn('a:srgbClr'), {'val': LINE})
                sf.append(clr)
                ln.append(sf)
            else:
                ln.append(ln.makeelement(qn('a:noFill'), {}))
            tcPr.append(ln)
        p = c.text_frame.paragraphs[0]
        p.alignment = {'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT, 'c': PP_ALIGN.CENTER}[aligns[j]]
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)

    for j, hd in enumerate(header):
        cell(0, j, hd.upper(), NAVY, 'FFFFFF', True, size - 2)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            fill = fills[i - 1][j] if fills and fills[i - 1][j] else PAPER
            col = colors[i - 1][j] if colors and colors[i - 1][j] else INK
            b = bolds[i - 1][j] if bolds else (j == 0)
            cell(i, j, v, fill, col, b, size)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
    return gt


def bar_chart(slide, x, y, w, h, cats, vals, size=11, label_fmt='+#,##0.0;-#,##0.0;0.0'):
    """Barres horizontales natives (modifiables), vert si positif, rouge si négatif, gris si non significatif."""
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series('Écart de marge (M FCFA)', [v / 1e6 for v in vals])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(size)
    ch.font.color.rgb = rgb(INK)
    va, ca = ch.value_axis, ch.category_axis
    va.visible = False
    va.has_major_gridlines = False
    lim = max(abs(v) for v in vals) / 1e6 * 1.25
    va.minimum_scale, va.maximum_scale = -lim, lim
    ca.reverse_order = True
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.color.rgb = rgb(LINE)
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(size)
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = label_fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(size)
    dl.font.bold = True
    ser = plot.series[0]
    ser.invert_if_negative = False
    for i, v in enumerate(vals):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(SMALL if abs(v) < E.SEUIL_SIGNIF else (NEG if v < 0 else POS))
        d = pt.data_label
        d.has_text_frame = True
        d.text_frame.text = E.fdm(v)
        d.position = XL_LABEL_POSITION.OUTSIDE_END
        r = d.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = rgb(sc(v, abs(v) < E.SEUIL_SIGNIF))
    return gf


def heat(v):
    if abs(v) < E.SEUIL_SIGNIF:
        return SOFT, MUTE
    t = min(abs(v) / 5e6, 1)
    base = (192, 57, 43) if v < 0 else (30, 132, 73)
    a = 0.12 + 0.6 * t
    fill = '%02X%02X%02X' % tuple(int(255 + (c - 255) * a) for c in base)
    return fill, ('FFFFFF' if t > 0.7 else (NEG_T if v < 0 else POS_T))


# ------------------------------------------------------------------ pages
def _title(slide, light, bold):
    tf = slide.shapes.title.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = (light + ' : ').upper() if light else ''
    r1.font.bold = False
    r2 = p.add_run()
    r2.text = bold.upper()
    r2.font.bold = True
    size = 26 if len(r1.text + r2.text) <= 46 else 22
    r1.font.size = r2.font.size = Pt(size)


def _content(prs, light, bold, footer, R):
    lay = _layout(prs, 'Titre seul')
    s = prs.slides.add_slide(lay)
    _title(s, light, bold)
    _clone_footer(s, lay, footer)
    if R.get('demo'):
        tb(s, 9.9, 1.25, 2.8, 0.22, [[('DONNÉES FICTIVES · DÉMONSTRATION', {'bold': True, 'size': 8.5, 'color': ORANGE})]], align=PP_ALIGN.RIGHT)
    return s


def _cover(prs, R):
    s = prs.slides.add_slide(_layout(prs, '14_Couverture'))
    tf = s.shapes.title.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    for txt, b in (('REVUE DE PERFORMANCE ', False), ('COMMERCIALE', True)):
        r = p.add_run()
        r.text = txt
        r.font.bold = b
    m = R['meta']
    box = tb(s, 0.94, 4.78, 7.4, 0.38, [[(f"Semaine {m['week']} · {m['periode']}", {'italic': True, 'color': NAVY, 'size': 14, 'bold': True})]],
             anchor=MSO_ANCHOR.MIDDLE)
    box.fill.solid()
    box.fill.fore_color.rgb = rgb('FFFFFF')
    meta = [[(f"Vs N-1 · M FCFA HT · {m['n_sites']} magasins", {'size': 11, 'color': 'FFFFFF', 'bold': True})],
            [(f"Données au lundi {m['data_date']:%d/%m/%Y}, 08:00", {'size': 10, 'color': 'C5C8E0'})],
            [('Confidentiel · usage interne', {'size': 10, 'color': 'C5C8E0'})]]
    if R.get('demo'):
        meta.append([('Données fictives · démonstration', {'size': 11, 'color': ORANGE, 'bold': True})])
    tb(s, 0.94, 5.3, 6.2, 1.2, meta, spacing=2)


def _p2(prs, R, ftr):
    g, t = R['global'], R['texts']
    s = _content(prs, 'Synthèse', t['p2_title'], ftr, R)
    # 5 indicateurs : gros chiffres, écart en pastille, N-1 en petit
    K = [("Chiffre d'affaires", E.fm(g['CA']), 'M FCFA', g['ca'], E.fpct(g['ca']), f"N-1 {E.fm(g['CA1'])}"),
         ('Marge', E.fm(g['M']), 'M FCFA', g['dM'], f"{E.fdm(g['dM'])} M", f"N-1 {E.fm(g['M1'])}"),
         ('Taux de marge', E.fnum(g['t'] * 100), '%', g['dt'], E.fpt(g['dt']), f"N-1 {E.fnum(g['t1'] * 100)} %"),
         ('Trafic', E.fnum(g['T'], 0), 'tickets', g['tk'], E.fpct(g['tk']), f"N-1 {E.fnum(g['T1'], 0)}"),
         ('Panier moyen', E.fnum(g['panier'], 0), 'FCFA', g['pan'], E.fpct(g['pan']), f"N-1 {E.fnum(g['panier1'], 0)}")]
    cw, gap = 2.26, 0.2
    for i, (lab, v, u, d, ds, n1) in enumerate(K):
        x = X0 + i * (cw + gap)
        tb(s, x, Y0, cw, 0.24, [[(lab.upper(), {'bold': True, 'size': 9.5, 'color': MUTE})]])
        tb(s, x, Y0 + 0.26, cw, 0.62, [[(v, {'bold': True, 'size': 30, 'color': NAVY}), ('  ' + u, {'size': 11, 'color': MUTE})]])
        pill(s, x, Y0 + 0.98, ('▼ ' if d < 0 else '▲ ') + ds, d)
        tb(s, x + 1.2, Y0 + 1.0, cw - 1.2, 0.28, [[(n1, {'size': 9.5, 'color': MUTE})]], align=PP_ALIGN.RIGHT)
    hline(s, X0, Y0 + 1.5, W)
    # ce qu'il faut retenir
    yb = Y0 + 1.75
    eyebrow(s, X0, yb, 5.9, 'Ce qu’il faut retenir')
    pts = t['faits'] + [t['taux']]
    for i, txt in enumerate(pts[:4]):
        yy = yb + 0.4 + i * 0.66
        c = rect(s, X0, yy + 0.02, 0.34, 0.34, NAVY, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i + 1)
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = rgb('FFFFFF')
        head, sep, rest = txt.partition(' : ')
        runs = [(head + ' : ', {'bold': True}), (rest, {})] if sep and len(head) < 30 else [(txt, {})]
        tb(s, X0 + 0.5, yy, 5.5, 0.62, [runs], size=12.5, anchor=MSO_ANCHOR.MIDDLE)
    # priorités de la semaine
    xr = 6.95
    eyebrow(s, xr, yb, 5.75, 'Familles à surveiller')
    flops = sorted([(sc_, f_) for sc_ in E.SECTEUR_ORDER for f_ in R['familles'][sc_]['flop'][:3]], key=lambda x: x[1]['dM'])[:3]
    for i, (sct, f_) in enumerate(flops):
        yy = yb + 0.4 + i * 0.8
        rect(s, xr, yy, 5.76, 0.68, SOFT)
        rect(s, xr, yy, 0.06, 0.68, NEG if f_['dM'] < 0 else POS)
        ca = f"CA {E.fpct(f_['ca'])} · " if f_['CA1'] else ''
        tb(s, xr + 0.2, yy + 0.06, 4.3, 0.6, [[(f_['label'], {'bold': True, 'size': 12}), (f"  ·  {sct}", {'size': 10, 'color': MUTE})],
                                                [(f"{ca}dont volume {E.fdm(f_['effV'])} M · dont taux {E.fdm(f_['effT'])} M", {'size': 10, 'color': MUTE})]])
        tb(s, xr + 4.5, yy, 1.15, 0.68, [[(f"{E.fdm(f_['dM'])} M", {'bold': True, 'size': 14, 'color': sc(f_['dM'])})]],
           align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # contexte en une ligne
    ctx = (f"Contexte : {t['contexte']} · Ex-Supeco convertis : " + ', '.join(f"{n.replace('Market ', '')} {d[:5]}" for n, d in E.CONVERSIONS)
           + ' · À vérifier : ' + ', '.join(E.SITES_A_VERIFIER) + ' · Vs N-1 uniquement, budget au reporting mensuel.')
    tb(s, X0, 6.62, W, 0.3, [[(ctx, {'size': 9, 'color': MUTE})]])


def _p3(prs, R, ftr):
    g, t = R['global'], R['texts']
    s = _content(prs, 'Marge', t['p3_title'], ftr, R)
    eyebrow(s, X0, Y0, 6.6, 'Passage de la marge N-1 → N (M FCFA)')
    steps = [('Marge N-1', g['M1'], 'tot')] + [('Non-alim.' if x == 'Non-alimentaire' else x, R['sect'][x]['dM'], 'd') for x in E.SECTEUR_ORDER] + [('Marge N', g['M'], 'tot')]
    cats, base, up, cols, labels = [], [], [], [], []
    cur = g['M1']
    for n, v, k in steps:
        cats.append(n)
        if k == 'tot':
            base.append(0)
            up.append(v / 1e6)
            cols.append(N1_BAR if n == 'Marge N-1' else NAVY)
            labels.append(E.fm(v))
        else:
            base.append(min(cur, cur + v) / 1e6)
            up.append(abs(v) / 1e6)
            cols.append(SMALL if abs(v) < E.SEUIL_SIGNIF else (NEG if v < 0 else POS))
            labels.append(E.fdm(v))
            cur += v
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series('base', base)
    cd.add_series('valeur', up)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(X0), Inches(Y0 + 0.35), Inches(6.6), Inches(3.95), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(11)
    lows = [b for b in base if b > 0] + [g['M1'] / 1e6, g['M'] / 1e6]
    ch.value_axis.minimum_scale = int(min(lows) * 0.95 // 5 * 5)
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.format.line.color.rgb = rgb(LINE)
    ch.category_axis.tick_labels.font.size = Pt(11)
    ch.category_axis.tick_labels.font.color.rgb = rgb(MUTE)
    plot = ch.plots[0]
    plot.gap_width = 55
    sb, sv = plot.series
    sb.format.fill.background()
    for i, c in enumerate(cols):
        pt = sv.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        dl = pt.data_label
        dl.has_text_frame = True
        dl.text_frame.text = labels[i]
        dl.position = XL_LABEL_POSITION.INSIDE_END
        r = dl.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = rgb('FFFFFF' if c not in (SMALL,) else INK)
    tb(s, X0, Y0 + 4.32, 6.6, 0.25, [[(f"≈ Axe tronqué à {ch.value_axis.minimum_scale:.0f} M FCFA", {'size': 9, 'color': MUTE, 'italic': True})]])
    # heatmap
    xr = 7.55
    eyebrow(s, xr, Y0, 5.15, 'Où se joue l’écart · secteur × format (M FCFA)')
    hdr = ['', 'Hyper', 'Market', 'Ex-Sup.', 'Supeco', 'Total']
    rows, fills, colors, bolds = [], [], [], []
    for sct in E.SECTEUR_ORDER + ['Total']:
        vals = [R['mtx'][sct][f] for f in E.FORMAT_ORDER] + [R['sect'][sct]['dM']] if sct != 'Total' else [R['fmt'][f]['dM'] for f in E.FORMAT_ORDER] + [g['dM']]
        rows.append(['Non-alim.' if sct == 'Non-alimentaire' else sct] + [E.fdm(v) for v in vals])
        hs = [heat(v) for v in vals]
        fills.append([None] + [h[0] for h in hs])
        colors.append([None] + [h[1] for h in hs])
        bolds.append([True] + [True] * len(vals))
    table(s, xr, Y0 + 0.4, 5.15, rows, [1.15, 0.8, 0.8, 0.8, 0.8, 0.8], hdr, size=12, row_h=0.52, fills=fills, colors=colors, bolds=bolds)
    tb(s, xr, Y0 + 3.1, 5.15, 0.3, [[('Gris : écart inférieur à 0,5 M FCFA, non significatif.', {'size': 9, 'color': MUTE})]])
    worst = min(((sct, f, R['mtx'][sct][f]) for sct in E.SECTEUR_ORDER for f in E.FORMAT_ORDER), key=lambda x: x[2])
    best = max(((sct, f, R['mtx'][sct][f]) for sct in E.SECTEUR_ORDER for f in E.FORMAT_ORDER), key=lambda x: x[2])
    takeaway(s, X0, 6.2, W, 0.55, f"Plus forte perte : {worst[0]} en {worst[1]} ({E.fdm(worst[2])} M). "
                                   f"Plus fort gain : {best[0]} en {best[1]} ({E.fdm(best[2])} M).")


def _p4(prs, R, ftr):
    s = _content(prs, 'Écarts', R['texts']['p4_title'], ftr, R)
    C = R['contrib']
    for k, x, lab in (('rayon', X0, 'Par rayon'), ('site', 6.95, 'Par magasin')):
        eyebrow(s, x, Y0, 5.75, lab)
        rows = C[k]
        bar_chart(s, x, Y0 + 0.3, 5.75, 3.85, [r['n'] for r in rows], [r['v'] for r in rows], size=11)
        tb(s, x, Y0 + 4.2, 5.75, 0.4, [[(' · '.join(f"{r['n']} : {r['why']}" for r in rows if r['v'] < 0)[:150], {'size': 9, 'color': MUTE})]])
    wr, ws = R['rayons'][0], R['sites'][0]
    takeaway(s, X0, 6.2, W, 0.55, f"{wr['label']} ({E.fdm(wr['dM'])} M) et {ws['Site']} ({E.fdm(ws['dM'])} M) concentrent le recul. "
                                   'Les deux listes ne s’additionnent pas.')


def _annex_sector(prs, R, sct, code, ftr):
    s = _content(prs, 'Non-alim.' if sct == 'Non-alimentaire' else sct, R['texts']['sect_title'][sct], ftr, R)
    o = R['sect_detail'][sct]
    fam = R['familles'][sct]
    # 3 indicateurs
    K = [("CA", E.fm(o['CA']), 'M FCFA', o['ca'], E.fpct(o['ca'])), ('Marge', E.fm(o['M']), 'M FCFA', o['dM'], f"{E.fdm(o['dM'])} M"),
         ('Taux de marge', E.fnum(o['t'] * 100), '%', o['dt'], E.fpt(o['dt']))]
    for i, (lab, v, u, d, ds) in enumerate(K):
        x = X0 + i * 1.95
        tb(s, x, Y0, 1.9, 0.22, [[(lab.upper(), {'bold': True, 'size': 9, 'color': MUTE})]])
        tb(s, x, Y0 + 0.22, 1.9, 0.5, [[(v, {'bold': True, 'size': 22, 'color': NAVY}), (' ' + u, {'size': 9.5, 'color': MUTE})]])
        pill(s, x, Y0 + 0.78, ('▼ ' if d < 0 else '▲ ') + ds, d, size=10)
    hline(s, X0, Y0 + 1.25, 5.7)
    eyebrow(s, X0, Y0 + 1.42, 5.7, 'Écart de marge par format (M FCFA)')
    for i, f in enumerate(E.FORMAT_ORDER):
        v = fam['fmt'][f]
        small = abs(v) < E.SEUIL_SIGNIF
        x = X0 + i * 1.44
        rect(s, x, Y0 + 1.78, 1.36, 0.78, SOFT if small else (NEG_BG if v < 0 else POS_BG))
        tb(s, x, Y0 + 1.82, 1.36, 0.26, [[(f, {'size': 10, 'color': MUTE})]], align=PP_ALIGN.CENTER)
        tb(s, x, Y0 + 2.08, 1.36, 0.42, [[(E.fdm(v), {'size': 17, 'bold': True, 'color': sc(v, small)})]], align=PP_ALIGN.CENTER)
    tb(s, X0, Y0 + 2.8, 5.7, 1.4, [[(R['texts']['lecture'][sct], {'size': 12, 'color': INK})]], line=1.2)
    # familles
    xr = 6.85
    eyebrow(s, xr, Y0, 5.85, 'Familles : contributeurs et détracteurs (M FCFA)')
    sel = fam['top'][:3] + fam['flop'][:3][::-1]
    bar_chart(s, xr, Y0 + 0.28, 5.85, 3.9, [f_['label'] for f_ in sel], [f_['dM'] for f_ in sel], size=11)
    tb(s, xr, Y0 + 4.3, 5.85, 0.6, [[(' · '.join(f"{f_['label']} : volume {E.fdm(f_['effV'])}, taux {E.fdm(f_['effT'])}" for f_ in fam['flop'][:3]),
                                       {'size': 9, 'color': MUTE})]])


def _annex_lexique(prs, R, ftr):
    s = _content(prs, 'Annexe', 'lexique et méthode', ftr, R)
    lex = [('CA', 'ventes HT de la semaine, en M FCFA'), ('Marge', 'CA − coût d’achat ; marge N-1 = hors promo + promo'),
           ('Taux de marge', 'marge ÷ CA'), ('Trafic', 'nombre de tickets de caisse (débit)'), ('Panier moyen', 'CA ÷ tickets, en FCFA'),
           ('Vs N-1', 'même semaine de l’année précédente'), ('Δ (écart)', 'valeur N − valeur N-1'), ('pt (point)', 'écart entre deux taux'),
           ('Effet volume', '(CA N − CA N-1) × taux N-1 ; inclut l’effet prix'), ('Effet taux', 'CA N × (taux N − taux N-1)'),
           ('Non significatif', 'écart de marge inférieur à 0,5 M FCFA, affiché en gris'),
           ('Ex-Supeco', 'Supeco convertis en Market : ' + ', '.join(f"{n.replace('Market ', '')} ({d})" for n, d in E.CONVERSIONS)),
           ('Familles', 'classées par écart de marge en valeur, CA ≥ 1 M FCFA en N ou N-1'),
           ('Sources', 'exports Power BI A (famille × magasin), B (rayon × magasin), C (magasin) ; casse et budget hors périmètre')]
    half = (len(lex) + 1) // 2
    for col, chunk in enumerate([lex[:half], lex[half:]]):
        x = X0 + col * 6.2
        for i, (k, v) in enumerate(chunk):
            y = Y0 + i * 0.72
            tb(s, x, y, 5.8, 0.66, [[(k, {'bold': True, 'size': 12, 'color': NAVY})], [(v, {'size': 11})]])
            hline(s, x, y + 0.68, 5.8)


def build_report(R: dict, masque, out) -> None:
    prs = Presentation(masque)
    _drop_all_slides(prs)
    sem = f"S{R['meta']['week']}"
    ftr = f'Revue de performance · {sem}' + (' · données fictives' if R.get('demo') else '')
    _cover(prs, R)
    _p2(prs, R, ftr)
    _p3(prs, R, ftr)
    _p4(prs, R, ftr)
    for code, sct in zip(['A1', 'A2', 'A3'], E.SECTEUR_ORDER):
        _annex_sector(prs, R, sct, code, f'Annexe {code} · {ftr}')
    _annex_lexique(prs, R, f'Annexe A4 · {ftr}')
    prs.save(out)
