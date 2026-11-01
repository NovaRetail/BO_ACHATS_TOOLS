"""Moteur PGC Radar : lecture de l export PBI Département / Rayon / Site, agrégations, alertes et export Excel (repris tels quels)."""
import re
import io

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ============================================================
# Constantes métier
# ============================================================

FORMAT_MAP = {
    "10301": "Hyper", "10202": "Hyper", "10203": "Hyper",
    "10705": "Market", "10208": "Market", "10209": "Market",
    "10206": "Market", "10605": "Market", "10604": "Market",
    "10601": "Supeco", "10602": "Supeco", "10603": "Supeco",
}
SEUIL_CA_CRITIQUE = -0.05
SEUIL_CA_ATTENTION = -0.02
SEUIL_MARGE_CRITIQUE_PT = -1.0
SEUIL_MARGE_ATTENTION_PT = -0.5
MATERIALITE_FCFA = 1_000_000
SEUIL_POIDS_SITE = 0.10

CAUSES = ["— À sélectionner —", "Rupture stock fournisseur", "Rupture stock magasin",
          "Prix non compétitif", "Exécution magasin", "Action concurrente",
          "Effet promo / cannibalisation", "Saisonnalité", "Autre"]

BLEU, ROUGE, VERT, ORANGE, VIOLET = "#007AFF", "#FF3B30", "#34C759", "#FF9500", "#AF52DE"
GRIS, GRIS_CLAIR, NOIR = "#8E8E93", "#F2F2F7", "#1C1C1E"

NAVY_XL, ORANGE_XL = "24235C", "F09F39"
RED_F, RED_T = "FFC7CE", "9C0006"
AMB_F, AMB_T = "FFEB9C", "9C6500"
GRN_F, GRN_T = "C6EFCE", "006100"

# ============================================================
# Chargement et calculs
# ============================================================

@st.cache_data(show_spinner=False)
def load_export(file):
    df = pd.read_excel(file, sheet_name="Export", header=0)
    df = df[~df["Département"].astype(str).str.startswith("Filtres appliqués", na=False)]
    df = df.dropna(how="all")
    df["Département"] = df["Département"].ffill()
    df["Rayon"] = df["Rayon"].ffill()

    ca_reseau = df.loc[df["Département"] == "Total", "CA"].sum()
    df = df[df["Département"] != "Total"]

    dept = df[df["Rayon"] == "Total"].copy()
    site = df[(df["Site"] != "Total") & (df["Rayon"] != "Total") & df["Site"].notna()
              & (df["Département"] == "01 - PGC")].copy()
    site["code_site"] = site["Site"].str.split(" - ").str[0].str.strip()
    site["site_court"] = site["Site"].str.split(" - ").str[1].str.replace(
        r"^(Hyper|Market|Supeco)\s+", "", regex=True)
    site["format"] = site["code_site"].map(FORMAT_MAP).fillna("Autre")
    return dept, site, ca_reseau


def alerte(vs_n1, marge_pt):
    md = marge_pt if pd.notna(marge_pt) else 0
    if (pd.notna(vs_n1) and vs_n1 < SEUIL_CA_CRITIQUE) or md < SEUIL_MARGE_CRITIQUE_PT:
        return "Critique"
    if (pd.notna(vs_n1) and vs_n1 < SEUIL_CA_ATTENTION) or md < SEUIL_MARGE_ATTENTION_PT:
        return "Attention"
    return "OK"


BADGE = {"Critique": ("badge-crit", "Critique"), "Attention": ("badge-att", "Attention"),
         "OK": ("badge-ok", "OK")}


def fmt_m(x):
    return "—" if pd.isna(x) else f"{x/1_000_000:,.0f} M".replace(",", " ")

def fmt_pct(x, d=1):
    return "—" if pd.isna(x) else f"{x*100:+.{d}f}%".replace(".", ",")

def fmt_pt(x):
    return "—" if pd.isna(x) else f"{x:+.2f} pt".replace(".", ",")

def cls(x):
    if pd.isna(x):
        return "mut"
    return "neg" if x < 0 else "pos"


def agg_niveau(site_df, by):
    g = site_df.groupby(by).agg(
        nb_sites=("code_site", "nunique"),
        CA=("CA", "sum"), CA_N1=("CA N-1", "sum"),
        Budget=("Budget", lambda s: np.nan if s.isna().all() else s.sum()),
        Marge=("Marge", "sum"), Marge_N1=("Marge N-1", "sum"),
    ).reset_index()
    g["ecart_n1"] = g["CA"] - g["CA_N1"]
    g["Vs_N1"] = g["ecart_n1"] / g["CA_N1"]
    g["Vs_Bgt"] = np.where(g["Budget"].isna(), np.nan, (g["CA"] - g["Budget"]) / g["Budget"])
    g["Marge_pt"] = (g["Marge"] / g["CA"] - g["Marge_N1"] / g["CA_N1"]) * 100
    pool = g.loc[g["ecart_n1"] < 0, "ecart_n1"].sum()
    g["poids"] = g["ecart_n1"].apply(lambda x: x / pool if (x < 0 and pool < 0) else 0.0)
    return g


# ============================================================
# Export Excel — 100% calculé en Python, aucune formule
# ============================================================

def build_excel(pgc, dept, ca_reseau, fmt_lvl, rayon_lvl, site_lvl, al_actives):
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor=NAVY_XL)
    head_font = Font(name="Arial", color="FFFFFF", bold=True, size=10)
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title_font = Font(name="Arial", size=12, bold=True, color=NAVY_XL)
    center = Alignment(horizontal="center")

    def header(ws, row, headers):
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=i, value=h)
            c.fill, c.font, c.alignment = head_fill, head_font, center

    def paint(cell, level):
        f, t = {"Critique": (RED_F, RED_T), "Attention": (AMB_F, AMB_T),
                "OK": (GRN_F, GRN_T)}[level]
        cell.fill = PatternFill("solid", fgColor=f)
        cell.font = Font(name="Arial", size=10, bold=True, color=t)

    # ---- Synthèse ----
    ws = wb.active
    ws.title = "Synthèse"
    ws.sheet_properties.tabColor = ORANGE_XL
    ws.cell(row=1, column=1, value="PGC Radar — Synthèse hebdomadaire").font = title_font
    header(ws, 3, ["Indicateur", "Valeur", "Vs N-1", "Vs Budget"])
    rows = [
        ("CA PGC (M)", round(pgc["CA"] / 1e6, 1), pgc["Vs N-1 (%)"], pgc["Vs Bgt (%)"]),
        ("Marge PGC (M)", round(pgc["Marge"] / 1e6, 1), None, None),
        ("Taux de marge", pgc["Taux de Marge"], pgc["Taux de Marge N Vs N-1"], None),
    ]
    r = 4
    for label, val, v1, vb in rows:
        ws.cell(row=r, column=1, value=label).font = base
        c = ws.cell(row=r, column=2, value=val)
        c.font = bold
        if label == "Taux de marge":
            c.number_format = "0.0%"
            c2 = ws.cell(row=r, column=3, value=round(v1, 2) if pd.notna(v1) else None)
            c2.font = base
            c2.number_format = '0.00" pt"'
        else:
            if v1 is not None:
                c2 = ws.cell(row=r, column=3, value=round(v1, 4))
                c2.number_format = "+0.0%;-0.0%"
                c2.font = base
            if vb is not None:
                c3 = ws.cell(row=r, column=4, value=round(vb, 4))
                c3.number_format = "+0.0%;-0.0%"
                c3.font = base
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Benchmark départements").font = title_font
    r += 1
    header(ws, r, ["Département", "CA (M)", "Vs N-1", "Vs Budget"])
    r += 1
    for _, d in dept.iterrows():
        ws.cell(row=r, column=1, value=d["Département"]).font = \
            bold if d["Département"] == "01 - PGC" else base
        ws.cell(row=r, column=2, value=round(d["CA"] / 1e6, 1)).font = base
        c = ws.cell(row=r, column=3, value=round(d["Vs N-1 (%)"], 4))
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        c = ws.cell(row=r, column=4, value=round(d["Vs Bgt (%)"], 4)
                    if pd.notna(d["Vs Bgt (%)"]) else None)
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        r += 1
    ws.cell(row=r, column=1, value="Total magasin").font = bold
    ws.cell(row=r, column=2, value=round(ca_reseau / 1e6, 1)).font = bold
    for col, w in zip(range(1, 5), [28, 12, 12, 12]):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ---- Format ----
    ws = wb.create_sheet("Format")
    ws.sheet_properties.tabColor = ORANGE_XL
    ws.cell(row=1, column=1, value="Performance par format et par site").font = title_font
    header(ws, 3, ["Format / Site", "Nb sites", "CA (M)", "Vs N-1", "Vs Budget",
                   "Marge (pt vs N-1)", "Poids ds le recul", "Alerte"])
    r = 4
    for _, f in fmt_lvl.iterrows():
        lvl = alerte(f["Vs_N1"], f["Marge_pt"])
        ws.cell(row=r, column=1, value=f["format"]).font = bold
        ws.cell(row=r, column=2, value=int(f["nb_sites"])).font = base
        ws.cell(row=r, column=3, value=round(f["CA"] / 1e6, 1)).font = bold
        c = ws.cell(row=r, column=4, value=round(f["Vs_N1"], 4))
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        c = ws.cell(row=r, column=5, value=round(f["Vs_Bgt"], 4) if pd.notna(f["Vs_Bgt"]) else None)
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        c = ws.cell(row=r, column=6, value=round(f["Marge_pt"], 2))
        c.font = base
        paint(ws.cell(row=r, column=8, value=lvl), lvl)
        r += 1
        sites_f = site_lvl[site_lvl["format"] == f["format"]].sort_values("ecart_n1")
        for _, s_ in sites_f.iterrows():
            ws.cell(row=r, column=1, value="    " + s_["site_court"]).font = base
            ws.cell(row=r, column=3, value=round(s_["CA"] / 1e6, 1)).font = base
            c = ws.cell(row=r, column=4, value=round(s_["Vs_N1"], 4))
            c.number_format = "+0.0%;-0.0%"
            c.font = base
            c = ws.cell(row=r, column=5, value=round(s_["Vs_Bgt"], 4)
                        if pd.notna(s_["Vs_Bgt"]) else None)
            c.number_format = "+0.0%;-0.0%"
            c.font = base
            c = ws.cell(row=r, column=7, value=round(s_["poids"], 3) if s_["poids"] > 0 else None)
            c.number_format = "0%"
            c.font = base
            r += 1
    for col, w in zip(range(1, 9), [26, 9, 10, 10, 10, 15, 15, 12]):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ---- Rayon ----
    ws = wb.create_sheet("Rayon")
    ws.sheet_properties.tabColor = ORANGE_XL
    ws.cell(row=1, column=1,
            value="Performance par rayon (tous formats et sites confondus)").font = title_font
    header(ws, 3, ["Rayon", "Nb sites", "CA (M)", "Vs N-1", "Vs Budget",
                   "Marge (pt vs N-1)", "Poids ds le recul", "Alerte"])
    r = 4
    for _, ry in rayon_lvl.iterrows():
        lvl = alerte(ry["Vs_N1"], ry["Marge_pt"])
        nom = ry["Rayon"].split(" - ")[1].title() if " - " in ry["Rayon"] else ry["Rayon"]
        ws.cell(row=r, column=1, value=nom).font = bold
        ws.cell(row=r, column=2, value=int(ry["nb_sites"])).font = base
        ws.cell(row=r, column=3, value=round(ry["CA"] / 1e6, 1)).font = bold
        c = ws.cell(row=r, column=4, value=round(ry["Vs_N1"], 4))
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        c = ws.cell(row=r, column=5, value=round(ry["Vs_Bgt"], 4) if pd.notna(ry["Vs_Bgt"]) else None)
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        c = ws.cell(row=r, column=6, value=round(ry["Marge_pt"], 2))
        c.font = base
        c = ws.cell(row=r, column=7, value=round(ry["poids"], 3) if ry["poids"] > 0 else None)
        c.number_format = "0%"
        c.font = base
        paint(ws.cell(row=r, column=8, value=lvl), lvl)
        r += 1
    for col, w in zip(range(1, 9), [26, 9, 10, 10, 10, 15, 15, 12]):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ---- Alertes ----
    ws = wb.create_sheet("Alertes")
    ws.sheet_properties.tabColor = ORANGE_XL
    ws.cell(row=1, column=1,
            value=f"Alertes à traiter — déclencheur Vs N-1 + marge, écart ≥ "
                  f"{MATERIALITE_FCFA/1e6:.0f} M FCFA").font = title_font
    header(ws, 3, ["Site", "Rayon", "Alerte", "CA (M)", "Écart N-1 (M)", "Vs N-1",
                   "Vs Budget", "Marge (pt)", "Débit vs N-1", "Panier vs N-1",
                   "Poids ds le recul", "Cause", "Commentaire"])
    r = 4
    first_data_row = r
    for _, a in al_actives.iterrows():
        ws.cell(row=r, column=1, value=a["Site"].split(" - ")[1]).font = base
        ws.cell(row=r, column=2, value=a["Rayon"].split(" - ")[1].title()).font = base
        paint(ws.cell(row=r, column=3, value=a["Alerte"]), a["Alerte"])
        ws.cell(row=r, column=4, value=round(a["CA"] / 1e6, 1)).font = base
        ws.cell(row=r, column=5, value=round(a["ecart_n1"] / 1e6, 1)).font = bold
        for col, key in [(6, "Vs_N1"), (7, "Vs_Bgt"), (9, "Debit_vs_n1"), (10, "Panier_vs_n1")]:
            c = ws.cell(row=r, column=col,
                        value=round(a[key], 4) if pd.notna(a[key]) else None)
            c.number_format = "+0.0%;-0.0%"
            c.font = base
        c = ws.cell(row=r, column=8, value=round(a["Marge_pt"], 2) if pd.notna(a["Marge_pt"]) else None)
        c.font = base
        c = ws.cell(row=r, column=11, value=round(a["poids"], 3) if a["poids"] > 0 else None)
        c.number_format = "0%"
        c.font = base
        ws.cell(row=r, column=12, value="").font = base
        r += 1
    if r > first_data_row:
        dv = DataValidation(type="list",
                            formula1='"' + ",".join(CAUSES[1:]) + '"',
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"L{first_data_row}:L{r-1}")
    for col, w in zip(range(1, 14), [16, 20, 11, 9, 12, 10, 10, 10, 12, 12, 14, 26, 30]):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ---- Règles ----
    ws = wb.create_sheet("Règles")
    ws.sheet_properties.tabColor = NAVY_XL
    ws.cell(row=1, column=1, value="Règles de calcul et d'alerte").font = title_font
    regles = [
        ("Déclencheur d'alerte", "Vs N-1 du CA et évolution du taux de marge (points) — "
                                 "le Budget n'est jamais un déclencheur (référence seulement)"),
        ("CA — Critique", f"Vs N-1 < {SEUIL_CA_CRITIQUE:.0%}"),
        ("CA — Attention", f"Vs N-1 < {SEUIL_CA_ATTENTION:.0%}"),
        ("Marge — Critique", f"évolution < {SEUIL_MARGE_CRITIQUE_PT} pt vs N-1"),
        ("Marge — Attention", f"évolution < {SEUIL_MARGE_ATTENTION_PT} pt vs N-1"),
        ("Plancher de matérialité", f"écart en valeur ≥ {MATERIALITE_FCFA:,.0f} FCFA "
                                    "à la maille Rayon × Site".replace(",", " ")),
        ("Poids dans le recul", "part de la ligne dans la somme des écarts négatifs vs N-1, "
                                "calculée au même niveau d'agrégation"),
        ("Rayon Global", "vue agrégée tous formats et sites confondus, même logique "
                         "de poids et d'alerte que le niveau Format"),
        ("Budget Supeco", "absent de l'export source — affiché « — », exclu des calculs"),
        ("Débit / Panier", "fiables uniquement au niveau Rayon × Site ; jamais sommés "
                           "entre rayons ; pas de budget disponible pour ces indicateurs"),
        ("Montants", "tous arrondis au million (M FCFA)"),
    ]
    r = 3
    for k, v in regles:
        ws.cell(row=r, column=1, value=k).font = bold
        ws.cell(row=r, column=2, value=v).font = base
        r += 1
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============================================================
# Sidebar


def compute(up):
    """Charge l'export et renvoie tous les niveaux d'analyse."""
    dept, site, ca_reseau = load_export(up)
    pgc = dept[dept["Département"] == "01 - PGC"].iloc[0]
    fmt_lvl = agg_niveau(site, "format").set_index("format").reindex(["Hyper", "Market", "Supeco"]).reset_index()
    site_lvl = agg_niveau(site, ["code_site", "site_court", "format"])
    rayon_lvl = agg_niveau(site, "Rayon").sort_values("ecart_n1")
    al = site.copy()
    al["Vs_N1"] = al["Vs N-1 (%)"]
    al["Vs_Bgt"] = al["Vs Bgt (%)"]
    al["Marge_pt"] = al["Taux de Marge N Vs N-1"]
    al["Debit_vs_n1"] = al["Vs N-1 (%).1"]
    al["Panier_vs_n1"] = al["Panier N Vs N-1"]
    al["ecart_n1"] = al["CA"] - al["CA N-1"]
    al["Alerte"] = al.apply(lambda r: alerte(r["Vs_N1"], r["Marge_pt"]), axis=1)
    pool_al = al.loc[al["ecart_n1"] < 0, "ecart_n1"].sum()
    al["poids"] = al["ecart_n1"].apply(lambda x: x / pool_al if (x < 0 and pool_al < 0) else 0.0)
    al_actives = al[(al["Alerte"] != "OK") & (al["ecart_n1"].abs() >= MATERIALITE_FCFA)].sort_values("ecart_n1")
    return dict(dept=dept, site=site, ca_reseau=ca_reseau, pgc=pgc, fmt_lvl=fmt_lvl, site_lvl=site_lvl,
                rayon_lvl=rayon_lvl, al_actives=al_actives)
