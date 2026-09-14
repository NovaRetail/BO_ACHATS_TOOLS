import re
import io

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

st.set_page_config(page_title="PGC Radar", page_icon="📡", layout="wide")

def render_html(fragment):
    """Compacte le HTML en une ligne — évite l'interprétation markdown en bloc de code."""
    st.markdown(re.sub(r"\n\s*", " ", fragment), unsafe_allow_html=True)


# ============================================================
# Constantes métier
# ============================================================

FORMAT_MAP = {
    "10301": "Hyper", "10202": "Hyper", "10203": "Hyper",
    "10705": "Market", "10208": "Market", "10209": "Market",
    "10206": "Market", "10605": "Market", "10604": "Market",
    "10601": "Supeco", "10602": "Supeco", "10603": "Supeco",
}
SEUIL_CA_CRITIQUE = -0.05
SEUIL_CA_ATTENTION = -0.02
SEUIL_MARGE_CRITIQUE_PT = -1.0
SEUIL_MARGE_ATTENTION_PT = -0.5
MATERIALITE_FCFA = 1_000_000
SEUIL_POIDS_SITE = 0.10

CAUSES = ["— À sélectionner —", "Rupture stock fournisseur", "Rupture stock magasin",
          "Prix non compétitif", "Exécution magasin", "Action concurrente",
          "Effet promo / cannibalisation", "Saisonnalité", "Autre"]

BLEU, ROUGE, VERT, ORANGE, VIOLET = "#007AFF", "#FF3B30", "#34C759", "#FF9500", "#AF52DE"
GRIS, GRIS_CLAIR, NOIR = "#8E8E93", "#F2F2F7", "#1C1C1E"

NAVY_XL, ORANGE_XL = "24235C", "F09F39"
RED_F, RED_T = "FFC7CE", "9C0006"
AMB_F, AMB_T = "FFEB9C", "9C6500"
GRN_F, GRN_T = "C6EFCE", "006100"

# ============================================================
# Charte visuelle — Apple style, police arrondie
# ============================================================

render_html(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Nunito:wght@400;600;700;800;900&display=swap');

html, body, [class*="css"], .stApp {{
    font-family: 'Nunito', -apple-system, BlinkMacSystemFont, sans-serif;
}}
.stApp {{ background-color: #F5F6FA; }}

div[data-testid="stTabs"] [data-baseweb="tab-list"],
.stTabs [data-baseweb="tab-list"] {{
    background: #E9EAF0; border-radius: 12px; padding: 3px; gap: 0;
    border-bottom: none !important;
}}
div[data-testid="stTabs"] button[data-baseweb="tab"],
.stTabs [data-baseweb="tab"] {{
    flex: 1; justify-content: center; border-radius: 10px;
    font-weight: 700; color: {GRIS}; background: transparent; height: 36px;
    border: none !important;
}}
div[data-testid="stTabs"] button[data-baseweb="tab"] p,
.stTabs [data-baseweb="tab"] p {{
    font-weight: 700; color: {GRIS};
}}
div[data-testid="stTabs"] button[aria-selected="true"],
.stTabs [aria-selected="true"] {{
    background: #fff !important; color: {NOIR} !important;
}}
div[data-testid="stTabs"] button[aria-selected="true"] p,
.stTabs [aria-selected="true"] p {{
    color: {NOIR} !important;
}}
div[data-testid="stTabs"] [data-baseweb="tab-highlight"],
div[data-testid="stTabs"] [data-baseweb="tab-border"],
.stTabs [data-baseweb="tab-highlight"], .stTabs [data-baseweb="tab-border"] {{ display: none; }}

.stDownloadButton button {{
    background: #0A2540 !important; color: #fff !important;
    border-radius: 12px !important; font-weight: 700 !important; border: none !important;
}}

.hero {{
    background: linear-gradient(135deg, #0A2540 0%, #1D4E89 55%, #6A5ACD 100%);
    border-radius: 18px; padding: 18px 20px; margin-bottom: 14px;
}}
.hero-title {{ font-size: 19px; font-weight: 800; color: #fff; }}
.hero-sub {{ font-size: 11px; color: rgba(255,255,255,0.75); font-weight: 600; }}
.hero-chips {{ display: flex; gap: 8px; margin-top: 10px; flex-wrap: wrap; }}
.chip {{ background: rgba(255,255,255,0.14); color: #fff; font-size: 12px;
         font-weight: 700; padding: 4px 12px; border-radius: 11px; }}
.chip-alert {{ background: rgba(255,90,80,0.4); }}

.page-title {{ font-size: 26px; font-weight: 900; color: {NOIR}; padding: 2px 0; }}
.page-sub {{ font-size: 13px; color: {GRIS}; font-weight: 600; margin-bottom: 16px; }}
.section-title {{ font-size: 15px; font-weight: 800; color: {NOIR}; margin: 6px 0 10px 0; }}

.kpi-card {{
    border-radius: 18px; padding: 16px 18px;
    display: flex; flex-direction: column; gap: 7px; height: 100%;
}}
.kpi-blue   {{ background: #EAF2FF; }}
.kpi-violet {{ background: #F3EDFC; }}
.kpi-top {{ display: flex; justify-content: space-between; align-items: center; }}
.kpi-label {{ font-size: 13px; font-weight: 800; }}
.lab-blue   {{ color: #4A6FA5; }}
.lab-violet {{ color: #7A5FA8; }}
.kpi-ico {{
    width: 34px; height: 34px; border-radius: 11px;
    display: flex; align-items: center; justify-content: center; font-size: 16px;
}}
.ico-blue   {{ background: {BLEU}; }}
.ico-violet {{ background: {VIOLET}; }}
.kpi-value {{ font-size: 32px; font-weight: 900; line-height: 1.05; }}
.val-blue   {{ color: #0C2D5B; }}
.val-violet {{ color: #3E2467; }}
.pill {{
    display: inline-block; font-size: 12px; font-weight: 800;
    padding: 3px 12px; border-radius: 12px; width: fit-content;
}}
.pill-neg {{ background: #FFE0DD; color: #C82D24; }}
.pill-pos {{ background: #E4F8EA; color: #1E9E47; }}
.pill-mut {{ background: {GRIS_CLAIR}; color: {GRIS}; }}

.ref-line {{
    background: #fff; border-radius: 14px; padding: 10px 16px;
    display: flex; justify-content: space-between; align-items: center;
    font-size: 13px; color: {GRIS}; font-weight: 700; margin: 12px 0 18px 0;
    border: 1px dashed #E0E0E5;
}}

.bench-card {{ background: #fff; border-radius: 16px; padding: 14px 16px; margin-bottom: 18px;
               border: 1px solid rgba(0,0,0,0.03); }}
.bench-item {{ margin-bottom: 10px; }}
.bench-item:last-child {{ margin-bottom: 0; }}
.bench-line {{ display: flex; justify-content: space-between; font-size: 13px;
               font-weight: 700; color: {NOIR}; margin-bottom: 4px; }}
.bench-track {{ height: 6px; background: {GRIS_CLAIR}; border-radius: 3px; }}
.bench-fill {{ height: 100%; background: {BLEU}; border-radius: 3px; }}

.fmt-block {{ background: #fff; border-radius: 16px; padding: 14px 16px; margin-bottom: 14px;
              border: 1px solid rgba(0,0,0,0.03); }}
.fmt-head {{ display: flex; justify-content: space-between; align-items: center; margin-bottom: 6px; }}
.fmt-name {{ font-size: 15px; font-weight: 900; color: {NOIR}; }}
.fmt-sub {{ font-size: 12px; color: {GRIS}; font-weight: 700; }}
.contrib-bar {{ display: flex; height: 10px; border-radius: 5px; overflow: hidden;
                gap: 2px; margin: 8px 0 6px 0; }}
.site-row {{
    display: flex; justify-content: space-between; padding: 7px 0 7px 10px;
    font-size: 13px; color: {NOIR}; font-weight: 700;
    border-top: 0.5px solid #F0F0F3;
}}
.autres {{ font-size: 12px; color: {GRIS}; font-weight: 700; padding: 7px 0 0 10px; }}

.badge {{ font-size: 11px; font-weight: 900; padding: 3px 12px; border-radius: 12px; }}
.badge-crit {{ background: #FFE5E3; color: {ROUGE}; }}
.badge-att  {{ background: #FFF2DE; color: {ORANGE}; }}
.badge-ok   {{ background: #E4F8EA; color: #1E9E47; }}

.alert-card {{ background: #fff; padding: 14px 16px; margin-bottom: 12px;
               border: 1px solid rgba(0,0,0,0.03); border-left: 4px solid transparent;
               border-radius: 0 16px 16px 0; }}
.alert-crit {{ border-left-color: {ROUGE}; }}
.alert-att  {{ border-left-color: {ORANGE}; }}
.alert-flex {{ display: flex; align-items: center; gap: 10px; margin-bottom: 6px; }}
.avatar {{ width: 34px; height: 34px; border-radius: 50%; font-size: 13px; font-weight: 800;
           display: flex; align-items: center; justify-content: center; flex-shrink: 0; }}
.av-crit {{ background: #FFE5E3; color: {ROUGE}; }}
.av-att  {{ background: #FFF2DE; color: {ORANGE}; }}
.alert-head {{ display: flex; justify-content: space-between; align-items: center; }}
.alert-title {{ font-size: 14px; font-weight: 900; color: {NOIR}; }}
.alert-poids {{ font-size: 11px; color: {GRIS}; font-weight: 700; }}
.alert-line {{ font-size: 12px; color: {GRIS}; font-weight: 700; margin-top: 3px; }}

.landing-hero {{
    background: linear-gradient(135deg, #0A2540 0%, #1D4E89 55%, #6A5ACD 100%);
    border-radius: 18px; padding: 20px 22px; margin-bottom: 16px;
}}
.landing-card {{ background: #fff; border-radius: 16px; padding: 16px 18px; height: 100%;
                 border: 1px solid rgba(0,0,0,0.03); }}
.rule-row {{ display: flex; justify-content: space-between; padding: 7px 0;
             font-size: 13px; font-weight: 700; color: {NOIR};
             border-bottom: 0.5px solid #F0F0F3; }}
.rule-row:last-child {{ border-bottom: none; }}
.col-chip {{ display: inline-block; background: {GRIS_CLAIR}; color: {NOIR};
             font-size: 11px; font-weight: 700; padding: 3px 10px;
             border-radius: 10px; margin: 3px 3px 0 0; }}

.neg {{ color: {ROUGE}; font-weight: 800; }}
.pos {{ color: #1E9E47; font-weight: 800; }}
.mut {{ color: {GRIS}; font-weight: 700; }}
</style>
""")

# ============================================================
# Chargement et calculs
# ============================================================

@st.cache_data(show_spinner=False)
def load_export(file):
    df = pd.read_excel(file, sheet_name="Export", header=0)
    df = df[~df["Département"].astype(str).str.startswith("Filtres appliqués", na=False)]
    df = df.dropna(how="all")
    df["Département"] = df["Département"].ffill()
    df["Rayon"] = df["Rayon"].ffill()

    ca_reseau = df.loc[df["Département"] == "Total", "CA"].sum()
    df = df[df["Département"] != "Total"]

    dept = df[df["Rayon"] == "Total"].copy()
    site = df[(df["Site"] != "Total") & (df["Rayon"] != "Total") & df["Site"].notna()
              & (df["Département"] == "01 - PGC")].copy()
    site["code_site"] = site["Site"].str.split(" - ").str[0].str.strip()
    site["site_court"] = site["Site"].str.split(" - ").str[1].str.replace(
        r"^(Hyper|Market|Supeco)\s+", "", regex=True)
    site["format"] = site["code_site"].map(FORMAT_MAP).fillna("Autre")
    return dept, site, ca_reseau


def alerte(vs_n1, marge_pt):
    md = marge_pt if pd.notna(marge_pt) else 0
    if (pd.notna(vs_n1) and vs_n1 < SEUIL_CA_CRITIQUE) or md < SEUIL_MARGE_CRITIQUE_PT:
        return "Critique"
    if (pd.notna(vs_n1) and vs_n1 < SEUIL_CA_ATTENTION) or md < SEUIL_MARGE_ATTENTION_PT:
        return "Attention"
    return "OK"


BADGE = {"Critique": ("badge-crit", "Critique"), "Attention": ("badge-att", "Attention"),
         "OK": ("badge-ok", "OK")}


def fmt_m(x):
    return "—" if pd.isna(x) else f"{x/1_000_000:,.0f} M".replace(",", " ")

def fmt_pct(x, d=1):
    return "—" if pd.isna(x) else f"{x*100:+.{d}f}%".replace(".", ",")

def fmt_pt(x):
    return "—" if pd.isna(x) else f"{x:+.2f} pt".replace(".", ",")

def cls(x):
    if pd.isna(x):
        return "mut"
    return "neg" if x < 0 else "pos"


def agg_niveau(site_df, by):
    g = site_df.groupby(by).agg(
        nb_sites=("code_site", "nunique"),
        CA=("CA", "sum"), CA_N1=("CA N-1", "sum"),
        Budget=("Budget", lambda s: np.nan if s.isna().all() else s.sum()),
        Marge=("Marge", "sum"), Marge_N1=("Marge N-1", "sum"),
    ).reset_index()
    g["ecart_n1"] = g["CA"] - g["CA_N1"]
    g["Vs_N1"] = g["ecart_n1"] / g["CA_N1"]
    g["Vs_Bgt"] = np.where(g["Budget"].isna(), np.nan, (g["CA"] - g["Budget"]) / g["Budget"])
    g["Marge_pt"] = (g["Marge"] / g["CA"] - g["Marge_N1"] / g["CA_N1"]) * 100
    pool = g.loc[g["ecart_n1"] < 0, "ecart_n1"].sum()
    g["poids"] = g["ecart_n1"].apply(lambda x: x / pool if (x < 0 and pool < 0) else 0.0)
    return g


# ============================================================
# Export Excel — 100% calculé en Python, aucune formule
# ============================================================

def build_excel(pgc, dept, ca_reseau, fmt_lvl, site_lvl, al_actives):
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor=NAVY_XL)
    head_font = Font(name="Arial", color="FFFFFF", bold=True, size=10)
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title_font = Font(name="Arial", size=12, bold=True, color=NAVY_XL)
    center = Alignment(horizontal="center")

    def header(ws, row, headers):
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=i, value=h)
            c.fill, c.font, c.alignment = head_fill, head_font, center

    def paint(cell, level):
        f, t = {"Critique": (RED_F, RED_T), "Attention": (AMB_F, AMB_T),
                "OK": (GRN_F, GRN_T)}[level]
        cell.fill = PatternFill("solid", fgColor=f)
        cell.font = Font(name="Arial", size=10, bold=True, color=t)

    # ---- Synthèse ----
    ws = wb.active
    ws.title = "Synthèse"
    ws.sheet_properties.tabColor = ORANGE_XL
    ws.cell(row=1, column=1, value="PGC Radar — Synthèse hebdomadaire").font = title_font
    header(ws, 3, ["Indicateur", "Valeur", "Vs N-1", "Vs Budget"])
    rows = [
        ("CA PGC (M)", round(pgc["CA"] / 1e6, 1), pgc["Vs N-1 (%)"], pgc["Vs Bgt (%)"]),
        ("Marge PGC (M)", round(pgc["Marge"] / 1e6, 1), None, None),
        ("Taux de marge", pgc["Taux de Marge"], pgc["Taux de Marge N Vs N-1"], None),
    ]
    r = 4
    for label, val, v1, vb in rows:
        ws.cell(row=r, column=1, value=label).font = base
        c = ws.cell(row=r, column=2, value=val)
        c.font = bold
        if label == "Taux de marge":
            c.number_format = "0.0%"
            c2 = ws.cell(row=r, column=3, value=round(v1, 2) if pd.notna(v1) else None)
            c2.font = base
            c2.number_format = '0.00" pt"'
        else:
            if v1 is not None:
                c2 = ws.cell(row=r, column=3, value=round(v1, 4))
                c2.number_format = "+0.0%;-0.0%"
                c2.font = base
            if vb is not None:
                c3 = ws.cell(row=r, column=4, value=round(vb, 4))
                c3.number_format = "+0.0%;-0.0%"
                c3.font = base
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Benchmark départements").font = title_font
    r += 1
    header(ws, r, ["Département", "CA (M)", "Vs N-1", "Vs Budget"])
    r += 1
    for _, d in dept.iterrows():
        ws.cell(row=r, column=1, value=d["Département"]).font = \
            bold if d["Département"] == "01 - PGC" else base
        ws.cell(row=r, column=2, value=round(d["CA"] / 1e6, 1)).font = base
        c = ws.cell(row=r, column=3, value=round(d["Vs N-1 (%)"], 4))
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        c = ws.cell(row=r, column=4, value=round(d["Vs Bgt (%)"], 4)
                    if pd.notna(d["Vs Bgt (%)"]) else None)
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        r += 1
    ws.cell(row=r, column=1, value="Total magasin").font = bold
    ws.cell(row=r, column=2, value=round(ca_reseau / 1e6, 1)).font = bold
    for col, w in zip(range(1, 5), [28, 12, 12, 12]):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ---- Format ----
    ws = wb.create_sheet("Format")
    ws.sheet_properties.tabColor = ORANGE_XL
    ws.cell(row=1, column=1, value="Performance par format et par site").font = title_font
    header(ws, 3, ["Format / Site", "Nb sites", "CA (M)", "Vs N-1", "Vs Budget",
                   "Marge (pt vs N-1)", "Poids ds le recul", "Alerte"])
    r = 4
    for _, f in fmt_lvl.iterrows():
        lvl = alerte(f["Vs_N1"], f["Marge_pt"])
        ws.cell(row=r, column=1, value=f["format"]).font = bold
        ws.cell(row=r, column=2, value=int(f["nb_sites"])).font = base
        ws.cell(row=r, column=3, value=round(f["CA"] / 1e6, 1)).font = bold
        c = ws.cell(row=r, column=4, value=round(f["Vs_N1"], 4))
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        c = ws.cell(row=r, column=5, value=round(f["Vs_Bgt"], 4) if pd.notna(f["Vs_Bgt"]) else None)
        c.number_format = "+0.0%;-0.0%"
        c.font = base
        c = ws.cell(row=r, column=6, value=round(f["Marge_pt"], 2))
        c.font = base
        paint(ws.cell(row=r, column=8, value=lvl), lvl)
        r += 1
        sites_f = site_lvl[site_lvl["format"] == f["format"]].sort_values("ecart_n1")
        for _, s_ in sites_f.iterrows():
            ws.cell(row=r, column=1, value="    " + s_["site_court"]).font = base
            ws.cell(row=r, column=3, value=round(s_["CA"] / 1e6, 1)).font = base
            c = ws.cell(row=r, column=4, value=round(s_["Vs_N1"], 4))
            c.number_format = "+0.0%;-0.0%"
            c.font = base
            c = ws.cell(row=r, column=5, value=round(s_["Vs_Bgt"], 4)
                        if pd.notna(s_["Vs_Bgt"]) else None)
            c.number_format = "+0.0%;-0.0%"
            c.font = base
            c = ws.cell(row=r, column=7, value=round(s_["poids"], 3) if s_["poids"] > 0 else None)
            c.number_format = "0%"
            c.font = base
            r += 1
    for col, w in zip(range(1, 9), [26, 9, 10, 10, 10, 15, 15, 12]):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ---- Alertes ----
    ws = wb.create_sheet("Alertes")
    ws.sheet_properties.tabColor = ORANGE_XL
    ws.cell(row=1, column=1,
            value=f"Alertes à traiter — déclencheur Vs N-1 + marge, écart ≥ "
                  f"{MATERIALITE_FCFA/1e6:.0f} M FCFA").font = title_font
    header(ws, 3, ["Site", "Rayon", "Alerte", "CA (M)", "Écart N-1 (M)", "Vs N-1",
                   "Vs Budget", "Marge (pt)", "Débit vs N-1", "Panier vs N-1",
                   "Poids ds le recul", "Cause", "Commentaire"])
    r = 4
    first_data_row = r
    for _, a in al_actives.iterrows():
        ws.cell(row=r, column=1, value=a["Site"].split(" - ")[1]).font = base
        ws.cell(row=r, column=2, value=a["Rayon"].split(" - ")[1].title()).font = base
        paint(ws.cell(row=r, column=3, value=a["Alerte"]), a["Alerte"])
        ws.cell(row=r, column=4, value=round(a["CA"] / 1e6, 1)).font = base
        ws.cell(row=r, column=5, value=round(a["ecart_n1"] / 1e6, 1)).font = bold
        for col, key in [(6, "Vs_N1"), (7, "Vs_Bgt"), (9, "Debit_vs_n1"), (10, "Panier_vs_n1")]:
            c = ws.cell(row=r, column=col,
                        value=round(a[key], 4) if pd.notna(a[key]) else None)
            c.number_format = "+0.0%;-0.0%"
            c.font = base
        c = ws.cell(row=r, column=8, value=round(a["Marge_pt"], 2) if pd.notna(a["Marge_pt"]) else None)
        c.font = base
        c = ws.cell(row=r, column=11, value=round(a["poids"], 3) if a["poids"] > 0 else None)
        c.number_format = "0%"
        c.font = base
        ws.cell(row=r, column=12, value="").font = base
        r += 1
    if r > first_data_row:
        dv = DataValidation(type="list",
                            formula1='"' + ",".join(CAUSES[1:]) + '"',
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"L{first_data_row}:L{r-1}")
    for col, w in zip(range(1, 14), [16, 20, 11, 9, 12, 10, 10, 10, 12, 12, 14, 26, 30]):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ---- Règles ----
    ws = wb.create_sheet("Règles")
    ws.sheet_properties.tabColor = NAVY_XL
    ws.cell(row=1, column=1, value="Règles de calcul et d'alerte").font = title_font
    regles = [
        ("Déclencheur d'alerte", "Vs N-1 du CA et évolution du taux de marge (points) — "
                                 "le Budget n'est jamais un déclencheur (référence seulement)"),
        ("CA — Critique", f"Vs N-1 < {SEUIL_CA_CRITIQUE:.0%}"),
        ("CA — Attention", f"Vs N-1 < {SEUIL_CA_ATTENTION:.0%}"),
        ("Marge — Critique", f"évolution < {SEUIL_MARGE_CRITIQUE_PT} pt vs N-1"),
        ("Marge — Attention", f"évolution < {SEUIL_MARGE_ATTENTION_PT} pt vs N-1"),
        ("Plancher de matérialité", f"écart en valeur ≥ {MATERIALITE_FCFA:,.0f} FCFA "
                                    "à la maille Rayon × Site".replace(",", " ")),
        ("Poids dans le recul", "part de la ligne dans la somme des écarts négatifs vs N-1, "
                                "calculée au même niveau d'agrégation"),
        ("Budget Supeco", "absent de l'export source — affiché « — », exclu des calculs"),
        ("Débit / Panier", "fiables uniquement au niveau Rayon × Site ; jamais sommés "
                           "entre rayons ; pas de budget disponible pour ces indicateurs"),
        ("Montants", "tous arrondis au million (M FCFA)"),
    ]
    r = 3
    for k, v in regles:
        ws.cell(row=r, column=1, value=k).font = bold
        ws.cell(row=r, column=2, value=v).font = base
        r += 1
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============================================================
# Sidebar
# ============================================================

with st.sidebar:
    render_html(f"<div style='font-size:18px; font-weight:900; color:{NOIR};'>📡 PGC Radar</div>"
                f"<div style='font-size:12px; color:{GRIS}; font-weight:600; margin-bottom:12px;'>"
                "SmartBuyer Hub</div>")
    st.markdown("**Import fichier**")
    up = st.file_uploader("Export Département / Rayon / Site (avec Budget)", type=["xlsx"])
    st.caption("Un seul export PBI hebdomadaire suffit.")

# ============================================================
# Landing page
# ============================================================

if not up:
    render_html("<div class='page-title'>📡 PGC Radar</div>")
    render_html("<div class='page-sub'>Pilotage hebdomadaire de la performance PGC — "
                "Vs N-1 et marge en premier plan, budget en référence</div>")

    render_html("""<div class='landing-hero'>
        <div style='font-size:17px; font-weight:900; color:#fff; margin-bottom:4px;'>
            Où en est le réseau, et qui doit creuser quoi&nbsp;?</div>
        <div style='font-size:13px; color:rgba(255,255,255,0.85); font-weight:600;'>
            Dépose l'export PBI hebdomadaire dans la barre latérale : le module produit
            une vue direction, un détail par format, et la liste des alertes à transmettre
            aux acheteurs — avec export Excel prêt à partager.</div>
    </div>""")

    c1, c2 = st.columns(2)
    with c1:
        render_html(f"""<div class='landing-card'>
            <div class='section-title'>🧭 Ce que contient le module</div>
            <div class='rule-row'><span>Vue d'ensemble</span><span class='mut'>CA, marge, benchmark départements</span></div>
            <div class='rule-row'><span>Format</span><span class='mut'>Hyper / Market / Supeco + sites en recul</span></div>
            <div class='rule-row'><span>Alertes</span><span class='mut'>Rayon × Site, causes à saisir</span></div>
            <div class='rule-row'><span>Export Excel</span><span class='mut'>4 onglets, figé, prêt à diffuser</span></div>
        </div>""")
    with c2:
        render_html(f"""<div class='landing-card'>
            <div class='section-title'>⚙️ Comment ça marche</div>
            <div class='rule-row'><span>1. Exporter</span><span class='mut'>PBI, semaine en cours, Dépt/Rayon/Site</span></div>
            <div class='rule-row'><span>2. Déposer</span><span class='mut'>le fichier dans la barre latérale</span></div>
            <div class='rule-row'><span>3. Lire</span><span class='mut'>Vue d'ensemble → Format → Alertes</span></div>
            <div class='rule-row'><span>4. Transmettre</span><span class='mut'>l'Excel avec causes aux acheteurs</span></div>
        </div>""")

    render_html("<div style='height:14px'></div>")
    c3, c4 = st.columns(2)
    with c3:
        render_html(f"""<div class='landing-card'>
            <div class='section-title'>🚨 Règles d'alerte</div>
            <div class='rule-row'><span>Déclencheur</span><span class='mut'>Vs N-1 + marge — jamais le Budget</span></div>
            <div class='rule-row'><span>CA Critique</span><span class='neg'>&lt; -5% vs N-1</span></div>
            <div class='rule-row'><span>CA Attention</span><span class='att' style='color:{ORANGE};'>&lt; -2% vs N-1</span></div>
            <div class='rule-row'><span>Marge Critique</span><span class='neg'>&lt; -1,0 pt vs N-1</span></div>
            <div class='rule-row'><span>Marge Attention</span><span class='att' style='color:{ORANGE};'>&lt; -0,5 pt vs N-1</span></div>
            <div class='rule-row'><span>Matérialité</span><span class='mut'>écart ≥ 1 M FCFA (Rayon × Site)</span></div>
            <div class='rule-row'><span>Poids ds le recul</span><span class='mut'>part des écarts négatifs vs N-1, par niveau</span></div>
        </div>""")
    with c4:
        render_html("""<div class='landing-card'>
            <div class='section-title'>📄 Colonnes attendues (onglet « Export »)</div>
            <span class='col-chip'>Département</span><span class='col-chip'>Rayon</span>
            <span class='col-chip'>Site</span><span class='col-chip'>CA</span>
            <span class='col-chip'>CA N-1</span><span class='col-chip'>Budget</span>
            <span class='col-chip'>Vs N-1 (%)</span><span class='col-chip'>Vs Bgt (%)</span>
            <span class='col-chip'>Marge</span><span class='col-chip'>Marge N-1</span>
            <span class='col-chip'>Taux de Marge N Vs N-1</span>
            <span class='col-chip'>Débit</span><span class='col-chip'>Panier</span>
            <div class='fmt-sub' style='margin-top:10px;'>Budget absent pour Supeco : géré
            automatiquement (« — »). Débit et Panier : fiables au niveau Rayon × Site
            uniquement, Vs N-1 seulement.</div>
        </div>""")

    st.info("Dépose l'export PBI dans la barre latérale pour démarrer.")
    st.stop()

# ============================================================
# Données chargées — calculs
# ============================================================

dept, site, ca_reseau = load_export(up)
pgc = dept[dept["Département"] == "01 - PGC"].iloc[0]

fmt_lvl = agg_niveau(site, "format").set_index("format") \
    .reindex(["Hyper", "Market", "Supeco"]).reset_index()
site_lvl = agg_niveau(site, ["code_site", "site_court", "format"])

al = site.copy()
al["Vs_N1"] = al["Vs N-1 (%)"]
al["Vs_Bgt"] = al["Vs Bgt (%)"]
al["Marge_pt"] = al["Taux de Marge N Vs N-1"]
al["Debit_vs_n1"] = al["Vs N-1 (%).1"]
al["Panier_vs_n1"] = al["Panier N Vs N-1"]
al["ecart_n1"] = al["CA"] - al["CA N-1"]
al["Alerte"] = al.apply(lambda r: alerte(r["Vs_N1"], r["Marge_pt"]), axis=1)
pool_al = al.loc[al["ecart_n1"] < 0, "ecart_n1"].sum()
al["poids"] = al["ecart_n1"].apply(lambda x: x / pool_al if (x < 0 and pool_al < 0) else 0.0)
al_actives = al[(al["Alerte"] != "OK") & (al["ecart_n1"].abs() >= MATERIALITE_FCFA)] \
    .sort_values("ecart_n1")

with st.sidebar:
    st.divider()
    st.download_button(
        "⬇️ Export Excel",
        data=build_excel(pgc, dept, ca_reseau, fmt_lvl, site_lvl, al_actives),
        file_name="PGC_Radar.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
    )

render_html(f"""<div class='hero'>
    <div style='display:flex; justify-content:space-between; align-items:center;'>
        <span class='hero-title'>📡 PGC Radar</span>
        <span class='hero-sub'>Semaine en cours · Vs N-1 et marge en premier plan</span>
    </div>
    <div class='hero-chips'>
        <span class='chip'>CA {fmt_m(pgc["CA"])}</span>
        <span class='chip'>Marge {f"{pgc['Taux de Marge']*100:.1f}%".replace(".", ",")}</span>
        <span class='chip chip-alert'>{len(al_actives)} alertes</span>
    </div>
</div>""")

tab1, tab2, tab3 = st.tabs(["Vue d'ensemble", "Format", "Alertes"])

# ---------------- Écran 1 — Vue d'ensemble ----------------
with tab1:
    c1, c2 = st.columns(2)
    with c1:
        d = pgc["Vs N-1 (%)"]
        pill = "pill-neg" if d < 0 else "pill-pos"
        render_html(f"""<div class='kpi-card kpi-blue'>
            <div class='kpi-top'>
                <span class='kpi-label lab-blue'>CA PGC</span>
                <div class='kpi-ico ico-blue'>💶</div>
            </div>
            <div class='kpi-value val-blue'>{fmt_m(pgc["CA"])}</div>
            <span class='pill {pill}'>{'▼' if d < 0 else '▲'} {fmt_pct(abs(d))[1:]} vs N-1</span>
        </div>""")
    with c2:
        mpt = pgc["Taux de Marge N Vs N-1"]
        pill = "pill-neg" if mpt < 0 else "pill-pos"
        render_html(f"""<div class='kpi-card kpi-violet'>
            <div class='kpi-top'>
                <span class='kpi-label lab-violet'>Taux de marge</span>
                <div class='kpi-ico ico-violet'>📈</div>
            </div>
            <div class='kpi-value val-violet'>{f"{pgc['Taux de Marge']*100:.1f}%".replace(".", ",")}</div>
            <span class='pill {pill}'>{'▼' if mpt < 0 else '▲'} {fmt_pt(abs(mpt))[1:]} vs N-1</span>
        </div>""")

    render_html(f"""<div class='ref-line'>
        <span>Budget (référence — non déclencheur d'alerte)</span>
        <span>{fmt_m(pgc["Budget"])} · écart {fmt_pct(pgc["Vs Bgt (%)"])}</span>
    </div>""")

    autres_dept = dept[dept["Département"] != "01 - PGC"]
    ca_max = autres_dept["CA"].max()
    bench_html = ""
    for _, r in autres_dept.iterrows():
        d = r["Vs N-1 (%)"]
        width = max(4, r["CA"] / ca_max * 100)
        bench_html += (f"<div class='bench-item'>"
                       f"<div class='bench-line'><span>{r['Département'].title()} "
                       f"<span class='mut'>{fmt_m(r['CA'])}</span></span>"
                       f"<span class='{cls(d)}'>{fmt_pct(d)} N-1</span></div>"
                       f"<div class='bench-track'><div class='bench-fill' "
                       f"style='width:{width:.0f}%;'></div></div></div>")
    bench_html += (f"<div class='bench-line' style='margin-top:10px; padding-top:10px; "
                   f"border-top:0.5px solid #F0F0F3;'><span>Total magasin</span>"
                   f"<span class='mut'>{fmt_m(ca_reseau)}</span></div>")
    render_html("<div class='section-title'>Benchmark départements</div>")
    render_html(f"<div class='bench-card'>{bench_html}</div>")

# ---------------- Écran 2 — Format ----------------
with tab2:
    RAMP = ["#E24B4A", "#F0997B", "#F5C4B3", "#FAD9CE"]
    for _, f in fmt_lvl.iterrows():
        b_cls, b_txt = BADGE[alerte(f["Vs_N1"], f["Marge_pt"])]
        sites_f = site_lvl[site_lvl["format"] == f["format"]].sort_values("ecart_n1")
        gros = sites_f[sites_f["poids"] >= SEUIL_POIDS_SITE]
        autres = sites_f[sites_f["poids"] < SEUIL_POIDS_SITE]

        negs = sites_f[sites_f["ecart_n1"] < 0]
        bar_html = ""
        if len(negs):
            total_neg = negs["ecart_n1"].sum()
            segs, cum = "", 0.0
            for i, (_, s_) in enumerate(negs.iterrows()):
                part = s_["ecart_n1"] / total_neg * 100
                cum += part
                segs += (f"<div style='width:{part:.0f}%; "
                         f"background:{RAMP[min(i, len(RAMP)-1)]};'></div>")
            if cum < 99:
                segs += f"<div style='width:{100-cum:.0f}%; background:#E5E5EA;'></div>"
            bar_html = f"<div class='contrib-bar'>{segs}</div>"

        sites_html = ""
        for _, s_ in gros.iterrows():
            sites_html += (f"<div class='site-row'><span>{s_['site_court']}</span>"
                           f"<span><span class='{cls(s_['Vs_N1'])}'>{fmt_pct(s_['Vs_N1'])}</span>"
                           f" · {fmt_m(s_['ecart_n1'])}"
                           f" · <b>{s_['poids']*100:.0f}% du recul réseau</b></span></div>")
        if len(autres) and autres["poids"].sum() > 0:
            sites_html += (f"<div class='autres'>+ {len(autres)} autres sites, "
                           f"poids cumulé {autres['poids'].sum()*100:.0f}% du recul</div>")
        elif not len(gros):
            sites_html += "<div class='autres'>Aucun site significatif dans le recul</div>"

        render_html(f"""<div class='fmt-block'>
            <div class='fmt-head'>
                <span class='fmt-name'>{f['format']}
                    <span class='fmt-sub'>{int(f['nb_sites'])} sites · {fmt_m(f['CA'])}</span></span>
                <span class='badge {b_cls}'>{b_txt}</span>
            </div>
            <div class='fmt-sub'>Vs N-1 <span class='{cls(f['Vs_N1'])}'>{fmt_pct(f['Vs_N1'])}</span>
                · Vs Budget <span class='mut'>{fmt_pct(f['Vs_Bgt'])}</span></div>
            {bar_html}
            {sites_html}
        </div>""")

# ---------------- Écran 3 — Alertes ----------------
with tab3:
    render_html(f"<div class='fmt-sub' style='margin-bottom:10px;'>"
                f"{len(al_actives)} lignes en alerte · déclencheur N-1 + marge · "
                f"écart ≥ {MATERIALITE_FCFA/1e6:.0f} M FCFA</div>")

    if al_actives.empty:
        st.success("Aucune alerte cette semaine.")
    else:
        for _, r in al_actives.iterrows():
            b_cls, b_txt = BADGE[r["Alerte"]]
            crit = r["Alerte"] == "Critique"
            side = "alert-crit" if crit else "alert-att"
            av = "av-crit" if crit else "av-att"
            site_court = r["Site"].split(" - ")[1]
            nom_court = site_court.split(" ", 1)[1] if " " in site_court else site_court
            initiales = "".join(w[0] for w in nom_court.split()[:2]).upper()
            rayon_court = r["Rayon"].split(" - ")[1].title()
            render_html(f"""<div class='alert-card {side}'>
                <div class='alert-flex'>
                    <span class='avatar {av}'>{initiales}</span>
                    <div style='flex:1;'>
                        <div class='alert-head'>
                            <span class='alert-title'>{nom_court} — {rayon_court}</span>
                            <span class='badge {b_cls}'>{b_txt}</span>
                        </div>
                        <div class='alert-poids'>{r['poids']*100:.0f}% du recul réseau</div>
                    </div>
                </div>
                <div class='alert-line'>CA {fmt_m(r['CA'])} ·
                    <span class='{cls(r['Vs_N1'])}'>{fmt_pct(r['Vs_N1'])} N-1</span> ·
                    <span class='mut'>Budget {fmt_pct(r['Vs_Bgt'])}</span> ·
                    Marge <span class='{cls(r['Marge_pt'])}'>{fmt_pt(r['Marge_pt'])}</span></div>
                <div class='alert-line'>Débit <span class='{cls(r['Debit_vs_n1'])}'>{fmt_pct(r['Debit_vs_n1'])} N-1</span>
                    · Panier <span class='{cls(r['Panier_vs_n1'])}'>{fmt_pct(r['Panier_vs_n1'])} N-1</span></div>
            </div>""")

        with st.expander("✏️ Saisir les causes (à transmettre aux acheteurs)"):
            saisie = al_actives[["Site", "Rayon", "Alerte"]].copy()
            saisie["Écart (M)"] = (al_actives["ecart_n1"] / 1e6).round(1)
            saisie["Cause"] = CAUSES[0]
            edited = st.data_editor(
                saisie, hide_index=True, width="stretch",
                column_config={"Cause": st.column_config.SelectboxColumn(options=CAUSES)},
                key="causes_editor",
            )
            csv = edited.to_csv(index=False).encode("utf-8-sig")
            st.download_button("Télécharger la liste (CSV)", data=csv,
                               file_name="alertes_pgc_radar.csv", mime="text/csv")
