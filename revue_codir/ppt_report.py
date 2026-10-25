"""Rapport PowerPoint à la Direction Générale, généré sur le masque de l'entreprise.

build_report(R, masque_path, out) : 5 pages principales + annexe (A1 à A4).
Les chiffres viennent tous de `R` (engine.compute) ; les « Causes et actions » de engine.merge_comments.
"""
from __future__ import annotations

import copy

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import engine as E

NAVY, ORANGE, INK, MUTE = '24235C', 'F19F39', '1E2140', '6B7090'
NEG, POS, SMALL, GREY, LIGHT = 'C0392B', '1E8449', 'B5B9CC', 'ECEDF2', 'F3F4F8'
TNEG, TPOS = 'B3261E', '1B7A45'


def rgb(h):
    return RGBColor.from_string(h)


# ------------------------------------------------------------------ outils de mise en page
def _drop_all_slides(prs):
    sld = prs.slides._sldIdLst
    for s in list(sld):
        prs.part.drop_rel(s.get(qn('r:id')))
        sld.remove(s)


def _layout(prs, name):
    for l in prs.slide_layouts:
        if l.name == name:
            return l
    raise KeyError(name)


def _clone_footer(slide, layout, footer_text):
    """Recopie le pied de page et le numéro de page du masque (python-pptx ne les crée pas)."""
    for ph in layout.placeholders:
        t = ph.placeholder_format.type
        if t is None:
            continue
        name = str(t)
        if 'FOOTER' in name or 'SLIDE_NUMBER' in name:
            el = copy.deepcopy(ph._element)
            slide.shapes._spTree.append(el)
            if 'FOOTER' in name:
                sp = slide.shapes[-1]
                tf = sp.text_frame
                tf.text = footer_text


def _title(slide, light, bold):
    t = slide.shapes.title
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = (light + ' : ').upper() if light else ''
    r1.font.bold = False
    r2 = p.add_run()
    r2.text = bold.upper()
    r2.font.bold = True
    size = 26 if len(r1.text + r2.text) <= 48 else 22
    r1.font.size = r2.font.size = Pt(size)


def tb(slide, x, y, w, h, runs, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, wrap=True, spacing=None):
    """runs : texte simple, ou liste de paragraphes, chaque paragraphe = liste de (texte, {options})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for text, o in para:
            r = p.add_run()
            r.text = text
            f = r.font
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.italic = o.get('italic', italic)
            f.color.rgb = rgb(o.get('color', color))
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def label_caps(slide, x, y, w, text, color=NAVY, size=9.5):
    return tb(slide, x, y, w, 0.25, [[(text.upper(), {'bold': True, 'color': color, 'size': size})]])


def table(slide, x, y, w, rows, col_w, header, size=10, row_h=0.3, aligns=None, fills=None, colors=None, bolds=None):
    nr, nc = len(rows) + 1, len(header)
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for attr in ('bandRow', 'firstRow'):
        tblPr.set(attr, '0')
    style = tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # aucun style
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    aligns = aligns or (['l'] + ['r'] * (nc - 1))

    def cell(i, j, text, fill, color, bold, italic=False):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(fill)
        c.margin_left = c.margin_right = Inches(0.06)
        c.margin_top = c.margin_bottom = Inches(0.02)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = {'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT, 'c': PP_ALIGN.CENTER}[aligns[j]]
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = rgb(color)

    for j, h in enumerate(header):
        cell(0, j, h, ORANGE, 'FFFFFF', True, True)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            fill = (fills[i - 1][j] if fills and fills[i - 1][j] else (GREY if i % 2 == 1 else 'FFFFFF'))
            col = colors[i - 1][j] if colors and colors[i - 1][j] else INK
            b = bolds[i - 1][j] if bolds else (j == 0)
            cell(i, j, v, fill, col, b)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
    return gt


def sign_color(v, small=False):
    if small:
        return MUTE
    return TNEG if v < 0 else TPOS


def heat(v):
    """Couleur de fond et de texte pour un écart de marge (FCFA)."""
    if abs(v) < E.SEUIL_SIGNIF:
        return GREY, MUTE
    t = min(abs(v) / 5e6, 1)
    base = (192, 57, 43) if v < 0 else (30, 132, 73)
    a = 0.18 + 0.67 * t
    mix = tuple(int(255 + (c - 255) * a) for c in base)
    fill = '%02X%02X%02X' % mix
    return fill, ('FFFFFF' if abs(v) > 3e6 else ('8E1F17' if v < 0 else '145A32'))


def diverging_rows(slide, x, y, w, rows, vmax, name_w=2.2, val_w=0.65, row_h=0.42, size=10.5, sub_size=8.5):
    """rows : [{'n','why','v'}] ; barres horizontales de part et d'autre de zéro (dessinées en formes)."""
    bar_w = w - name_w - val_w
    zero = x + name_w + bar_w / 2
    for i, r in enumerate(rows):
        yy = y + i * row_h
        tb(slide, x, yy, name_w, row_h, [[(r['n'], {'bold': True, 'size': size})], [(r['why'], {'size': sub_size, 'color': MUTE})]], anchor=MSO_ANCHOR.MIDDLE)
        v = r['v']
        L = abs(v) / vmax * (bar_w / 2 - 0.05)
        small = abs(v) < E.SEUIL_SIGNIF
        col = SMALL if small else (NEG if v < 0 else POS)
        if L > 0.01:
            rect(slide, zero - L if v < 0 else zero, yy + row_h * 0.3, L, row_h * 0.4, col)
        tb(slide, x + name_w + bar_w, yy, val_w, row_h, [[(E.fdm(v), {'bold': True, 'color': sign_color(v, small), 'size': size})]],
           align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    ln = slide.shapes.add_connector(1, Inches(zero), Inches(y), Inches(zero), Inches(y + row_h * len(rows)))
    ln.line.color.rgb = rgb('9AA0BC')
    ln.line.width = Pt(0.75)


# ------------------------------------------------------------------ pages
def _cover(prs, R):
    lay = _layout(prs, '14_Couverture')
    s = prs.slides.add_slide(lay)
    t = s.shapes.title.text_frame
    t.clear()
    p = t.paragraphs[0]
    r = p.add_run()
    r.text = 'REVUE DE PERFORMANCE '
    r.font.bold = False
    r = p.add_run()
    r.text = 'COMMERCIALE'
    r.font.bold = True
    m = R['meta']
    box = tb(s, 0.94, 4.78, 7.6, 0.36, [[(f"Semaine {m['week']} · {m['periode']} · Rapport à la Direction Générale", {'italic': True, 'color': NAVY, 'size': 13})]],
             anchor=MSO_ANCHOR.MIDDLE)
    box.fill.solid()
    box.fill.fore_color.rgb = rgb('FFFFFF')
    box.text_frame.auto_size = None
    tb(s, 0.94, 5.3, 6.0, 0.6, [[(f"Données au lundi {m['data_date']:%d/%m/%Y}, 08:00 · Vs N-1 · Préparé par : Category management · Validé par : [nom]",
                                   {'size': 10, 'color': 'C5C8E0'})], [('Confidentiel · usage interne', {'size': 10, 'color': 'C5C8E0'})]])
    return s


def _content(prs, R, light, bold, footer):
    lay = _layout(prs, 'Titre seul')
    s = prs.slides.add_slide(lay)
    _title(s, light, bold)
    _clone_footer(s, lay, footer)
    return s


def _p2(prs, R, ftr):
    g, t = R['global'], R['texts']
    s = _content(prs, R, 'Synthèse', t['p2_title'], ftr)
    K = [("Chiffre d'affaires", f"{E.fm(g['CA'])} M", g['ca'], E.fpct(g['ca']), f"{E.fm(g['CA1'])} M"),
         ('Marge', f"{E.fm(g['M'])} M", g['dM'], f"{E.fdm(g['dM'])} M", f"{E.fm(g['M1'])} M"),
         ('Taux de marge', f"{E.fnum(g['t'] * 100)} %", g['dt'], E.fpt(g['dt']), f"{E.fnum(g['t1'] * 100)} %"),
         ('Trafic (tickets)', E.fnum(g['T'], 0), g['tk'], E.fpct(g['tk']), E.fnum(g['T1'], 0)),
         ('Panier moyen', f"{E.fnum(g['panier'], 0)} FCFA", g['pan'], E.fpct(g['pan']), f"{E.fnum(g['panier1'], 0)} FCFA")]
    x0, w, gap = 0.61, 2.34, 0.1
    for i, (lab, v, d, ds, n1) in enumerate(K):
        x = x0 + i * (w + gap)
        rect(s, x, 1.45, w, 1.1, LIGHT)
        rect(s, x, 2.5, w, 0.05, NEG if d < 0 else POS)
        tb(s, x + 0.08, 1.5, w - 0.16, 0.25, [[(lab.upper(), {'bold': True, 'size': 9, 'color': '4E5370'})]])
        tb(s, x + 0.08, 1.72, w - 0.16, 0.4, [[(v, {'bold': True, 'size': 19, 'color': NAVY})]])
        tb(s, x + 0.08, 2.12, w - 0.16, 0.36, [[(('▼ ' if d < 0 else '▲ ') + ds, {'bold': True, 'size': 11, 'color': sign_color(d)}),
                                                 (f'   N-1 : {n1}', {'size': 9, 'color': MUTE})]])
    # où en est-on / pourquoi
    label_caps(s, 0.61, 2.8, 5.9, 'Où en est-on, pourquoi')
    paras = []
    for f_ in t['faits']:
        head, _, rest = f_.partition(':') if f_.startswith('Pourquoi') or f_.startswith('Secteurs') else ('', '', f_)
        if head:
            paras.append([('●  ', {'color': ORANGE, 'size': 10}), (head + ' :', {'bold': True}), (rest, {})])
        else:
            paras.append([('●  ', {'color': ORANGE, 'size': 10}), (f_, {'bold': True})])
    paras.append([('●  ', {'color': ORANGE, 'size': 10}), (t['taux'], {})])
    tb(s, 0.61, 3.1, 5.9, 2.6, paras, size=12, spacing=7)
    # ce que l'on fait
    label_caps(s, 6.85, 2.8, 5.85, "Ce que l'on fait")
    flops = []
    for sct in E.SECTEUR_ORDER:
        for f_ in R['familles'][sct]['flop'][:3]:
            flops.append((sct, f_))
    flops = sorted(flops, key=lambda x: x[1]['dM'])[:3]
    yy = 3.12
    for sct, f_ in flops:
        c = f_.get('cdg', {})
        act = c.get('Action') or f"En attente : présenté en séance par {E.CHEFS[sct]}"
        who = c.get('Responsable') or E.CHEFS[sct]
        tb(s, 6.85, yy, 4.4, 0.55, [[(f"{f_['label']} · {E.fdm(f_['dM'])} M", {'bold': True, 'size': 11})],
                                     [(act, {'size': 9.5, 'italic': not c.get('Action'), 'color': INK if c.get('Action') else MUTE})]])
        tb(s, 11.25, yy, 1.45, 0.55, [[(who, {'size': 9, 'color': MUTE})]], align=PP_ALIGN.RIGHT)
        ln = s.shapes.add_connector(1, Inches(6.85), Inches(yy + 0.6), Inches(12.7), Inches(yy + 0.6))
        ln.line.color.rgb = rgb(GREY)
        yy += 0.66
    decs = [f_ for sct in E.SECTEUR_ORDER for f_ in R['familles'][sct]['flop'] if f_.get('cdg', {}).get('Décision DG', '').lower() == 'oui']
    box = rect(s, 6.85, 5.18, 5.85, 0.55, NAVY)
    dtxt = (f"{len(decs)} décision(s) demandée(s) : " + ', '.join(f_['label'] for f_ in decs[:3]) + ' (page 5).') if decs else 'Aucune décision demandée cette semaine (page 5).'
    tb(s, 6.95, 5.2, 5.7, 0.5, [[('De quoi avons-nous besoin : ', {'bold': True, 'color': ORANGE, 'size': 9.5}), (dtxt, {'color': 'FFFFFF', 'size': 9.5})]],
       anchor=MSO_ANCHOR.MIDDLE)
    # contexte
    rect(s, 0.61, 5.95, 12.1, 0.85, LIGHT)
    rect(s, 0.61, 5.95, 0.05, 0.85, ORANGE)
    verif = ', '.join(f"{k} ({v})" for k, v in E.SITES_A_VERIFIER.items())
    tb(s, 0.75, 5.99, 11.9, 0.8, [[('Contexte : ', {'bold': True}), (f"{t['contexte']}. ", {}),
                                    ('Périmètre : ', {'bold': True}), (f"12 magasins ouverts en N-1 ; Ex-Supeco convertis en Market : " + ', '.join(f"{n.replace('Market ', '')} {d[:5]}" for n, d in E.CONVERSIONS) + '. ', {}),
                                    ('À vérifier : ', {'bold': True, 'color': '8A4B00'}), (verif + '. ', {}),
                                    ('Référence : ', {'bold': True}), ('Vs N-1 uniquement, budget suivi au reporting mensuel.', {})]], size=10)


def _p3(prs, R, ftr):
    s = _content(prs, R, 'Marge', R['texts']['p3_title'], ftr)
    label_caps(s, 0.61, 1.45, 6.2, 'Passage de la marge N-1 → N (M FCFA)')
    g = R['global']
    steps = [('Marge N-1', g['M1'], 'tot')] + [(sct if sct != 'Non-alimentaire' else 'Non-alim.', R['sect'][sct]['dM'], 'd') for sct in E.SECTEUR_ORDER] + [('Marge N', g['M'], 'tot')]
    cats, base, up, cols, labels = [], [], [], [], []
    cur = g['M1']
    for n, v, k in steps:
        cats.append(n)
        if k == 'tot':
            base.append(0)
            up.append(v / 1e6)
            cols.append('758CC0' if n == 'Marge N-1' else NAVY)
            labels.append(E.fm(v))
        else:
            lo = min(cur, cur + v)
            base.append(lo / 1e6)
            up.append(abs(v) / 1e6)
            cols.append(NEG if v < 0 else POS)
            labels.append(E.fdm(v))
            cur += v
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series('base', base)
    cd.add_series('valeur', up)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.61), Inches(1.75), Inches(6.2), Inches(4.9), cd)
    ch = gf.chart
    ch.has_legend = False
    lo_axis = min(g['M1'], g['M'], *(c for c in [g['M1'] + sum(R['sect'][x]['dM'] for x in E.SECTEUR_ORDER[:i]) for i in range(4)])) / 1e6
    ch.value_axis.minimum_scale = int(lo_axis * 0.95 // 5 * 5)
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.visible = False
    ch.category_axis.tick_labels.font.size = Pt(10)
    ch.category_axis.format.line.color.rgb = rgb('9AA0BC')
    plot = ch.plots[0]
    plot.gap_width = 60
    sb, sv = plot.series
    sb.format.fill.background()
    for i, c in enumerate(cols):
        pt = sv.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        dl = pt.data_label
        dl.has_text_frame = True
        dl.text_frame.text = labels[i]
        dl.position = XL_LABEL_POSITION.INSIDE_END if i in (0, len(cols) - 1) else XL_LABEL_POSITION.CENTER
        r = dl.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = rgb('FFFFFF')
    tb(s, 0.61, 6.6, 6.2, 0.3, [[(f"≈ Axe tronqué à {ch.value_axis.minimum_scale:.0f} M FCFA pour rendre les écarts visibles", {'size': 9, 'color': MUTE, 'italic': True})]])
    # heatmap
    label_caps(s, 7.1, 1.45, 5.6, 'Écart de marge secteur × format (M FCFA)')
    hdr = ['Secteur'] + E.FORMAT_ORDER + ['Total']
    rows, fills, colors, bolds = [], [], [], []
    for sct in E.SECTEUR_ORDER + ['Total']:
        vals = [R['mtx'][sct][f] for f in E.FORMAT_ORDER] if sct != 'Total' else [R['fmt'][f]['dM'] for f in E.FORMAT_ORDER]
        tot = R['sect'][sct]['dM'] if sct != 'Total' else g['dM']
        vals.append(tot)
        rows.append([sct if sct != 'Non-alimentaire' else 'Non-alim.'] + [E.fdm(v) for v in vals])
        hs = [heat(v) for v in vals]
        fills.append([None] + [h[0] for h in hs])
        colors.append([None] + [h[1] for h in hs])
        bolds.append([sct == 'Total'] + [True] * len(vals))
    table(s, 7.1, 1.8, 5.6, rows, [1.25, 0.87, 0.87, 0.87, 0.87, 0.87], hdr, size=11, row_h=0.5, fills=fills, colors=colors, bolds=bolds)
    tb(s, 7.1, 4.4, 5.6, 0.5, [[('Gris : écart inférieur à 0,5 M FCFA, non significatif. Ex-Supeco : Cité Verte, Aboboté, Toit Rouge.', {'size': 9, 'color': MUTE})]])
    # tableau formats
    label_caps(s, 7.1, 4.95, 5.6, 'Formats : CA, trafic et panier Vs N-1')
    frows = [[f, E.fpct(R['fmt'][f]['ca']), E.fpct(R['fmt'][f]['tk']), E.fpct(R['fmt'][f]['pan'])] for f in E.FORMAT_ORDER]
    fcol = [[None] + [sign_color(R['fmt'][f][k]) for k in ('ca', 'tk', 'pan')] for f in E.FORMAT_ORDER]
    table(s, 7.1, 5.25, 5.6, frows, [1.7, 1.3, 1.3, 1.3], ['Format', 'CA', 'Trafic', 'Panier'], size=10, row_h=0.3, colors=fcol)


def _p4(prs, R, ftr):
    s = _content(prs, R, 'Écarts', R['texts']['p4_title'], ftr)
    C = R['contrib']
    vmax = max(abs(r['v']) for r in C['rayon'] + C['site']) * 1.05
    label_caps(s, 0.61, 1.5, 5.8, 'Par rayon')
    diverging_rows(s, 0.61, 1.85, 5.8, C['rayon'], vmax, name_w=2.5, row_h=0.72, size=12, sub_size=9.5)
    label_caps(s, 6.9, 1.5, 5.8, 'Par magasin')
    diverging_rows(s, 6.9, 1.85, 5.8, C['site'], vmax, name_w=2.5, row_h=0.72, size=12, sub_size=9.5)
    tb(s, 0.61, 6.35, 12.1, 0.4, [[('Écart de marge Vs N-1, M FCFA. Les deux listes ne s’additionnent pas : un rayon et un magasin peuvent porter la même perte.', {'size': 9.5, 'color': MUTE, 'italic': True})]])


def _p5(prs, R, ftr):
    s = _content(prs, R, "Plan d'actions", 'engagements et décisions attendues', ftr)
    rows, cols = [], []
    filled, empty = [], []
    for sct in E.SECTEUR_ORDER:
        for i, f_ in enumerate(R['familles'][sct]['flop']):
            c = f_.get('cdg', {})
            item = (sct, f_, c)
            if c.get('Action') or c.get('Commentaire'):
                filled.append(item)
            elif i < 2:
                empty.append(item)
    items = sorted(filled, key=lambda x: (x[2].get('Décision DG', '').lower() != 'oui', x[1]['dM']))[:8]
    items += sorted(empty, key=lambda x: x[1]['dM'])[:max(0, 8 - len(items))]
    for sct, f_, c in items:
        short = {'Non-alimentaire': 'Non-alim.'}.get(sct, sct)
        rows.append([f"{f_['label']} ({short})", E.fdm(f_['dM']), c.get('Cause') or '—', c.get('Action') or f"En attente : {E.CHEFS[sct]}",
                     c.get('Responsable') or '—', c.get('Échéance') or '—', c.get('Décision DG') or '—'])
        grey = MUTE
        cols.append([None, sign_color(f_['dM']), None if c.get('Cause') else grey, None if c.get('Action') else grey,
                     None if c.get('Responsable') else grey, None if c.get('Échéance') else grey,
                     (TNEG if (c.get('Décision DG') or '').lower() == 'oui' else None) if c.get('Décision DG') else grey])
    table(s, 0.61, 1.5, 12.1, rows, [2.8, 0.85, 1.6, 3.6, 1.35, 0.95, 0.95],
          ['Famille', 'Écart (M)', 'Cause', 'Action', 'Responsable', 'Échéance', 'Décision DG'], size=10, row_h=0.5,
          aligns=['l', 'r', 'l', 'l', 'l', 'c', 'c'], colors=cols)
    tb(s, 0.61, 1.6 + 0.5 * (len(rows) + 1), 12.1, 0.4,
       [[('Repris des « Causes et actions » saisies par les chefs de groupe, gel à 10:10. Les lignes « Décision DG : Oui » sont en tête et remontent en page 2.', {'size': 9.5, 'color': MUTE, 'italic': True})]])


def _annex_sector(prs, R, sct, code, ftr):
    s = _content(prs, R, sct if sct != 'Non-alimentaire' else 'Non-alim.', R['texts']['sect_title'][sct], ftr)
    o = R['sect_detail'][sct]
    label_caps(s, 0.61, 1.45, 5.0, "Chiffre d'affaires et marge (Vs N-1)")
    rows = [['CA (M FCFA)', E.fm(o['CA']), E.fm(o['CA1']), E.fdm(o['dCA'])],
            ['Marge (M FCFA)', E.fm(o['M']), E.fm(o['M1']), E.fdm(o['dM'])],
            ['Taux de marge', f"{E.fnum(o['t'] * 100)} %", f"{E.fnum(o['t1'] * 100)} %", E.fpt(o['dt'])]]
    cols = [[None, None, MUTE, sign_color(o['dCA'])], [None, None, MUTE, sign_color(o['dM'])], [None, None, MUTE, sign_color(o['dt'])]]
    bolds = [[True, True, False, True]] * 3
    table(s, 0.61, 1.75, 5.0, rows, [1.7, 1.1, 1.1, 1.1], ['', 'Réel', 'N-1', 'Δ'], size=11, row_h=0.36, colors=cols, bolds=bolds)
    label_caps(s, 0.61, 3.35, 5.0, 'Écart de marge par format (M FCFA)')
    for i, f in enumerate(E.FORMAT_ORDER):
        v = R['familles'][sct]['fmt'][f]
        small = abs(v) < E.SEUIL_SIGNIF
        fill = GREY if small else ('FBEAE8' if v < 0 else 'E7F4EC')
        x = 0.61 + i * 1.27
        rect(s, x, 3.65, 1.2, 0.72, fill)
        tb(s, x, 3.67, 1.2, 0.3, [[(f, {'size': 9.5, 'color': '3A3F5C'})]], align=PP_ALIGN.CENTER)
        tb(s, x, 3.92, 1.2, 0.4, [[(E.fdm(v), {'size': 14, 'bold': True, 'color': sign_color(v, small)})]], align=PP_ALIGN.CENTER)
    tb(s, 0.61, 4.5, 5.0, 0.9, [[(R['texts']['lecture'][sct], {'size': 10, 'color': '33385A'})]])
    # familles
    label_caps(s, 6.0, 1.45, 6.7, 'Familles · écart de marge (dont volume / dont taux)')
    fam = R['familles'][sct]
    sel = fam['top'][:3] + fam['flop'][:3][::-1]
    vmax = max([abs(f_['dM']) for f_ in sel] + [1]) * 1.05
    rws = [{'n': f_['label'], 'why': f"volume {E.fdm(f_['effV'])} · taux {E.fdm(f_['effT'])}", 'v': f_['dM']} for f_ in sel]
    diverging_rows(s, 6.0, 1.78, 6.7, rws, vmax, name_w=2.9, row_h=0.6, size=11, sub_size=9)
    # causes et actions
    rect(s, 0.61, 5.55, 12.1, 1.3, 'F7F8FB', line='C9CCDB')
    gen = R.get('general', {}).get(sct) if isinstance(R.get('general'), dict) else None
    paras = [[(f'Causes et actions · {E.CHEFS[sct]}', {'bold': True, 'color': NAVY, 'size': 11})]]
    lines = [f_ for f_ in fam['flop'][:3] if f_.get('cdg', {}).get('Commentaire') or f_.get('cdg', {}).get('Action')]
    if gen:
        paras.append([(gen, {'size': 10})])
    for f_ in lines:
        c = f_['cdg']
        paras.append([(f"{f_['label']} : ", {'bold': True, 'size': 10}), (' '.join(x for x in [c.get('Cause') and f"[{c['Cause']}]", c.get('Commentaire'), c.get('Action') and f"Action : {c['Action']}"] if x), {'size': 10})])
    if not gen and not lines:
        paras.append([(f'En attente : présenté en séance par {E.CHEFS[sct]}.', {'italic': True, 'color': MUTE, 'size': 10.5})])
    tb(s, 0.72, 5.6, 11.9, 1.22, paras, spacing=2)


def _annex_lexique(prs, R, ftr):
    s = _content(prs, R, 'Annexe', 'lexique et méthode', ftr)
    lex = [('CA', 'ventes HT de la semaine, en M FCFA'), ('Marge', 'CA − coût d’achat des produits vendus ; marge N-1 = hors promo + promo'),
           ('Taux de marge', 'marge ÷ CA'), ('Trafic', 'nombre de tickets de caisse (débit)'), ('Panier moyen', 'CA ÷ tickets, en FCFA'),
           ('Vs N-1', 'même semaine de l’année précédente'), ('Δ (écart)', 'valeur N − valeur N-1'), ('pt (point)', 'écart entre deux taux'),
           ('Effet volume', '(CA N − CA N-1) × taux N-1 ; inclut l’effet prix'), ('Effet taux', 'CA N × (taux N − taux N-1)'),
           ('Non significatif', 'écart de marge inférieur à 0,5 M FCFA, affiché en gris'),
           ('Ex-Supeco', 'anciens Supeco convertis en Market : ' + ', '.join(f"{n.replace('Market ', '')} ({d})" for n, d in E.CONVERSIONS)),
           ('Familles', 'classées par écart de marge en valeur, CA ≥ 1 M FCFA en N ou N-1'),
           ('Sources', 'exports Power BI A (famille × site), B (rayon × site), C (site) ; casse et budget hors périmètre')]
    half = (len(lex) + 1) // 2
    for col, chunk in enumerate([lex[:half], lex[half:]]):
        tb(s, 0.61 + col * 6.2, 1.55, 5.9, 5.2, [[(k + ' : ', {'bold': True, 'color': NAVY}), (v, {})] for k, v in chunk], size=12, spacing=9)


def build_report(R: dict, masque, out) -> None:
    prs = Presentation(masque)
    _drop_all_slides(prs)
    sem = f"S{R['meta']['week']}"
    ftr = f'Revue de performance · {sem}'
    _cover(prs, R)
    _p2(prs, R, ftr)
    _p3(prs, R, ftr)
    _p4(prs, R, ftr)
    _p5(prs, R, ftr)
    for code, sct in zip(['A1', 'A2', 'A3'], E.SECTEUR_ORDER):
        _annex_sector(prs, R, sct, code, f'Annexe {code} · {ftr}')
    _annex_lexique(prs, R, f'Annexe A4 · {ftr}')
    prs.save(out)
